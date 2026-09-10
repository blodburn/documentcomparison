from __future__ import annotations

import base64
import html
import io
import json
import os
import re
import threading
import uuid
import unicodedata
from dataclasses import dataclass, asdict
from difflib import SequenceMatcher

try:
    from rapidfuzz.distance import Indel as _RapidFuzzIndel
except Exception:
    _RapidFuzzIndel = None
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from typing import List, Optional, Dict, Tuple, Any

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
import xlsxwriter

HOST = "127.0.0.1"
PORT = 8765
MAX_BODY = 80 * 1024 * 1024  # 80 MB total request
RESULT_CACHE: Dict[str, Dict[str, Any]] = {}

ARTICLE_RE = re.compile(r"^\s*제\s*(\d+)\s*조(?:\s*의\s*(\d+))?\s*(?:\(([^\n\)]{1,120})\))?\s*(.*)$")
SECTION_RE = re.compile(r"^\s*제\s*(\d+)\s*(장|절|관)\s*(.*)$")
CIRCLED_RE = re.compile(r"[①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳]")
HO_RE = re.compile(r"(?m)^\s*\d+\.\s*")
MOK_RE = re.compile(r"(?m)^\s*[가-하]\.\s*")
NUMBER_RE = re.compile(
    r"(?:\d{4}\s*[./-]\s*\d{1,2}(?:\s*[./-]\s*\d{1,2})?|\d+(?:,\d{3})*(?:\.\d+)?\s*(?:%|퍼센트|원|만원|억원|조원|일|개월|년|시간|분|회|배))"
)

LEGAL_TERMS = [
    "하여야 한다", "해야 한다", "할 수 있다", "할 수 없다", "아니한다", "금지", "책임",
    "면책", "손해배상", "위약", "해지", "해제", "자동갱신", "관할", "준거법", "통지",
    "동의", "승인", "사전", "사후", "및", "또는", "이상", "초과", "이하", "미만"
]

@dataclass
class Unit:
    index: int
    kind: str
    number: str
    title: str
    header: str
    body: str
    text: str
    section: str
    norm_title: str
    norm_body: str
    structure: Tuple[int, int, int]


def iter_docx_blocks(doc: Document):
    """Yield paragraphs and tables in document order."""
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield Table(child, doc)


def extract_docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    lines: List[str] = []
    for block in iter_docx_blocks(doc):
        if isinstance(block, Paragraph):
            txt = block.text.strip()
            if txt:
                lines.append(txt)
        else:
            for row in block.rows:
                cells = [re.sub(r"\s+", " ", c.text.strip()) for c in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
    return "\n".join(lines)


def decode_txt(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp949", "euc-kr"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def read_document(name: str, data: bytes) -> str:
    ext = os.path.splitext(name)[1].lower()
    if ext == ".txt":
        return decode_txt(data)
    if ext == ".docx":
        return extract_docx_text(data)
    raise ValueError(f"지원하지 않는 파일 형식입니다: {ext or '(확장자 없음)'}")


def normalize_text(s: str) -> str:
    s = re.sub(r"제\s*\d+\s*조(?:\s*의\s*\d+)?", " ", s)
    s = re.sub(r"\s+", "", s)
    s = re.sub(r"[^0-9A-Za-z가-힣%]", "", s)
    return s.lower()


def token_words(s: str) -> List[str]:
    return re.findall(r"[가-힣A-Za-z]+|\d+(?:,\d{3})*(?:\.\d+)?%?", s.lower())


def structure_sig(text: str) -> Tuple[int, int, int]:
    return (len(CIRCLED_RE.findall(text)), len(HO_RE.findall(text)), len(MOK_RE.findall(text)))


def parse_units(text: str) -> List[Unit]:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", x).strip() for x in text.split("\n")]
    lines = [x for x in lines if x]

    units: List[Unit] = []
    section_stack: Dict[str, str] = {}
    current: Optional[Dict[str, Any]] = None
    preamble: List[str] = []

    def current_section() -> str:
        parts = []
        for key in ("장", "절", "관"):
            if section_stack.get(key):
                parts.append(section_stack[key])
        return " > ".join(parts)

    def flush_current():
        nonlocal current
        if not current:
            return
        body = "\n".join(current["body"]).strip()
        header = current["header"].strip()
        full = header if not body else header + "\n" + body
        units.append(Unit(
            index=len(units), kind="article", number=current["number"], title=current["title"],
            header=header, body=body, text=full, section=current["section"],
            norm_title=normalize_text(current["title"]), norm_body=normalize_text(body),
            structure=structure_sig(body)
        ))
        current = None

    for line in lines:
        sm = SECTION_RE.match(line)
        if sm:
            flush_current()
            level = sm.group(2)
            label = f"제{sm.group(1)}{level}" + (f" {sm.group(3).strip()}" if sm.group(3).strip() else "")
            section_stack[level] = label
            if level == "장":
                section_stack.pop("절", None); section_stack.pop("관", None)
            elif level == "절":
                section_stack.pop("관", None)
            continue

        am = ARTICLE_RE.match(line)
        if am:
            flush_current()
            num = am.group(1) + (f"의{am.group(2)}" if am.group(2) else "")
            title = (am.group(3) or "").strip()
            tail = (am.group(4) or "").strip()
            header = f"제{num}조" + (f"({title})" if title else "")
            current = {"number": num, "title": title, "header": header, "body": [], "section": current_section()}
            if tail:
                current["body"].append(tail)
        else:
            if current:
                current["body"].append(line)
            else:
                preamble.append(line)

    flush_current()

    if preamble:
        pre = "\n".join(preamble).strip()
        units.insert(0, Unit(
            index=0, kind="preamble", number="", title="문서 머리말", header="문서 머리말", body=pre,
            text=pre, section="", norm_title="문서머리말", norm_body=normalize_text(pre), structure=structure_sig(pre)
        ))
        for i, u in enumerate(units):
            u.index = i

    # Fallback for documents without formal article headers: compare by paragraphs.
    if not units:
        paras = [x.strip() for x in text.split("\n") if x.strip()]
        for i, p in enumerate(paras):
            units.append(Unit(i, "paragraph", str(i+1), f"문단 {i+1}", f"문단 {i+1}", p, p, "", f"문단{i+1}", normalize_text(p), structure_sig(p)))

    return units


def seq_ratio(a: str, b: str) -> float:
    """Fast normalized similarity used by the alignment engine.

    RapidFuzz's Indel similarity is implemented in native code and is normally very close
    to difflib's ratio for document text, while being much faster.  Keep the old
    SequenceMatcher path as a compatibility fallback when RapidFuzz is unavailable.
    """
    if not a and not b:
        return 1.0
    if not a or not b:
        return 0.0
    if _RapidFuzzIndel is not None:
        return float(_RapidFuzzIndel.normalized_similarity(a, b))
    return SequenceMatcher(None, a, b, autojunk=False).ratio()


def jaccard(a: str, b: str) -> float:
    sa, sb = set(token_words(a)), set(token_words(b))
    if not sa and not sb:
        return 1.0
    if not sa or not sb:
        return 0.0
    return len(sa & sb) / len(sa | sb)


def struct_similarity(a: Tuple[int,int,int], b: Tuple[int,int,int]) -> float:
    dif = sum(abs(x-y) for x, y in zip(a,b))
    base = max(1, sum(a)+sum(b))
    return max(0.0, 1.0 - dif/base)


def similarity(a: Unit, b: Unit) -> float:
    if a.kind != b.kind and "preamble" in (a.kind, b.kind):
        return 0.0
    title_r = seq_ratio(a.norm_title, b.norm_title) if a.title and b.title else 0.0
    body_r = seq_ratio(a.norm_body, b.norm_body)
    jac = jaccard(a.body, b.body)
    st = struct_similarity(a.structure, b.structure)
    # Body gets most weight; title helps moved / renumbered clauses significantly.
    if a.title and b.title:
        score = 0.28*title_r + 0.47*body_r + 0.18*jac + 0.07*st
    else:
        score = 0.58*body_r + 0.30*jac + 0.12*st
    if a.norm_body and a.norm_body == b.norm_body:
        score = max(score, 0.96)
    if a.norm_title and a.norm_title == b.norm_title:
        score += 0.04
    return min(score, 1.0)


def match_units(base: List[Unit], other: List[Unit], threshold: float = 0.48) -> Tuple[Dict[int, Tuple[int,float]], set]:
    candidates: List[Tuple[float,int,int,float]] = []
    for i, a in enumerate(base):
        for j, b in enumerate(other):
            s = similarity(a,b)
            if s >= threshold:
                # tiny proximity tie-break only; never enough to override content.
                proximity = 1.0 - min(abs(i-j)/max(len(base),len(other),1), 1.0)
                rank = s + proximity*0.015
                candidates.append((rank, i, j, s))
    candidates.sort(reverse=True)
    used_a, used_b = set(), set()
    mapping: Dict[int, Tuple[int,float]] = {}
    for _, i, j, s in candidates:
        if i in used_a or j in used_b:
            continue
        mapping[i] = (j, s)
        used_a.add(i); used_b.add(j)
    return mapping, used_b


def build_groups(docs_units: List[List[Unit]]) -> List[Dict[str, Any]]:
    base = docs_units[0]
    maps = []
    used_sets = []
    for units in docs_units[1:]:
        mp, used = match_units(base, units)
        maps.append(mp); used_sets.append(used)

    groups: List[Dict[str, Any]] = []
    for i, a in enumerate(base):
        members = [a]
        scores = [1.0]
        for d_idx, units in enumerate(docs_units[1:]):
            if i in maps[d_idx]:
                j, s = maps[d_idx][i]
                members.append(units[j]); scores.append(s)
            else:
                members.append(None); scores.append(None)
        groups.append({"members": members, "scores": scores})

    # Unmatched additions. For 3 docs, pair B/C additions if they are likely the same new clause.
    if len(docs_units) == 3:
        b_un = [j for j in range(len(docs_units[1])) if j not in used_sets[0]]
        c_un = [j for j in range(len(docs_units[2])) if j not in used_sets[1]]
        b_units = [docs_units[1][j] for j in b_un]
        c_units = [docs_units[2][j] for j in c_un]
        local_map, c_used_local = match_units(b_units, c_units, threshold=0.50)
        matched_b_local = set(local_map.keys())
        for bi_local, (ci_local, s) in local_map.items():
            groups.append({"members": [None, b_units[bi_local], c_units[ci_local]], "scores": [None, 1.0, s]})
        for bi_local, u in enumerate(b_units):
            if bi_local not in matched_b_local:
                groups.append({"members": [None, u, None], "scores": [None, 1.0, None]})
        for ci_local, u in enumerate(c_units):
            if ci_local not in c_used_local:
                groups.append({"members": [None, None, u], "scores": [None, None, 1.0]})
    else:
        for j, u in enumerate(docs_units[1]):
            if j not in used_sets[0]:
                groups.append({"members": [None, u], "scores": [None, 1.0]})

    return groups


def display_tokens(text: str) -> List[str]:
    # Preserve the ORIGINAL token text/length for marker offsets, but include Hangul Jamo
    # so canonically equivalent NFC/NFD Korean text is not silently skipped by the lexer.
    return re.findall(
        r"\s+|[가-힣\u1100-\u11ff\u3130-\u318fA-Za-z]+|\d+(?:,\d{3})*(?:\.\d+)?%?|[^\w\s]",
        text or "", re.UNICODE
    )


def token_key(tok: str) -> str:
    if tok.isspace():
        return " "
    # Compare canonical Unicode, but keep the raw token itself for display/offsets.
    return unicodedata.normalize("NFC", tok).lower()


def pair_changed_indices(a_text: str, b_text: str) -> Tuple[set,set]:
    a = display_tokens(a_text); b = display_tokens(b_text)
    ak = [token_key(x) for x in a]; bk = [token_key(x) for x in b]
    sm = SequenceMatcher(None, ak, bk, autojunk=False)
    ac, bc = set(), set()
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag != "equal":
            ac.update(range(i1,i2)); bc.update(range(j1,j2))
    return ac, bc


def _merge_styled_tokens_v19(tokens: List[str], styles: List[str]) -> List[Tuple[str,str]]:
    if not tokens: return []
    out=[]; cur_style=styles[0] if styles else 'normal'; buf=[]
    for tok,sty in zip(tokens,styles):
        sty=sty or 'normal'
        if sty!=cur_style and buf:
            out.append((''.join(buf),cur_style)); buf=[tok]; cur_style=sty
        else:
            buf.append(tok)
    if buf: out.append((''.join(buf),cur_style))
    return out


def _directional_pair_marks_v19(base_text: str, other_text: str):
    """Return baseline deletion/replacement marks and comparison insertion/replacement marks."""
    bt=display_tokens(base_text); ot=display_tokens(other_text)
    bk=[token_key(x) for x in bt]; ok=[token_key(x) for x in ot]
    bm=['normal']*len(bt); om=['normal']*len(ot)
    sm=SequenceMatcher(None,bk,ok,autojunk=False)
    for tag,i1,i2,j1,j2 in sm.get_opcodes():
        if tag in ('replace','delete'):
            for k in range(i1,i2): bm[k]='delete'
        if tag in ('replace','insert'):
            for k in range(j1,j2): om[k]='insert'
    return bt,bm,ot,om


def directional_segments_v19(members: List[Optional[Unit]], base_index: int, attr: str='body') -> List[List[Tuple[str,str]]]:
    """Style changes relative to the explicitly selected baseline.

    Baseline removed/replaced text = red strike-through (delete).
    Comparison added/replacement text = blue underline (insert).
    """
    n=len(members); result=[[] for _ in range(n)]; base=members[base_index]
    if base is None:
        for i,m in enumerate(members):
            if m:
                txt=getattr(m,attr) or ''
                toks=display_tokens(txt)
                result[i]=_merge_styled_tokens_v19(toks,['insert']*len(toks))
        return result

    base_text=getattr(base,attr) or ''
    base_tokens=display_tokens(base_text); base_marks=['normal']*len(base_tokens)
    for i,m in enumerate(members):
        if i==base_index: continue
        if m is None:
            # Entire baseline clause/content was deleted from this comparison document.
            for k in range(len(base_marks)): base_marks[k]='delete'
            continue
        bt,bm,ot,om=_directional_pair_marks_v19(base_text,getattr(m,attr) or '')
        # display_tokens(base_text) is deterministic; union deletion/replacement marks across B/C.
        for k,sty in enumerate(bm):
            if sty=='delete' and k < len(base_marks): base_marks[k]='delete'
        result[i]=_merge_styled_tokens_v19(ot,om)
    result[base_index]=_merge_styled_tokens_v19(base_tokens,base_marks)
    return result


def multi_changed(members: List[Optional[Unit]]) -> List[set]:
    change_sets = [set() for _ in members]
    for i in range(len(members)):
        if not members[i]:
            continue
        for j in range(i+1,len(members)):
            if not members[j]:
                continue
            ai, bj = pair_changed_indices(members[i].text, members[j].text)
            change_sets[i].update(ai); change_sets[j].update(bj)
    # Entire content of unilateral added/deleted articles should be highlighted in the versions where it exists.
    present = [i for i,m in enumerate(members) if m]
    if len(present) == 1:
        idx = present[0]
        change_sets[idx].update(range(len(display_tokens(members[idx].text))))
    return change_sets


def segments_from_changes(text: str, changed: set) -> List[Tuple[str,bool]]:
    toks = display_tokens(text)
    out: List[Tuple[str,bool]] = []
    for i,t in enumerate(toks):
        ch = i in changed and not t.isspace()
        if out and out[-1][1] == ch:
            out[-1] = (out[-1][0] + t, ch)
        else:
            out.append((t,ch))
    return out


def render_segments_html(segments: List[Tuple[str,bool]]) -> str:
    parts=[]
    for text, ch in segments:
        esc = html.escape(text)
        if ch:
            parts.append(f'<span class="changed">{esc}</span>')
        else:
            parts.append(esc)
    return "".join(parts).replace("\n", "<br>")


def number_values(text: str) -> List[str]:
    return [re.sub(r"\s+", "", x) for x in NUMBER_RE.findall(text)]


def term_presence(text: str) -> set:
    return {t for t in LEGAL_TERMS if t in text}


def summarize_group(members: List[Optional[Unit]], scores: List[Optional[float]]) -> Dict[str, Any]:
    present = [i for i,m in enumerate(members) if m]
    msgs: List[str] = []
    tags: List[str] = []
    risk = "LOW"

    if 0 not in present:
        msgs.append("기준문서에는 없고 비교문서에 새로 등장한 조항입니다.")
        tags.append("신설")
        risk = "MEDIUM"
    elif len(present) == 1 and present[0] == 0:
        msgs.append("기준문서에는 있으나 모든 비교문서에서 사라진 조항입니다.")
        tags.append("삭제")
        risk = "HIGH"
    else:
        for i in range(1, len(members)):
            a, b = members[0], members[i]
            label = chr(ord('A')+i)
            if a and not b:
                msgs.append(f"문서 {label}: 해당 조항이 확인되지 않습니다(삭제 또는 대규모 재편 가능).")
                if "삭제" not in tags: tags.append("삭제")
                risk = "HIGH"
            elif not a and b:
                msgs.append(f"문서 {label}: 신규 조항입니다.")
            elif a and b:
                if a.number != b.number and a.kind == "article" and b.kind == "article":
                    msgs.append(f"문서 {label}: 제{a.number}조 → 제{b.number}조로 이동/재번호화되었습니다.")
                    if "이동" not in tags: tags.append("이동")
                if a.norm_title != b.norm_title and a.title and b.title:
                    msgs.append(f"문서 {label}: 조문 제목이 ‘{a.title}’ → ‘{b.title}’로 변경되었습니다.")
                    if "제목변경" not in tags: tags.append("제목변경")
                if a.section != b.section and (a.section or b.section):
                    msgs.append(f"문서 {label}: 상위 편제가 ‘{a.section or '없음'}’ → ‘{b.section or '없음'}’로 변경되었습니다.")
                    if "편제이동" not in tags: tags.append("편제이동")
                if a.norm_body != b.norm_body:
                    if "내용변경" not in tags: tags.append("내용변경")
                    if risk == "LOW": risk = "MEDIUM"

    # Numeric differences across versions.
    nums = [number_values(m.body) if m else [] for m in members]
    if len({tuple(x) for x in nums if x}) > 1:
        vals = []
        for i,n in enumerate(nums):
            if n:
                vals.append(f"{chr(65+i)}: {', '.join(n[:8])}")
        if vals:
            msgs.append("수치/기간/금액 차이: " + " / ".join(vals))
            tags.append("수치변경")
            risk = "HIGH"

    # Material legal term presence changes.
    terms = [term_presence(m.text) if m else set() for m in members]
    present_term_sets = [t for t in terms if t]
    if present_term_sets and any(t != present_term_sets[0] for t in present_term_sets[1:]):
        changed_terms = sorted(set.union(*present_term_sets) - set.intersection(*present_term_sets)) if len(present_term_sets)>1 else []
        if changed_terms:
            msgs.append("법적 효과에 영향을 줄 수 있는 표현 차이: " + ", ".join(changed_terms[:10]))
            tags.append("법률표현")
            risk = "HIGH"

    # B/C divergence in three-way comparison.
    if len(members) == 3 and members[1] and members[2]:
        bc = similarity(members[1], members[2])
        if members[0] and members[1].norm_body != members[2].norm_body and bc < 0.93:
            msgs.append("문서 B와 C의 수정 내용이 서로 다릅니다. 충돌 여부를 확인하세요.")
            tags.append("B/C차이")
            risk = "HIGH"

    if not msgs:
        msgs.append("실질적인 내용 차이가 확인되지 않았습니다.")
        tags.append("동일")

    valid_scores = [s for i,s in enumerate(scores) if i>0 and s is not None]
    confidence = round(min(valid_scores)*100) if valid_scores else (100 if len(present)==1 else None)

    return {"messages": msgs, "tags": tags, "risk": risk, "confidence": confidence}


def compare_documents(names: List[str], texts: List[str]) -> Dict[str, Any]:
    units = [parse_units(t) for t in texts]
    groups = build_groups(units)
    rows=[]
    counts={"total":0,"changed":0,"added":0,"deleted":0,"moved":0,"high":0}
    for ridx,g in enumerate(groups):
        members=g["members"]; scores=g["scores"]
        changes=multi_changed(members)
        segs=[]; htmls=[]; raw=[]
        for i,m in enumerate(members):
            if m:
                sg=segments_from_changes(m.text, changes[i]); segs.append(sg); htmls.append(render_segments_html(sg)); raw.append(m.text)
            else:
                segs.append([]); htmls.append('<span class="missing">[해당 조항 없음]</span>'); raw.append("")
        summary=summarize_group(members,scores)
        changed = any(changes[i] for i,m in enumerate(members) if m) or len([m for m in members if m]) != len(members)
        counts["total"] += 1
        if changed: counts["changed"] += 1
        if "신설" in summary["tags"]: counts["added"] += 1
        if "삭제" in summary["tags"]: counts["deleted"] += 1
        if "이동" in summary["tags"]: counts["moved"] += 1
        if summary["risk"] == "HIGH": counts["high"] += 1
        rows.append({
            "id":ridx, "raw":raw, "html":htmls, "segments":segs,
            "members":[asdict(m) if m else None for m in members],
            "summary":summary, "changed":changed
        })
    return {"names":names,"rows":rows,"counts":counts,"unit_counts":[len(x) for x in units]}


def make_xlsx(result: Dict[str,Any]) -> bytes:
    out=io.BytesIO()
    wb=xlsxwriter.Workbook(out, {"in_memory":True})
    ws=wb.add_worksheet("조문비교")
    summary_ws=wb.add_worksheet("요약")

    header_fmt=wb.add_format({"bold":True,"bg_color":"#1F4E78","font_color":"#FFFFFF","align":"center","valign":"vcenter","border":1})
    cell_fmt=wb.add_format({"text_wrap":True,"valign":"top","border":1,"font_size":10})
    missing_fmt=wb.add_format({"text_wrap":True,"valign":"top","border":1,"font_color":"#777777","italic":True,"bg_color":"#F2F2F2"})
    changed_fmt=wb.add_format({"font_color":"#C00000","underline":True,"bold":True})
    changed_cell_fmt=wb.add_format({"text_wrap":True,"valign":"top","border":1,"font_color":"#C00000","underline":True,"bold":True})
    high_fmt=wb.add_format({"text_wrap":True,"valign":"top","border":1,"bg_color":"#FCE8E6"})
    med_fmt=wb.add_format({"text_wrap":True,"valign":"top","border":1,"bg_color":"#FFF4CE"})
    title_fmt=wb.add_format({"bold":True,"font_size":16,"font_color":"#1F4E78"})
    stat_label=wb.add_format({"bold":True,"bg_color":"#D9EAF7","border":1})
    stat_val=wb.add_format({"border":1,"align":"center"})

    headers=[f"{chr(65+i)} · {n}" for i,n in enumerate(result["names"])]+["변경사항"]
    ws.write_row(0,0,headers,header_fmt)
    ws.freeze_panes(1,0)
    ws.autofilter(0,0,max(1,len(result["rows"])),len(headers)-1)
    doc_width=42 if len(result["names"])==2 else 32
    for c in range(len(result["names"])):
        ws.set_column(c,c,doc_width)
    ws.set_column(len(result["names"]),len(result["names"]),44)

    def write_segments(row,col,segments):
        if not segments:
            ws.write(row,col,"[해당 조항 없음]",missing_fmt); return
        if all(ch for text,ch in segments if text.strip()):
            ws.write(row,col,"".join(t for t,_ in segments),changed_cell_fmt); return
        args=[]
        for text,ch in segments:
            if not text: continue
            if ch and text.strip():
                args.extend([changed_fmt,text])
            else:
                args.append(text)
        # rich_string needs at least one format and two string fragments; otherwise normal write.
        if any(ch and t.strip() for t,ch in segments) and len(args)>=3:
            try:
                ws.write_rich_string(row,col,*args,cell_fmt)
                return
            except Exception:
                pass
        ws.write(row,col,"".join(t for t,_ in segments),cell_fmt)

    for r_idx,row in enumerate(result["rows"],start=1):
        for c,segments in enumerate(row["segments"]):
            write_segments(r_idx,c,segments)
        sm=row["summary"]
        lines=[f"[{sm['risk']}] " + " / ".join(sm["tags"])] + sm["messages"]
        fmt=high_fmt if sm["risk"]=="HIGH" else med_fmt if sm["risk"]=="MEDIUM" else cell_fmt
        ws.write(r_idx,len(result["names"]),"\n".join(lines),fmt)
        ws.set_row(r_idx, min(150, max(42, 15*(max(2, len(lines)+2)))))

    summary_ws.write("A1","법률문서 비교 요약",title_fmt)
    summary_ws.write("A3","비교 문서",stat_label)
    summary_ws.write("B3"," / ".join(result["names"]),stat_val)
    summary_ws.set_column("A:A",20); summary_ws.set_column("B:B",55)
    labels=[("전체 대응행","total"),("변경행","changed"),("신설","added"),("삭제","deleted"),("이동/재번호화","moved"),("HIGH 위험","high")]
    for i,(label,key) in enumerate(labels,start=5):
        summary_ws.write(i-1,0,label,stat_label); summary_ws.write(i-1,1,result["counts"][key],stat_val)
    summary_ws.write("A13","표시 규칙",stat_label)
    summary_ws.write("B13","붉은 글씨 + 밑줄 = 다른 문서와 비교해 변경된 텍스트",cell_fmt)
    summary_ws.write("A15","주의",stat_label)
    summary_ws.write("B15","자동 매칭 결과는 검토 보조용입니다. 특히 대규모 조문 분할·병합 또는 전면 개정은 사람이 최종 확인해야 합니다.",cell_fmt)
    wb.close()
    out.seek(0)
    return out.read()





# ---------------- V1.4: faster matching + explicit baseline + richer change summary ----------------
import logging
from pathlib import Path
from collections import defaultdict
from statistics import median

APP_VERSION = "1.5"
LOG_PATH = Path(__file__).with_name("legal_doc_compare.log")
logging.basicConfig(
    filename=str(LOG_PATH), level=logging.INFO,
    format="%(asctime)s %(levelname)s %(message)s", encoding="utf-8"
)

class ComparisonCancelled(Exception):
    pass


def _check_cancel(cancel_event):
    if cancel_event is not None and cancel_event.is_set():
        raise ComparisonCancelled("비교가 취소되었습니다.")


def _candidate_indices(a: Unit, i: int, base_len: int, other: List[Unit], title_idx, number_idx, token_idx) -> set:
    n=len(other)
    if n <= 90:
        return set(range(n))
    cand=set()
    if a.norm_title:
        cand.update(title_idx.get(a.norm_title, ()))
    if a.number:
        cand.update(number_idx.get(a.number, ()))
    # Nearby position candidates catch modest reordering cheaply.
    projected = round(i * max(0, n-1) / max(1, base_len-1))
    for j in range(max(0, projected-14), min(n, projected+15)):
        cand.add(j)
    # Rare/long tokens catch clauses moved far away even when article numbers changed.
    toks=sorted({t for t in token_words((a.title or '')+' '+a.body) if len(t) >= 2}, key=len, reverse=True)[:12]
    for tok in toks:
        hits=token_idx.get(tok, ())
        if len(hits) <= max(35, n//8):
            cand.update(hits)
    if len(cand) < 8:
        for j in range(max(0, projected-28), min(n, projected+29)):
            cand.add(j)
    return cand


def match_units_fast(base: List[Unit], other: List[Unit], threshold: float = 0.48,
                     cancel_event=None, progress_cb=None) -> Tuple[Dict[int, Tuple[int,float]], set]:
    title_idx=defaultdict(set); number_idx=defaultdict(set); token_idx=defaultdict(set)
    for j,b in enumerate(other):
        if b.norm_title: title_idx[b.norm_title].add(j)
        if b.number: number_idx[b.number].add(j)
        for tok in set(token_words((b.title or '')+' '+b.body)):
            if len(tok)>=2: token_idx[tok].add(j)
    candidates=[]
    total=max(1,len(base))
    for i,a in enumerate(base):
        _check_cancel(cancel_event)
        for j in _candidate_indices(a,i,len(base),other,title_idx,number_idx,token_idx):
            s=similarity(a,other[j])
            if s>=threshold:
                proximity=1.0-min(abs(i-j)/max(len(base),len(other),1),1.0)
                candidates.append((s+proximity*0.015,i,j,s))
        if progress_cb and i % max(1,total//20)==0:
            progress_cb(i/total)
    candidates.sort(reverse=True)
    used_a=set(); used_b=set(); mapping={}
    for _,i,j,s in candidates:
        if i in used_a or j in used_b: continue
        mapping[i]=(j,s); used_a.add(i); used_b.add(j)
    return mapping, used_b


def build_groups(docs_units: List[List[Unit]], base_index: int = 0, cancel_event=None, progress_cb=None) -> List[Dict[str, Any]]:
    base=docs_units[base_index]
    maps={}; used_sets={}
    others=[i for i in range(len(docs_units)) if i != base_index]
    for pos,doc_idx in enumerate(others):
        def subprogress(x, pos=pos):
            if progress_cb:
                progress_cb((pos+x)/max(1,len(others)))
        mp,used=match_units_fast(base, docs_units[doc_idx], cancel_event=cancel_event, progress_cb=subprogress)
        maps[doc_idx]=mp; used_sets[doc_idx]=used

    groups=[]
    for bi,a in enumerate(base):
        members=[None]*len(docs_units); scores=[None]*len(docs_units)
        members[base_index]=a; scores[base_index]=1.0
        for doc_idx in others:
            if bi in maps[doc_idx]:
                j,s=maps[doc_idx][bi]
                members[doc_idx]=docs_units[doc_idx][j]; scores[doc_idx]=s
        groups.append({"members":members,"scores":scores})

    # Pair unmatched additions between the two non-baseline documents in a 3-way comparison.
    if len(docs_units)==3:
        d1,d2=others
        idx1=[j for j in range(len(docs_units[d1])) if j not in used_sets[d1]]
        idx2=[j for j in range(len(docs_units[d2])) if j not in used_sets[d2]]
        u1=[docs_units[d1][j] for j in idx1]; u2=[docs_units[d2][j] for j in idx2]
        local,used2=match_units_fast(u1,u2,threshold=0.50,cancel_event=cancel_event) if u1 and u2 else ({},set())
        used1=set(local.keys())
        for i1,(i2,s) in local.items():
            members=[None]*3; scores=[None]*3
            members[d1]=u1[i1]; members[d2]=u2[i2]; scores[d1]=1.0; scores[d2]=s
            groups.append({"members":members,"scores":scores})
        for i,u in enumerate(u1):
            if i not in used1:
                members=[None]*3; scores=[None]*3; members[d1]=u; scores[d1]=1.0
                groups.append({"members":members,"scores":scores})
        for i,u in enumerate(u2):
            if i not in used2:
                members=[None]*3; scores=[None]*3; members[d2]=u; scores[d2]=1.0
                groups.append({"members":members,"scores":scores})
    elif len(docs_units)==2:
        d=others[0]
        for j,u in enumerate(docs_units[d]):
            if j not in used_sets[d]:
                members=[None]*2; scores=[None]*2; members[d]=u; scores[d]=1.0
                groups.append({"members":members,"scores":scores})

    # V1.9: the selected baseline is a truly immutable vertical axis.
    # Baseline rows are emitted first, in their physical source order, with no comparison-only
    # additions interleaved.  Unmatched additions are collected afterwards in a dedicated
    # baseline-external section.  This guarantees Article 4 -> Article 5 can never be broken by
    # an Article 27 that exists only in B/C.
    base_groups=groups[:len(base)]
    addition_groups=groups[len(base):]

    # Preserve natural order for comparison-only additions without allowing them to disturb the
    # baseline axis.  Attach a marker used by the GUI/summary.
    def add_sort_key(g):
        pos=[]; nums=[]
        for di,m in enumerate(g['members']):
            if di==base_index or m is None: continue
            pos.append(m.index)
            if m.kind=='article':
                v=_article_num_value_v16(m.number)
                if v is not None: nums.append(v)
        return (median(pos) if pos else 10**9, median(nums) if nums else 10**9)

    for g in base_groups:
        g['axis_extra']=False
    for g in addition_groups:
        g['axis_extra']=True
    addition_groups=sorted(addition_groups,key=add_sort_key)
    return base_groups + addition_groups


def _clean_diff_piece(tokens):
    s=''.join(tokens)
    s=re.sub(r'\s+',' ',s).strip()
    return s


def compact_text_diff(a_text: str, b_text: str, limit: int = 6) -> List[str]:
    a=display_tokens(a_text); b=display_tokens(b_text)
    ak=[token_key(x) for x in a]; bk=[token_key(x) for x in b]
    sm=SequenceMatcher(None,ak,bk,autojunk=False)
    out=[]
    for tag,i1,i2,j1,j2 in sm.get_opcodes():
        if tag=='equal': continue
        old=_clean_diff_piece(a[i1:i2]); new=_clean_diff_piece(b[j1:j2])
        if not old and not new: continue
        if tag=='replace': msg=f'“{old[:120]}” → “{new[:120]}”'
        elif tag=='delete': msg=f'삭제: “{old[:140]}”'
        else: msg=f'추가: “{new[:140]}”'
        out.append(msg)
        if len(out)>=limit: break
    return out


def summarize_group(members: List[Optional[Unit]], scores: List[Optional[float]], base_index: int=0) -> Dict[str, Any]:
    present=[i for i,m in enumerate(members) if m]
    msgs=[]; tags=[]; risk='LOW'
    base=members[base_index]
    if base is None:
        docs=', '.join(chr(65+i) for i in present)
        msgs.append(f'기준문서에는 없고 문서 {docs}에 새로 등장한 조항입니다.')
        tags.append('신설'); risk='MEDIUM'
    elif len(present)==1:
        msgs.append('기준문서에는 있으나 모든 비교문서에서 사라진 조항입니다.')
        tags.append('삭제'); risk='HIGH'
    else:
        for i,m in enumerate(members):
            if i==base_index: continue
            label=chr(65+i)
            if base and not m:
                msgs.append(f'문서 {label}: 해당 조항이 확인되지 않습니다(삭제 또는 대규모 재편 가능).')
                if '삭제' not in tags: tags.append('삭제')
                risk='HIGH'; continue
            if not base or not m: continue
            if base.number != m.number and base.kind=='article' and m.kind=='article':
                msgs.append(f'문서 {label}: 제{base.number}조 → 제{m.number}조 이동/재번호화')
                if '이동' not in tags: tags.append('이동')
            if base.norm_title != m.norm_title and base.title and m.title:
                msgs.append(f'문서 {label}: 제목 “{base.title}” → “{m.title}”')
                if '제목변경' not in tags: tags.append('제목변경')
            if base.section != m.section and (base.section or m.section):
                msgs.append(f'문서 {label}: 편제 “{base.section or "없음"}” → “{m.section or "없음"}”')
                if '편제이동' not in tags: tags.append('편제이동')
            if base.norm_body != m.norm_body:
                if '내용변경' not in tags: tags.append('내용변경')
                if risk=='LOW': risk='MEDIUM'
                for d in compact_text_diff(base.body,m.body,limit=5):
                    msgs.append(f'문서 {label}: {d}')

    nums=[number_values(m.body) if m else [] for m in members]
    nonempty=[tuple(x) for x in nums if x]
    if len(set(nonempty))>1:
        vals=[f'{chr(65+i)}: {", ".join(n[:8])}' for i,n in enumerate(nums) if n]
        msgs.append('수치/기간/금액: '+' / '.join(vals))
        if '수치변경' not in tags: tags.append('수치변경')
        risk='HIGH'

    terms=[term_presence(m.text) if m else set() for m in members]
    nonempty_terms=[t for t in terms if t]
    if len(nonempty_terms)>1 and any(t != nonempty_terms[0] for t in nonempty_terms[1:]):
        union=set.union(*nonempty_terms); inter=set.intersection(*nonempty_terms)
        changed_terms=sorted(union-inter)
        if changed_terms:
            msgs.append('법률상 주의 표현: '+', '.join(changed_terms[:10]))
            if '법률표현' not in tags: tags.append('법률표현')
            risk='HIGH'

    if len(members)==3:
        others=[i for i in range(3) if i!=base_index]
        m1,m2=members[others[0]],members[others[1]]
        if m1 and m2 and m1.norm_body != m2.norm_body and similarity(m1,m2)<0.93:
            msgs.append(f'문서 {chr(65+others[0])}와 {chr(65+others[1])}의 수정 내용이 서로 다릅니다.')
            if '비교본차이' not in tags: tags.append('비교본차이')
            risk='HIGH'

    if not msgs:
        msgs=['실질적인 내용 차이가 확인되지 않았습니다.']; tags=['동일']

    valid=[s for i,s in enumerate(scores) if i!=base_index and s is not None]
    confidence=round(min(valid)*100) if valid else (100 if len(present)==1 else None)
    return {'messages':msgs,'tags':tags,'risk':risk,'confidence':confidence}


def compare_documents(names: List[str], texts: List[str], base_index: int=0, progress_cb=None, cancel_event=None) -> Dict[str, Any]:
    if progress_cb: progress_cb(3,'문서를 조문 구조로 나누는 중...')
    units=[]
    for i,t in enumerate(texts):
        _check_cancel(cancel_event)
        units.append(parse_units(t))
        if progress_cb: progress_cb(5+int(15*(i+1)/len(texts)), f'문서 {chr(65+i)} 구조 분석 완료 · {len(units[-1])}개 단위')
    if progress_cb: progress_cb(22,'대응 조항을 찾는 중...')
    def match_progress(x):
        if progress_cb: progress_cb(22+int(33*x),'이동·재번호화 조항 매칭 중...')
    groups=build_groups(units,base_index=base_index,cancel_event=cancel_event,progress_cb=match_progress)
    rows=[]; counts={'total':0,'changed':0,'added':0,'deleted':0,'moved':0,'high':0}
    total=max(1,len(groups))
    for ridx,g in enumerate(groups):
        _check_cancel(cancel_event)
        members=g['members']; scores=g['scores']
        changes=multi_changed(members)
        segs=[]; htmls=[]; raw=[]
        for i,m in enumerate(members):
            if m:
                sg=segments_from_changes(m.text,changes[i]); segs.append(sg); htmls.append(render_segments_html(sg)); raw.append(m.text)
            else:
                segs.append([]); htmls.append('<span class="missing">[해당 조항 없음]</span>'); raw.append('')
        summary=summarize_group(members,scores,base_index=base_index)
        changed=summary['tags'] != ['동일']
        counts['total']+=1
        if changed: counts['changed']+=1
        if '신설' in summary['tags']: counts['added']+=1
        if '삭제' in summary['tags']: counts['deleted']+=1
        if '이동' in summary['tags']: counts['moved']+=1
        if summary['risk']=='HIGH': counts['high']+=1
        rows.append({'id':ridx,'raw':raw,'html':htmls,'segments':segs,
                     'members':[asdict(m) if m else None for m in members],
                     'summary':summary,'changed':changed})
        if progress_cb and ridx % max(1,total//25)==0:
            progress_cb(58+int(40*(ridx+1)/total), f'변경 내용 분석 중... {ridx+1}/{total}')
    if progress_cb: progress_cb(100,'비교 완료')
    return {'names':names,'rows':rows,'counts':counts,'unit_counts':[len(x) for x in units],'base_index':base_index}


# Excel exporter mirrors the GUI's directional revision marks.
def make_xlsx(result: Dict[str,Any]) -> bytes:
    out=io.BytesIO(); wb=xlsxwriter.Workbook(out, {'in_memory':True})
    ws=wb.add_worksheet('조문비교'); sw=wb.add_worksheet('요약')
    hf=wb.add_format({'bold':True,'bg_color':'#1F4E78','font_color':'#FFFFFF','align':'center','valign':'vcenter','border':1})
    cf=wb.add_format({'text_wrap':True,'valign':'top','border':1,'font_size':10})
    mf=wb.add_format({'text_wrap':True,'valign':'top','border':1,'font_color':'#777777','italic':True,'bg_color':'#F2F2F2'})
    delf=wb.add_format({'font_color':'#C00000','font_strikeout':True})
    insf=wb.add_format({'font_color':'#1565C0','underline':True})
    delcell=wb.add_format({'text_wrap':True,'valign':'top','border':1,'font_color':'#C00000','font_strikeout':True})
    inscell=wb.add_format({'text_wrap':True,'valign':'top','border':1,'font_color':'#1565C0','underline':True})
    high=wb.add_format({'text_wrap':True,'valign':'top','border':1,'bg_color':'#FCE8E6'})
    med=wb.add_format({'text_wrap':True,'valign':'top','border':1,'bg_color':'#FFF4CE'})
    title=wb.add_format({'bold':True,'font_size':16,'font_color':'#1F4E78'})
    sl=wb.add_format({'bold':True,'bg_color':'#D9EAF7','border':1}); sv=wb.add_format({'border':1,'align':'center','text_wrap':True})
    bi=result.get('base_index',0)
    headers=[]
    for i,n in enumerate(result['names']):
        headers.append(f'{chr(65+i)} · {n}' + (' · 기준' if i==bi else ''))
    headers.append('변경사항')
    ws.write_row(0,0,headers,hf); ws.freeze_panes(1,0); ws.autofilter(0,0,max(1,len(result['rows'])),len(headers)-1)
    if len(result['names'])==3:
        for c in range(3): ws.set_column(c,c,30)
        ws.set_column(3,3,46)
    else:
        for c in range(2): ws.set_column(c,c,40)
        ws.set_column(2,2,48)

    def write_segments(r,c,segs):
        if not segs: ws.write(r,c,'[해당 조항 없음]',mf); return
        nonempty=[(txt,sty) for txt,sty in segs if txt]
        if not nonempty: ws.write(r,c,'',cf); return
        styles={sty for txt,sty in nonempty if txt.strip()}
        if styles=={'delete'}: ws.write(r,c,''.join(t for t,_ in nonempty),delcell); return
        if styles=={'insert'}: ws.write(r,c,''.join(t for t,_ in nonempty),inscell); return
        args=[]; has_fmt=False
        for txt,sty in nonempty:
            if sty=='delete' and txt.strip(): args.extend([delf,txt]); has_fmt=True
            elif sty=='insert' and txt.strip(): args.extend([insf,txt]); has_fmt=True
            else: args.append(txt)
        if has_fmt:
            try:
                ws.write_rich_string(r,c,*args,cf); return
            except Exception:
                pass
        ws.write(r,c,''.join(t for t,_ in nonempty),cf)

    for r,row in enumerate(result['rows'],1):
        for c,segs in enumerate(row['segments']): write_segments(r,c,segs)
        s=row['summary']; lines=(s.get('messages') or ['변경 없음'])
        ws.write(r,len(result['names']),'\n'.join(lines),cf)
        ws.set_row(r,min(240,max(45,15*(len(lines)+2))))
    sw.write('A1','법률문서 비교 요약',title); sw.set_column('A:A',20); sw.set_column('B:B',70)
    sw.write('A3','기준 문서',sl); sw.write('B3',result['names'][bi],sv)
    sw.write('A4','비교 문서',sl); sw.write('B4',' / '.join(n for i,n in enumerate(result['names']) if i!=bi),sv)
    labels=[('전체 대응행','total'),('변경행','changed'),('신설','added'),('삭제','deleted'),('이동/재번호화','moved')]
    for i,(label,key) in enumerate(labels,start=6): sw.write(i-1,0,label,sl); sw.write(i-1,1,result['counts'][key],sv)
    sw.write('A14','표시 규칙',sl)
    sw.write('B14','삭제/대체된 기존 문구 = 붉은색 취소선\n신설/대체된 새 문구 = 파란색 밑줄',cf)
    sw.write('A15','정렬 규칙',sl)
    sw.write('B15','선택한 기준문서의 원래 조항 순서를 절대축으로 유지하며, 기준문서에 없는 신설 조항은 하단에 별도 표시',cf)
    wb.close(); out.seek(0); return out.read()


# ---------------- V1.5: article-first parser + per-part change analysis ----------------
KOR_ARTICLE_RE_V15 = re.compile(
    r"^\s*제\s*(\d+)\s*조(?:\s*의\s*(\d+))?\s*(?:\(([^\n\)]{1,160})\)|\[([^\n\]]{1,160})\])?\s*(.*)$",
    re.IGNORECASE,
)
ENG_ARTICLE_RE_V15 = re.compile(
    r"^\s*Article\s+((?:\d+(?:[-.]\d+)*[A-Za-z]?)|(?:[IVXLCDM]+))\s*(?:\(([^\n\)]{1,180})\))?\s*(?:[:\-–—]\s*)?(.*)$",
    re.IGNORECASE,
)
ENG_CHAPTER_RE_V15 = re.compile(
    r"^\s*Chapter\s+((?:\d+)|(?:[IVXLCDM]+))\b\s*(.*)$", re.IGNORECASE
)
LEADING_SUBCLAUSE_RE_V15 = re.compile(
    r"^\s*((?:[①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳]|\(\d+\)|\d+[.)]|[가-하][.)]|[A-Za-z][.)]))\s*"
)


def _article_match_v15(line: str):
    m = KOR_ARTICLE_RE_V15.match(line)
    if m:
        number = m.group(1) + (f"의{m.group(2)}" if m.group(2) else "")
        title = (m.group(3) or m.group(4) or "").strip()
        tail = (m.group(5) or "").strip()
        header = f"제{number}조" + (f"({title})" if title else "")
        return {"number": number, "title": title, "tail": tail, "header": header, "lang": "ko"}
    m = ENG_ARTICLE_RE_V15.match(line)
    if m:
        number = (m.group(1) or "").strip()
        title = (m.group(2) or "").strip()
        tail = (m.group(3) or "").strip()
        # A bare short tail without terminal punctuation is commonly an unparenthesized article title.
        if not title and tail and len(tail) <= 100 and not re.search(r"[.!?;:]", tail):
            title, tail = tail, ""
        header = f"Article {number}" + (f" ({title})" if title else "")
        return {"number": number, "title": title, "tail": tail, "header": header, "lang": "en"}
    return None


def parse_units(text: str) -> List[Unit]:
    """Parse a legal document into article-level units.

    V1.5 recognizes both Korean 제N조 and English Article N headings.  The comparison
    therefore remains article-by-article even when a DOCX stores an entire article body
    as one paragraph.  Text before the first article is kept as one preamble unit.
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Paragraph.text preserves manual line breaks. Normalize only horizontal whitespace.
    raw_lines = text.split("\n")
    lines = [re.sub(r"[ \t]+", " ", x).strip() for x in raw_lines]
    lines = [x for x in lines if x]

    units: List[Unit] = []
    section_stack: Dict[str, str] = {}
    current: Optional[Dict[str, Any]] = None
    preamble: List[str] = []

    def current_section() -> str:
        parts=[]
        for key in ("장","절","관","chapter"):
            if section_stack.get(key): parts.append(section_stack[key])
        return " > ".join(parts)

    def flush_current():
        nonlocal current
        if not current: return
        body="\n".join(current["body"]).strip()
        header=current["header"].strip()
        full=header if not body else header+"\n"+body
        units.append(Unit(
            index=len(units), kind="article", number=current["number"], title=current["title"],
            header=header, body=body, text=full, section=current["section"],
            norm_title=normalize_text(current["title"]), norm_body=normalize_text(body),
            structure=structure_sig(body)
        ))
        current=None

    for line in lines:
        sm=SECTION_RE.match(line)
        if sm:
            flush_current()
            level=sm.group(2)
            label=f"제{sm.group(1)}{level}"+(f" {sm.group(3).strip()}" if sm.group(3).strip() else "")
            section_stack[level]=label
            if level=="장": section_stack.pop("절",None); section_stack.pop("관",None)
            elif level=="절": section_stack.pop("관",None)
            continue
        em=ENG_CHAPTER_RE_V15.match(line)
        if em:
            flush_current()
            section_stack["chapter"] = f"Chapter {em.group(1)}" + (f" {em.group(2).strip()}" if em.group(2).strip() else "")
            continue

        am=_article_match_v15(line)
        if am:
            flush_current()
            current={"number":am["number"],"title":am["title"],"header":am["header"],"body":[],"section":current_section()}
            if am["tail"]: current["body"].append(am["tail"])
        else:
            if current: current["body"].append(line)
            else: preamble.append(line)

    flush_current()

    if preamble:
        pre="\n".join(preamble).strip()
        units.insert(0,Unit(
            index=0,kind="preamble",number="",title="문서 머리말",header="문서 머리말",body=pre,text=pre,
            section="",norm_title="문서머리말",norm_body=normalize_text(pre),structure=structure_sig(pre)
        ))
    for i,u in enumerate(units): u.index=i

    if not units:
        paras=[x.strip() for x in text.split("\n") if x.strip()]
        for i,p in enumerate(paras):
            units.append(Unit(i,"paragraph",str(i+1),f"문단 {i+1}",f"문단 {i+1}",p,p,"",f"문단{i+1}",normalize_text(p),structure_sig(p)))
    return units


def _meaningful_piece_v15(s: str) -> bool:
    return bool(re.search(r"[0-9A-Za-z가-힣]", s or ""))


def _short_v15(s: str, n: int=180) -> str:
    s=re.sub(r"\s+"," ",s or "").strip()
    if len(s)<=n: return s
    return s[:n-1].rstrip()+"…"


def _split_sentences_v15(line: str) -> List[str]:
    """Split legal prose into sentence-sized units regardless of paragraph length.

    Word files often store two visually separate sentences in one paragraph in one
    version and as separate paragraphs in another.  The old implementation skipped
    sentence splitting for short paragraphs, which made a leading phrase insertion
    look like 'sentence deleted + sentence added'.
    """
    line=re.sub(r"\s+"," ",line).strip()
    if not line: return []
    # Protect a few common abbreviations so they do not create false sentence breaks.
    protected=line
    repl={
        'e.g.':'e<prd>g<prd>', 'i.e.':'i<prd>e<prd>', 'etc.':'etc<prd>',
        'No.':'No<prd>', 'Art.':'Art<prd>', 'Sec.':'Sec<prd>'
    }
    for k,v in repl.items():
        protected=protected.replace(k,v)
    parts=re.split(r"(?<=[.!?])\s+(?=(?:[\"'“‘(\[]?[A-Z0-9가-힣]))", protected)
    parts=[x.replace('<prd>','.').strip() for x in parts if x.strip()]
    if len(parts)==1 and len(line)>420:
        parts=[x.strip() for x in re.split(r"(?<=;)\s+",line) if x.strip()]
    return parts


def split_legal_subparts_v15(text: str) -> List[Tuple[str,str]]:
    parts=[]; seq=1
    for raw in (text or "").splitlines():
        line=raw.strip()
        if not line: continue
        # Split embedded enumerators when Word stored several numbered clauses in one paragraph.
        line=re.sub(r"\s+(?=(?:[①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳]|\(\d+\)|\d+[.)]\s+|[가-하][.)]\s+))","\n",line)
        for chunk in line.split("\n"):
            for sent in _split_sentences_v15(chunk):
                lm=LEADING_SUBCLAUSE_RE_V15.match(sent)
                if lm:
                    label=lm.group(1)
                else:
                    label=f"문장 {seq}"
                parts.append((label,sent.strip())); seq+=1
    return parts


def _inline_delta_v15(old: str, new: str, max_fragments: int=4) -> str:
    a=display_tokens(old); b=display_tokens(new)
    ak=[token_key(x) for x in a]; bk=[token_key(x) for x in b]
    sm=SequenceMatcher(None,ak,bk,autojunk=False)
    frags=[]
    for tag,i1,i2,j1,j2 in sm.get_opcodes():
        if tag=='equal': continue
        o=_clean_diff_piece(a[i1:i2]); n=_clean_diff_piece(b[j1:j2])
        if not _meaningful_piece_v15(o) and not _meaningful_piece_v15(n): continue
        if tag=='replace': frags.append(f'“{_short_v15(o,90)}” → “{_short_v15(n,90)}”')
        elif tag=='delete': frags.append(f'삭제 “{_short_v15(o,110)}”')
        elif tag=='insert': frags.append(f'추가 “{_short_v15(n,110)}”')
        if len(frags)>=max_fragments: break
    if frags: return ' / '.join(frags)
    return f'“{_short_v15(old,130)}” → “{_short_v15(new,130)}”'


def _subpart_similarity_v22(a: Tuple[str,str], b: Tuple[str,str]) -> float:
    """Similarity used for monotonic sentence/clause alignment."""
    la,ta=a; lb,tb=b
    na=normalize_text(ta); nb=normalize_text(tb)
    if not na or not nb: return 0.0
    char=SequenceMatcher(None,na,nb,autojunk=False).ratio()
    wa=set(token_words(ta)); wb=set(token_words(tb))
    jac=len(wa & wb)/max(1,len(wa | wb))
    score=0.72*char+0.28*jac
    # Matching enumerators ((1), 1., ①...) are strong structural evidence.
    if la==lb and not la.startswith('문장'):
        score=min(1.0,score+0.15)
    return score


def _align_subparts_v22(ap: List[Tuple[str,str]], bp: List[Tuple[str,str]]):
    """Needleman-Wunsch style monotonic alignment with fuzzy match costs.

    A similar sentence is paired even if text was inserted at its beginning/end, so
    revisions are summarized as an in-sentence insertion/deletion instead of as a
    whole sentence deletion followed by a new sentence.
    """
    n,m=len(ap),len(bp); gap=0.62
    dp=[[0.0]*(m+1) for _ in range(n+1)]
    prev=[[None]*(m+1) for _ in range(n+1)]
    for i in range(1,n+1): dp[i][0]=dp[i-1][0]+gap; prev[i][0]='D'
    for j in range(1,m+1): dp[0][j]=dp[0][j-1]+gap; prev[0][j]='I'
    for i in range(1,n+1):
        for j in range(1,m+1):
            sim=_subpart_similarity_v22(ap[i-1],bp[j-1])
            # Pair only when there is meaningful similarity; otherwise gaps are cheaper.
            match_cost=(2.0*(1.0-sim)) if sim>=0.38 else 1.5
            choices=[(dp[i-1][j-1]+match_cost,'M'),(dp[i-1][j]+gap,'D'),(dp[i][j-1]+gap,'I')]
            dp[i][j],prev[i][j]=min(choices,key=lambda x:x[0])
    ops=[]; i=n; j=m
    while i or j:
        op=prev[i][j]
        if op=='M': ops.append(('M',i-1,j-1)); i-=1; j-=1
        elif op=='D': ops.append(('D',i-1,None)); i-=1
        else: ops.append(('I',None,j-1)); j-=1
    return list(reversed(ops))


def detailed_subpart_diff_v15(a_text: str, b_text: str, limit: int=14) -> Dict[str,Any]:
    ap=split_legal_subparts_v15(a_text); bp=split_legal_subparts_v15(b_text)
    lines=[]; total=0
    for op,ai,bi in _align_subparts_v22(ap,bp):
        if op=='M':
            old=ap[ai]; new=bp[bi]
            if normalize_text(old[1])==normalize_text(new[1]):
                continue
            total+=1
            label=old[0] if old[0]==new[0] else f'{old[0]}→{new[0]}'
            if len(lines)<limit:
                lines.append(f'[{label}] 수정: {_inline_delta_v15(old[1],new[1])}')
        elif op=='D':
            label,txt=ap[ai]; total+=1
            if len(lines)<limit: lines.append(f'[{label}] 삭제: “{_short_v15(txt)}”')
        elif op=='I':
            label,txt=bp[bi]; total+=1
            if len(lines)<limit: lines.append(f'[{label}] 추가: “{_short_v15(txt)}”')
    return {"lines":lines,"count":total,"truncated":max(0,total-len(lines))}


def multi_changed(members: List[Optional[Unit]]) -> List[set]:
    """Highlight body differences only; article headings are shown separately in the GUI."""
    change_sets=[set() for _ in members]
    for i in range(len(members)):
        if not members[i]: continue
        for j in range(i+1,len(members)):
            if not members[j]: continue
            ai,bj=pair_changed_indices(members[i].body,members[j].body)
            change_sets[i].update(ai); change_sets[j].update(bj)
    present=[i for i,m in enumerate(members) if m]
    if len(present)==1:
        idx=present[0]; change_sets[idx].update(range(len(display_tokens(members[idx].body))))
    return change_sets


def summarize_group(members: List[Optional[Unit]], scores: List[Optional[float]], base_index: int=0) -> Dict[str,Any]:
    present=[i for i,m in enumerate(members) if m]
    msgs=[]; tags=[]; risk='LOW'; base=members[base_index]
    if base is None:
        docs=', '.join(chr(65+i) for i in present)
        msgs.append(f'기준문서에는 없고 문서 {docs}에 새 조항/부분이 등장했습니다.')
        tags.append('신설'); risk='MEDIUM'
    elif len(present)==1:
        msgs.append('기준문서의 이 조항이 모든 비교문서에서 확인되지 않습니다.')
        tags.append('삭제'); risk='HIGH'
    else:
        for i,m in enumerate(members):
            if i==base_index: continue
            label=chr(65+i)
            if base and not m:
                msgs.append(f'문서 {label}: 대응 조항 없음 — 삭제 또는 대규모 재편 가능')
                if '삭제' not in tags: tags.append('삭제')
                risk='HIGH'; continue
            if not base or not m: continue
            if base.number != m.number and base.kind=='article' and m.kind=='article':
                msgs.append(f'문서 {label}: {base.header} → {m.header} (이동/재번호화)')
                if '이동' not in tags: tags.append('이동')
            if base.norm_title != m.norm_title and base.title and m.title:
                msgs.append(f'문서 {label}: 조 제목 “{base.title}” → “{m.title}”')
                if '제목변경' not in tags: tags.append('제목변경')
            if base.section != m.section and (base.section or m.section):
                msgs.append(f'문서 {label}: 상위 편제 “{base.section or "없음"}” → “{m.section or "없음"}”')
                if '편제이동' not in tags: tags.append('편제이동')
            if base.norm_body != m.norm_body:
                if '내용변경' not in tags: tags.append('내용변경')
                if risk=='LOW': risk='MEDIUM'
                detail=detailed_subpart_diff_v15(base.body,m.body,limit=14)
                msgs.append(f'문서 {label}: 이 조항의 세부 변경 {detail["count"]}건')
                for line in detail['lines']:
                    msgs.append(f'문서 {label} · {line}')
                if detail['truncated']:
                    msgs.append(f'문서 {label}: 화면에는 주요 {len(detail["lines"])}건만 표시 · 나머지 {detail["truncated"]}건')
                # Same-number clauses with semantically related headings are the same lineage
                # even when one version consolidates a detailed enumeration into a broad rule.
                if base.kind=='article' and m.kind=='article' and base.number.lower()==m.number.lower():
                    tr_sem=_title_semantic_ratio_v17(base,m)
                    old_parts=split_legal_subparts_v15(base.body); new_parts=split_legal_subparts_v15(m.body)
                    old_enum=sum(1 for lab,_ in old_parts if lab.startswith('(') or re.match(r'^[0-9①-⑳가-하A-Za-z]',lab or ''))
                    new_enum=sum(1 for lab,_ in new_parts if lab.startswith('(') or re.match(r'^[0-9①-⑳가-하A-Za-z]',lab or ''))
                    len_ratio=(len(m.norm_body)/max(1,len(base.norm_body)))
                    if tr_sem>=0.58 and (detail['count']>=4 or len_ratio<0.62):
                        if '대폭개정' not in tags: tags.append('대폭개정')
                        msgs.append(f'문서 {label}: 동일 계보 조항으로 판단되나 본문이 대폭 개정/축약되었습니다.')
                        if old_enum>=3 and new_enum < old_enum:
                            msgs.append(f'문서 {label}: 기존 세부 열거항목 {old_enum}개가 {new_enum}개 수준으로 축소되어 포괄규정으로 통합되었을 가능성이 있습니다.')
                        risk='HIGH'

    nums=[number_values(m.body) if m else [] for m in members]
    nonempty=[tuple(x) for x in nums if x]
    if len(set(nonempty))>1:
        vals=[f'{chr(65+i)}: {", ".join(n[:12])}' for i,n in enumerate(nums) if n]
        msgs.append('수치/기간/금액 차이: '+' / '.join(vals))
        if '수치변경' not in tags: tags.append('수치변경')
        risk='HIGH'

    terms=[term_presence(m.text) if m else set() for m in members]
    nonempty_terms=[t for t in terms if t]
    if len(nonempty_terms)>1 and any(t != nonempty_terms[0] for t in nonempty_terms[1:]):
        union=set.union(*nonempty_terms); inter=set.intersection(*nonempty_terms)
        changed_terms=sorted(union-inter)
        if changed_terms:
            msgs.append('법률상 주의 표현 차이: '+', '.join(changed_terms[:12]))
            if '법률표현' not in tags: tags.append('법률표현')
            risk='HIGH'

    if len(members)==3:
        others=[i for i in range(3) if i!=base_index]
        m1,m2=members[others[0]],members[others[1]]
        if m1 and m2 and m1.norm_body != m2.norm_body and similarity(m1,m2)<0.93:
            msgs.append(f'문서 {chr(65+others[0])}와 {chr(65+others[1])}가 이 조항을 서로 다르게 수정했습니다.')
            if '비교본차이' not in tags: tags.append('비교본차이')
            risk='HIGH'

    if not msgs:
        msgs=['이 조항은 실질적인 내용 차이가 확인되지 않았습니다.']; tags=['동일']
    valid=[s for i,s in enumerate(scores) if i!=base_index and s is not None]
    confidence=round(min(valid)*100) if valid else (100 if len(present)==1 else None)
    return {'messages':msgs,'tags':tags,'risk':risk,'confidence':confidence}


def compare_documents(names: List[str], texts: List[str], base_index: int=0, progress_cb=None, cancel_event=None) -> Dict[str,Any]:
    if progress_cb: progress_cb(3,'문서를 조(Article) 기준으로 자동 분할하는 중...')
    units=[]
    for i,t in enumerate(texts):
        _check_cancel(cancel_event); parsed=parse_units(t); units.append(parsed)
        article_count=sum(1 for u in parsed if u.kind=='article')
        if progress_cb: progress_cb(5+int(15*(i+1)/len(texts)),f'문서 {chr(65+i)} · 조항 {article_count}개 / 전체 비교단위 {len(parsed)}개')
    if progress_cb: progress_cb(22,'조 번호가 달라도 같은 조항인지 대응 관계를 찾는 중...')
    def match_progress(x):
        if progress_cb: progress_cb(22+int(33*x),'이동·재번호화 조항 매칭 중...')
    groups=build_groups(units,base_index=base_index,cancel_event=cancel_event,progress_cb=match_progress)
    rows=[]; counts={'total':0,'changed':0,'added':0,'deleted':0,'moved':0,'high':0}
    total=max(1,len(groups))
    for ridx,g in enumerate(groups):
        _check_cancel(cancel_event); members=g['members']; scores=g['scores']
        body_segs=directional_segments_v19(members,base_index,'body')
        header_segs=directional_segments_v19(members,base_index,'header')
        full_segs=[]; htmls=[]; raw=[]
        for i,m in enumerate(members):
            if m:
                hs=header_segs[i] or [(m.header,'normal')]
                fs=hs+[('\n','normal')]+body_segs[i] if m.kind!='preamble' else [('문서 머리말\n','normal')]+body_segs[i]
                full_segs.append(fs); htmls.append(render_segments_html_v19(fs)); raw.append(m.text)
            else:
                full_segs.append([]); htmls.append('<span class="missing">[해당 조항 없음]</span>'); raw.append('')
        summary=summarize_group(members,scores,base_index=base_index)
        if g.get('axis_extra') and members[base_index] is None:
            summary['messages'].insert(0,'기준문서 외 신설 조항입니다. 기준문서 순서를 보존하기 위해 하단에 별도 표시합니다.')
            if '신설' not in summary['tags']: summary['tags'].append('신설')
        changed=summary['tags']!=['동일']
        counts['total']+=1
        if changed: counts['changed']+=1
        if '신설' in summary['tags']: counts['added']+=1
        if '삭제' in summary['tags']: counts['deleted']+=1
        if '이동' in summary['tags']: counts['moved']+=1
        if summary['risk']=='HIGH': counts['high']+=1
        rows.append({'id':ridx,'raw':raw,'html':htmls,'segments':full_segs,'body_segments':body_segs,'header_segments':header_segs,
                     'members':[asdict(m) if m else None for m in members],'summary':summary,'changed':changed,'axis_extra':bool(g.get('axis_extra'))})
        if progress_cb and ridx % max(1,total//25)==0:
            progress_cb(58+int(40*(ridx+1)/total),f'조별 세부 변경 분석 중... {ridx+1}/{total}')
    if progress_cb: progress_cb(100,'조별 비교 완료')
    return {'names':names,'rows':rows,'counts':counts,'unit_counts':[len(x) for x in units],'base_index':base_index}



# ---------------- Native Windows GUI (Tkinter + drag & drop) ----------------
def _safe_import_tk():
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox
    return tk,ttk,filedialog,messagebox


class NativeGui:
    def __init__(self, root, dnd_available=False):
        tk,ttk,filedialog,messagebox=_safe_import_tk()
        self.tk,self.ttk=tk,ttk; self.filedialog,self.messagebox=filedialog,messagebox
        self.root=root; self.dnd_available=dnd_available
        self.result=None; self.current_filter='all'; self.cancel_event=threading.Event()
        self.file_vars=[tk.StringVar(),tk.StringVar(),tk.StringVar()]
        self.base_var=tk.IntVar(value=0)
        self.status_var=tk.StringVar(value='TXT 또는 DOCX 2~3개를 선택하거나 아래 영역에 끌어놓으세요.')
        self.progress_var=tk.DoubleVar(value=0)
        self.search_var=tk.StringVar(); self.filter_buttons={}
        self.compare_mode_var=tk.StringVar(value='자동')
        self.stats_vars={k:tk.StringVar(value='0') for k in ('total','changed','added','deleted','moved','high')}
        root.title(f'문서 비교기 V{APP_VERSION}'); root.geometry('1650x950'); root.minsize(1150,720)
        try: root.state('zoomed')
        except Exception: pass
        self._configure_style(); self._build_top(); self._build_results()
        self.search_var.trace_add('write',lambda *_: self.render_rows())

    def _configure_style(self):
        s=self.ttk.Style()
        try: s.theme_use('vista')
        except Exception: pass
        s.configure('Title.TLabel',font=('Malgun Gothic',16,'bold'))
        s.configure('Sub.TLabel',font=('Malgun Gothic',9),foreground='#5f6b76')
        s.configure('Stat.TLabel',font=('Malgun Gothic',10,'bold'),padding=(10,7))
        s.configure('Primary.TButton',font=('Malgun Gothic',10,'bold'),padding=(12,7))
        s.configure('Filter.TButton',font=('Malgun Gothic',9),padding=(9,5))
        s.configure('ActiveFilter.TButton',font=('Malgun Gothic',9,'bold'),padding=(9,5))
        s.configure('Drop.TLabel',font=('Malgun Gothic',10,'bold'),padding=(10,10),anchor='center')

    def _register_drop(self, widget, idx=None):
        if not self.dnd_available: return
        try:
            from tkinterdnd2 import DND_FILES
            widget.drop_target_register(DND_FILES)
            widget.dnd_bind('<<Drop>>', (lambda e:self._drop_files(e,idx)))
        except Exception:
            logging.exception('drag-drop registration failed')

    def _build_top(self):
        tk,ttk=self.tk,self.ttk
        top=ttk.Frame(self.root,padding=(14,12)); top.pack(fill='x')
        ttk.Label(top,text='문서 비교기',style='Title.TLabel').pack(anchor='w')
        ttk.Label(top,text='일반 문서와 법률·규정 문서를 자동 인식해 구조를 맞추고, 삭제=붉은 취소선 / 추가=파란 밑줄로 비교합니다.',style='Sub.TLabel').pack(anchor='w',pady=(2,8))
        drop=ttk.Label(top,text='📄  TXT / DOCX 2~3개를 여기에 끌어놓으세요  ·  클릭해서 여러 파일 선택',style='Drop.TLabel',relief='groove')
        drop.pack(fill='x',pady=(0,9)); drop.bind('<Button-1>',lambda e:self.choose_multiple())
        self._register_drop(drop,None)

        files=ttk.Frame(top); files.pack(fill='x')
        for i in range(3):
            title=f'문서 {chr(65+i)}' + (' (선택)' if i==2 else '')
            box=ttk.LabelFrame(files,text=title,padding=(8,7)); box.grid(row=0,column=i,sticky='nsew',padx=(0 if i==0 else 5,0 if i==2 else 5)); files.grid_columnconfigure(i,weight=1)
            ent=ttk.Entry(box,textvariable=self.file_vars[i],state='readonly'); ent.pack(fill='x',pady=(0,6)); self._register_drop(ent,i); self._register_drop(box,i)
            row=ttk.Frame(box); row.pack(fill='x')
            ttk.Radiobutton(row,text='기준문서로 지정',variable=self.base_var,value=i).pack(side='left')
            ttk.Button(row,text='파일 선택',command=lambda idx=i:self.choose_file(idx)).pack(side='right')

        actions=ttk.Frame(top); actions.pack(fill='x',pady=(10,0))
        self.compare_btn=ttk.Button(actions,text='비교 시작',style='Primary.TButton',command=self.start_compare); self.compare_btn.pack(side='left')
        self.cancel_btn=ttk.Button(actions,text='취소',command=self.cancel_compare,state='disabled'); self.cancel_btn.pack(side='left',padx=(7,0))
        self.export_btn=ttk.Button(actions,text='Excel 내보내기',command=self.export_excel,state='disabled'); self.export_btn.pack(side='left',padx=(7,0))
        self.word_export_btn=ttk.Button(actions,text='Word 변경추적',command=self.export_word_tracked,state='disabled'); self.word_export_btn.pack(side='left',padx=(7,0))
        ttk.Label(actions,text='비교 방식:').pack(side='left',padx=(14,4))
        self.compare_mode_combo=ttk.Combobox(actions,textvariable=self.compare_mode_var,values=('자동','일반 문서','법률·규정'),state='readonly',width=11)
        self.compare_mode_combo.pack(side='left')
        ttk.Label(actions,text='검색:').pack(side='left',padx=(14,4))
        ttk.Entry(actions,textvariable=self.search_var,width=34).pack(side='left')
        self.progress=ttk.Progressbar(actions,variable=self.progress_var,maximum=100,length=180,mode='determinate'); self.progress.pack(side='left',padx=(14,8))
        ttk.Label(actions,textvariable=self.status_var).pack(side='left',fill='x',expand=True)

    def _build_results(self):
        tk,ttk=self.tk,self.ttk
        # V2.4: header and document body MUST share the exact same geometry column.
        # In older builds the header filled the whole outer frame, while the body canvas
        # was narrower by the vertical-scrollbar width. That made every vertical divider
        # drift a few pixels. Put canvas/header in column 0 and the scrollbar alone in
        # column 1 so A/B/C/summary boundaries are structurally identical.
        outer=ttk.Frame(self.root,padding=(8,0,8,8)); outer.pack(fill='both',expand=True)
        outer.grid_columnconfigure(0,weight=1)
        outer.grid_columnconfigure(1,weight=0)
        outer.grid_rowconfigure(1,weight=1)

        self.header_frame=ttk.Frame(outer)
        self.header_frame.grid(row=0,column=0,sticky='ew')
        self.meta_frame=ttk.Frame(outer)  # compatibility only; never packed in V2.4

        self.canvas=tk.Canvas(outer,highlightthickness=0,borderwidth=0,bg='#f5f7fa')
        self.vscroll=ttk.Scrollbar(outer,orient='vertical',command=self.canvas.yview)
        self.canvas.configure(yscrollcommand=self.vscroll.set)
        self.canvas.grid(row=1,column=0,sticky='nsew')
        self.vscroll.grid(row=1,column=1,sticky='ns')

        # Header spacer occupies exactly the same scrollbar column. It is intentionally
        # empty; its sole purpose is to keep header width == canvas viewport width.
        self.header_scroll_spacer=ttk.Frame(outer)
        self.header_scroll_spacer.grid(row=0,column=1,sticky='nsew')

        self.rows_frame=ttk.Frame(self.canvas)
        self.canvas_window=self.canvas.create_window((0,0),window=self.rows_frame,anchor='nw')
        self.rows_frame.bind('<Configure>',lambda e:self.canvas.configure(scrollregion=self.canvas.bbox('all')))
        self.canvas.bind('<Configure>',self._sync_canvas_width)
        self.canvas.bind('<MouseWheel>',self._on_mousewheel); self.root.bind('<Prior>',lambda e:self._page_scroll(-1)); self.root.bind('<Next>',lambda e:self._page_scroll(1))

    def _sync_canvas_width(self,event):
        # The inner grid always uses the visible canvas viewport width, never the outer
        # frame width. This remains correct when Windows DPI/theme changes scrollbar size.
        width=max(1,int(event.width))
        self.canvas.itemconfigure(self.canvas_window,width=width)
        self.root.after_idle(self._sync_header_to_body)

    def _page_scroll(self,d): self.canvas.yview_scroll(d,'pages'); return 'break'
    def _on_mousewheel(self,event):
        delta=-1 if event.delta>0 else 1
        steps=max(1,abs(int(event.delta/120))) if event.delta else 1
        self.canvas.yview_scroll(delta*steps*3,'units'); return 'break'
    def _child_wheel(self,event): return self._on_mousewheel(event)

    def choose_file(self,idx):
        p=self.filedialog.askopenfilename(title=f'문서 {chr(65+idx)} 선택',filetypes=[('지원 문서','*.txt *.docx'),('텍스트','*.txt'),('Word','*.docx')])
        if p: self.file_vars[idx].set(p)
    def choose_multiple(self):
        ps=self.filedialog.askopenfilenames(title='비교할 문서 2~3개 선택',filetypes=[('지원 문서','*.txt *.docx'),('텍스트','*.txt'),('Word','*.docx')])
        if ps: self._assign_files(list(ps))
    def _drop_files(self,event,idx=None):
        try: paths=list(self.root.tk.splitlist(event.data))
        except Exception: paths=[event.data]
        valid=[p for p in paths if Path(p).suffix.lower() in ('.txt','.docx') and Path(p).is_file()]
        if not valid: self.messagebox.showwarning('드래그앤드롭','TXT 또는 DOCX 파일만 지원합니다.'); return 'break'
        if idx is not None:
            self.file_vars[idx].set(valid[0])
        else: self._assign_files(valid)
        return 'break'
    def _assign_files(self,paths):
        paths=paths[:3]
        for i in range(3): self.file_vars[i].set(paths[i] if i<len(paths) else '')
        self.base_var.set(0)
        if len(paths)<2: self.status_var.set('문서를 최소 2개 넣어야 합니다.')
        else: self.status_var.set(f'{len(paths)}개 문서 준비 완료 · 기준문서를 확인한 뒤 비교 시작을 누르세요.')

    def start_compare(self):
        global ACTIVE_COMPARE_MODE
        ACTIVE_COMPARE_MODE={'자동':'auto','일반 문서':'general','법률·규정':'legal'}.get(self.compare_mode_var.get(),'auto')
        indexed=[(i,v.get().strip()) for i,v in enumerate(self.file_vars) if v.get().strip()]
        if len(indexed) not in (2,3): self.messagebox.showwarning('문서 선택','TXT/DOCX 문서를 2개 또는 3개 선택하세요.'); return
        base_slot=self.base_var.get(); slots=[i for i,_ in indexed]
        if base_slot not in slots: self.messagebox.showwarning('기준문서','선택된 문서 중 하나를 기준문서로 지정하세요.'); return
        paths=[p for _,p in indexed]; base_index=slots.index(base_slot)
        self._last_compare_paths=list(paths)
        self.cancel_event.clear(); self.compare_btn.configure(state='disabled'); self.cancel_btn.configure(state='normal'); self.export_btn.configure(state='disabled'); self.word_export_btn.configure(state='disabled')
        self.progress_var.set(1); self.status_var.set('파일을 읽는 중...')
        threading.Thread(target=self._compare_worker,args=(paths,base_index),daemon=True).start()
    def cancel_compare(self): self.cancel_event.set(); self.status_var.set('취소 요청 중...')
    def _progress_from_worker(self,pct,msg): self.root.after(0,lambda:self._set_progress(pct,msg))
    def _set_progress(self,pct,msg): self.progress_var.set(pct); self.status_var.set(msg)
    def _compare_worker(self,paths,base_index):
        try:
            names=[]; texts=[]
            for k,p in enumerate(paths):
                _check_cancel(self.cancel_event); path=Path(p); raw=path.read_bytes()
                if len(raw)>40*1024*1024: raise ValueError(f'{path.name}: 파일당 40MB까지 지원합니다.')
                names.append(path.name); texts.append(read_document(path.name,raw)); self._progress_from_worker(2+int(8*(k+1)/len(paths)),f'{path.name} 읽기 완료')
            result=compare_documents(names,texts,base_index=base_index,progress_cb=self._progress_from_worker,cancel_event=self.cancel_event)
            self.root.after(0,lambda:self._compare_done(result))
        except ComparisonCancelled:
            self.root.after(0,self._compare_cancelled)
        except Exception as e:
            logging.exception('comparison failed'); self.root.after(0,lambda:self._compare_failed(str(e)))
    def _compare_cancelled(self):
        self.compare_btn.configure(state='normal'); self.cancel_btn.configure(state='disabled'); self.progress_var.set(0); self.status_var.set('비교가 취소되었습니다.')
    def _compare_done(self,result):
        self.result=result
        for k,v in result['counts'].items():
            if k in self.stats_vars: self.stats_vars[k].set(str(v))
        self.compare_btn.configure(state='normal'); self.cancel_btn.configure(state='disabled'); self.export_btn.configure(state='normal')
        self.word_export_btn.configure(state=('normal' if len(result.get('names',[]))==2 else 'disabled'))
        self.progress_var.set(100)
        parse_info=result.get('parse_info') or []
        parsed = ' / '.join(f"{chr(65+i)} {x.get('units',0)}블록" for i,x in enumerate(parse_info))
        suffix = f" · 파싱: {parsed}" if parsed else ''
        self.status_var.set(f"비교 완료 · 기준: {result['names'][result['base_index']]}{suffix}")
        self._render_headers(); self.render_rows(); self.canvas.yview_moveto(0)
    def _compare_failed(self,msg):
        self.compare_btn.configure(state='normal'); self.cancel_btn.configure(state='disabled'); self.progress_var.set(0); self.status_var.set('비교 중 오류가 발생했습니다.'); self.messagebox.showerror('비교 오류',msg)

    def _render_headers(self):
        for w in self.header_frame.winfo_children(): w.destroy()
        self.header_labels=[]
        if not self.result: return
        n=len(self.result['names']); bi=self.result.get('base_index',0)
        labels=[]
        for i,name in enumerate(self.result['names']):
            txt=f'{chr(65+i)} · {name}' + ('   [기준]' if i==bi else '')
            labels.append(self.ttk.Label(self.header_frame,text=txt,anchor='center',padding=(8,9),relief='solid',font=('Malgun Gothic',10,'bold' if i==bi else 'normal')))
        labels.append(self.ttk.Label(self.header_frame,text='변경사항',anchor='center',padding=(8,9),relief='solid',font=('Malgun Gothic',10,'bold')))
        self.header_labels=labels
        # place() is intentional: body grid is the source of truth for pixel boundaries.
        # A separate header grid can round columns differently because its children have
        # different requested widths. We copy the body's measured x/width instead.
        self.header_frame.configure(height=40)
        for w in labels:
            w.place(x=0,y=0,width=1,height=40)
        self.root.after_idle(self._sync_header_to_body)

    def _sync_header_to_body(self):
        if not self.result or not getattr(self,'header_labels',None): return
        try:
            self.rows_frame.update_idletasks()
            n=len(self.result['names']); cols=n+1
            if self.rows_frame.winfo_children():
                boxes=[self.rows_frame.grid_bbox(c,0) for c in range(cols)]
                if all(b[2] > 0 for b in boxes):
                    for w,(x,y,width,height) in zip(self.header_labels,boxes):
                        w.place_configure(x=x,y=0,width=width,height=40)
                    return
            # Fallback before rows exist: calculate against the actual canvas viewport.
            width=max(1,self.canvas.winfo_width()); weights=([27,27,27,19] if n==3 else [40,40,20]); total=sum(weights)
            x=0
            for idx,(w,weight) in enumerate(zip(self.header_labels,weights)):
                next_x=width if idx==len(weights)-1 else round(width*sum(weights[:idx+1])/total)
                w.place_configure(x=x,y=0,width=max(1,next_x-x),height=40); x=next_x
        except Exception:
            logging.exception('header/body boundary sync failed')

    def _render_preamble(self):
        """Render document title/chapter text as compact metadata above the article grid.

        Preamble must never behave like Article 0.  It is intentionally excluded from
        the synchronized article rows, because doing so created a large empty first row
        and a meaningless LOW/동일 summary cell.
        """
        for w in self.meta_frame.winfo_children():
            w.destroy()
        try:
            self.meta_frame.pack_forget()
        except Exception:
            pass
        if not self.result:
            return
        prow=None
        for r in self.result.get('rows',[]):
            members=r.get('members') or []
            if any(m and m.get('kind')=='preamble' for m in members):
                prow=r
                break
        if not prow:
            return
        # Place directly after the column-name header and before the scrolling article canvas.
        self.meta_frame.pack(fill='x', after=self.header_frame, pady=(0,4))
        changed=bool(prow.get('changed'))
        bar=self.ttk.Frame(self.meta_frame,padding=(8,4))
        bar.pack(fill='x')
        self.ttk.Label(bar,text='문서 헤더 / 머리말',font=('Malgun Gothic',9,'bold')).pack(side='left')
        self.ttk.Label(bar,text=('변경 있음' if changed else '동일'),style='Sub.TLabel').pack(side='left',padx=(8,0))
        body=self.ttk.Frame(self.meta_frame); body.pack(fill='x')
        n=len(self.result['names']); cols=n+1
        doc_weight=24 if n==3 else 34; sum_weight=28 if n==3 else 32
        for c in range(cols):
            body.grid_columnconfigure(c,weight=sum_weight if c==cols-1 else doc_weight,uniform='c')
        # Keep this strip compact. Long preambles remain readable via the Text widget's own selection/copy,
        # but they do not consume the main legal-comparison viewport.
        max_lines=1
        for m in prow.get('members') or []:
            if m:
                txt=(m.get('body') or m.get('text') or '').strip()
                max_lines=max(max_lines,min(5,max(1,len(txt.splitlines()))))
        h=max(2,min(5,max_lines))
        for i,m in enumerate(prow.get('members') or []):
            t=self._new_text(body,h,bg='#fbfcfe')
            t.configure(padx=9,pady=4)
            if not m:
                t.insert('end','[머리말 없음]','missing')
            else:
                segs=(prow.get('body_segments') or [[] for _ in range(n)])[i]
                if segs:
                    for txt,sty in segs:
                        t.insert('end',txt,sty if sty in ('delete','insert') else ())
                else:
                    t.insert('end',(m.get('body') or m.get('text') or '').strip())
            t.configure(state='disabled')
            t.grid(row=0,column=i,sticky='nsew',padx=(0,1))
        st=self._new_text(body,h,bg='#ffffff')
        st.configure(padx=9,pady=4)
        if changed:
            msgs=prow.get('summary',{}).get('messages') or []
            for msg in msgs[:4]:
                st.insert('end','• '+msg+'\n')
        else:
            st.insert('end','변경 없음','conf')
        st.configure(state='disabled')
        st.grid(row=0,column=cols-1,sticky='nsew')

    def set_filter(self,key):
        self.current_filter=key
        for k,b in self.filter_buttons.items(): b.configure(style='ActiveFilter.TButton' if k==key else 'Filter.TButton')
        self.render_rows(); self.canvas.yview_moveto(0)
    def _row_visible(self,row):
        q=self.search_var.get().strip().lower()
        return not q or q in (' '.join(row['raw'])+' '+' '.join(row['summary']['messages'])).lower()

    def render_rows(self):
        if not self.result: return
        for w in self.rows_frame.winfo_children(): w.destroy()
        rows=[r for r in self.result['rows'] if not any(m and m.get('kind')=='preamble' for m in (r.get('members') or [])) and self._row_visible(r)]
        if not rows:
            self.ttk.Label(self.rows_frame,text='검색 결과가 없습니다.',padding=20).pack(anchor='w'); return
        n=len(self.result['names']); cols=n+1
        doc_weight=27 if n==3 else 40; sum_weight=19 if n==3 else 20
        for c in range(cols):
            self.rows_frame.grid_columnconfigure(c,weight=sum_weight if c==cols-1 else doc_weight,uniform='c')

        grid_row=0
        last_sections=[None]*n
        for row in rows:
            members=row.get('members') or []
            sections=[(members[i].get('section') if i < len(members) and members[i] else '') or '' for i in range(n)]
            # Show a compact structural divider only when at least one document enters a new Chapter/장/절/관.
            changed_section=any(sections[i] and sections[i] != last_sections[i] for i in range(n))
            if changed_section:
                for c in range(cols):
                    bg='#edf2f7'
                    t=self.tk.Text(self.rows_frame,height=1,wrap='none',relief='solid',borderwidth=1,padx=9,pady=3,
                                   font=('Malgun Gothic',9,'bold'),bg=bg,cursor='arrow',takefocus=0)
                    t.bind('<MouseWheel>',self._child_wheel)
                    if c < n:
                        sec=sections[c]
                        if sec and sec != last_sections[c]: t.insert('end',sec)
                    else:
                        vals=[x for x in sections if x]
                        if vals and len(set(vals))>1: t.insert('end','편제/장 구성이 문서별로 다릅니다.')
                    t.configure(state='disabled')
                    t.grid(row=grid_row,column=c,sticky='nsew')
                grid_row+=1
                for i,sec in enumerate(sections):
                    if sec: last_sections[i]=sec

            h=self._estimate_row_height(row)
            for i in range(n): self._make_doc_cell(self.rows_frame,row,i,h).grid(row=grid_row,column=i,sticky='nsew')
            self._make_summary_cell(self.rows_frame,row,h).grid(row=grid_row,column=cols-1,sticky='nsew')
            grid_row+=1
        self.rows_frame.update_idletasks(); self.canvas.configure(scrollregion=self.canvas.bbox('all')); self._sync_header_to_body()

    def _estimate_row_height(self,row):
        vals=[]
        # Width-aware estimate. All cells receive exactly the same height, which is the basis of synchronized scrolling.
        for txt in row['raw']:
            if not txt: vals.append(4); continue
            vals.append(sum(max(1,(len(line)//38)+1) for line in txt.splitlines())+1)
        msgs=sum(max(1,(len(m)//34)+1) for m in row['summary']['messages'])+3
        return max(6,min(90,max(vals+[msgs])))
    def _new_text(self,parent,height,bg='#ffffff'):
        t=self.tk.Text(parent,height=height,wrap='word',relief='solid',borderwidth=1,padx=9,pady=7,font=('Malgun Gothic',10),bg=bg,undo=False,cursor='arrow',takefocus=0)
        t.tag_configure('delete',foreground='#c62828',overstrike=1,font=('Malgun Gothic',10),background='#fff4f4')
        t.tag_configure('insert',foreground='#1565c0',underline=1,font=('Malgun Gothic',10),background='#f2f7ff')
        t.tag_configure('article_header',foreground='#17365d',font=('Malgun Gothic',10,'bold'),spacing3=5)
        t.tag_configure('section',foreground='#667085',font=('Malgun Gothic',8)); t.tag_configure('missing',foreground='#777777',font=('Malgun Gothic',9,'italic'))
        t.tag_configure('high',foreground='#9f1c1c',font=('Malgun Gothic',11,'bold')); t.tag_configure('medium',foreground='#805b00',font=('Malgun Gothic',11,'bold')); t.tag_configure('low',foreground='#236a34',font=('Malgun Gothic',11,'bold'))
        t.tag_configure('tag',foreground='#425466',font=('Malgun Gothic',9,'bold')); t.tag_configure('conf',foreground='#667085',font=('Malgun Gothic',8))
        t.bind('<MouseWheel>',self._child_wheel); return t
    def _make_doc_cell(self,parent,row,i,h):
        t=self._new_text(parent,h); m=row['members'][i]
        if not m: t.insert('end','[해당 조항 없음]','missing')
        else:
            # Header and body both use directional diff: delete=red strike, insert=blue underline.
            hs=row.get('header_segments',[[] for _ in row['segments']])[i] if row.get('header_segments') else []
            if hs:
                for txt,sty in hs: t.insert('end',txt,sty if sty in ('delete','insert') else 'article_header')
                t.insert('end','\n')
            else:
                t.insert('end',m.get('header','')+'\n','article_header')
            for txt,sty in row.get('body_segments',row['segments'])[i]: t.insert('end',txt,sty if sty in ('delete','insert') else ())
        t.configure(state='disabled'); return t
    def _make_summary_cell(self,parent,row,h):
        s=row['summary']; t=self._new_text(parent,h,'#ffffff')
        msgs=s.get('messages') or []
        if s.get('tags')==['동일'] or not row.get('changed'):
            t.insert('end','변경 없음','conf')
        else:
            for msg in msgs:
                t.insert('end','• '+msg+'\n')
        t.configure(state='disabled'); return t

    def export_word_tracked(self):
        """Create a real Microsoft Word comparison document with tracked revisions.

        The selected base document is passed to Word as OriginalDocument and the
        other document as RevisedDocument. Word's own CompareDocuments engine is
        used so the output is a native tracked-changes DOCX, not merely colored text.
        """
        if not self.result:
            return
        if len(self.result.get('names',[])) != 2:
            self.messagebox.showwarning('Word 변경추적','Word 변경추적 내보내기는 2개 문서 비교에서만 사용할 수 있습니다.')
            return
        paths=list(getattr(self,'_last_compare_paths',[]) or [])
        if len(paths) != 2:
            self.messagebox.showwarning('Word 변경추적','비교에 사용한 원본 파일 경로를 찾을 수 없습니다. 다시 비교한 뒤 시도하세요.')
            return
        bi=int(self.result.get('base_index',0))
        oi=1-bi
        base_path=Path(paths[bi])
        revised_path=Path(paths[oi])
        initial=f'{base_path.stem}_vs_{revised_path.stem}_변경추적.docx'
        out=self.filedialog.asksaveasfilename(
            title='Word 변경 내용 추적 문서 저장',
            defaultextension='.docx',
            initialfile=initial,
            filetypes=[('Word 문서','*.docx')]
        )
        if not out:
            return
        self.word_export_btn.configure(state='disabled')
        self.status_var.set('Microsoft Word에서 변경 내용 추적 비교본을 만드는 중...')
        self.progress_var.set(15)
        threading.Thread(
            target=self._word_export_worker,
            args=(base_path,revised_path,Path(out),revised_path.stem),
            daemon=True
        ).start()

    def _word_export_worker(self, base_path, revised_path, out_path, revised_author):
        try:
            create_word_tracked_compare(base_path,revised_path,out_path,revised_author=revised_author)
            self.root.after(0,lambda:self._word_export_done(out_path))
        except Exception as exc:
            logging.exception('word tracked export failed')
            self.root.after(0,lambda:self._word_export_failed(str(exc)))

    def _word_export_done(self, out_path):
        self.progress_var.set(100)
        self.status_var.set(f'Word 변경추적 저장 완료: {Path(out_path).name}')
        if self.result and len(self.result.get('names',[]))==2:
            self.word_export_btn.configure(state='normal')
        self.messagebox.showinfo(
            '저장 완료',
            'Word 변경 내용 추적 문서를 저장했습니다.\n\n'
            f'{out_path}\n\n'
            'Word에서 열고 [검토] 탭의 변경 내용 추적/검토 창에서 삽입·삭제·이동 변경을 확인할 수 있습니다.'
        )

    def _word_export_failed(self, msg):
        self.progress_var.set(0)
        self.status_var.set('Word 변경추적 문서 생성에 실패했습니다.')
        if self.result and len(self.result.get('names',[]))==2:
            self.word_export_btn.configure(state='normal')
        self.messagebox.showerror(
            'Word 변경추적 오류',
            msg + '\n\nMicrosoft Word 데스크톱 앱이 설치되어 있어야 합니다.'
        )

    def export_excel(self):
        if not self.result: return
        p=self.filedialog.asksaveasfilename(title='Excel 비교 결과 저장',defaultextension='.xlsx',initialfile='법률문서_비교결과.xlsx',filetypes=[('Excel 통합 문서','*.xlsx')])
        if not p: return
        try: Path(p).write_bytes(make_xlsx(self.result)); self.status_var.set(f'Excel 저장 완료: {Path(p).name}'); self.messagebox.showinfo('저장 완료',f'Excel 파일을 저장했습니다.\n\n{p}')
        except Exception as e: logging.exception('excel export failed'); self.messagebox.showerror('저장 오류',str(e))



# ---------------- V1.6: robust whole-document article segmentation + staged matching ----------------
APP_VERSION = "1.9"

# English EULAs are common inputs too, so numeric/legal-effect detection covers both Korean and English.
NUMBER_RE = re.compile(
    r"(?:\d{4}\s*[./-]\s*\d{1,2}(?:\s*[./-]\s*\d{1,2})?|"
    r"\d+(?:,\d{3})*(?:\.\d+)?\s*(?:%|퍼센트|원|만원|억원|조원|일|개월|년|시간|분|회|배|"
    r"days?|months?|years?|hours?|minutes?|times?|USD|KRW|dollars?|won))",
    re.IGNORECASE
)
LEGAL_TERMS = list(dict.fromkeys(LEGAL_TERMS + [
    "shall", "must", "may", "shall not", "must not", "liability", "liable", "indemnify", "indemnity",
    "terminate", "termination", "cancel", "cancellation", "damages", "warranty", "disclaimer",
    "governing law", "jurisdiction", "consent", "notice", "obligation", "responsibility"
]))

# Strong legal headings are allowed even when Word stored them in the middle of one paragraph.
# Requiring a parenthesized title for inline English headings avoids splitting ordinary references
# such as "under Article 5" in running text.
ENG_STRONG_INLINE_V16 = re.compile(
    r"(?i)(?<![A-Za-z0-9])Article\s+((?:\d+(?:[-.]\d+)*[A-Za-z]?)|(?:[IVXLCDM]+))\s*\(([^)\n]{1,180})\)"
)
KOR_STRONG_INLINE_V16 = re.compile(
    r"(?<![가-힣A-Za-z0-9])제\s*(\d+)\s*조(?:\s*의\s*(\d+))?\s*(?:\(([^)\n]{1,180})\)|\[([^]\n]{1,180})\])"
)
ENG_LINE_START_V16 = re.compile(
    r"(?im)^[ \t]*Article\s+((?:\d+(?:[-.]\d+)*[A-Za-z]?)|(?:[IVXLCDM]+))\b"
)
KOR_LINE_START_V16 = re.compile(
    r"(?m)^[ \t]*제\s*(\d+)\s*조(?:\s*의\s*(\d+))?\b"
)

_CROSSREF_BEFORE_V16 = re.compile(
    r"(?i)(?:under|pursuant\s+to|according\s+to|in\s+accordance\s+with|defined\s+in|set\s+forth\s+in|see|refer(?:red)?\s+to)\s*$"
)


def _roman_to_int_v16(s: str):
    vals={'I':1,'V':5,'X':10,'L':50,'C':100,'D':500,'M':1000}
    s=(s or '').upper()
    if not s or any(c not in vals for c in s): return None
    total=0; prev=0
    for c in reversed(s):
        v=vals[c]
        if v<prev: total-=v
        else: total+=v; prev=v
    return total


def _article_num_value_v16(s: str):
    s=(s or '').strip()
    if re.fullmatch(r'[IVXLCDM]+',s,re.I): return float(_roman_to_int_v16(s) or 0)
    m=re.match(r'^(\d+)(?:[.-](\d+))?',s)
    if not m: return None
    return float(m.group(1)) + (float(m.group(2))/1000 if m.group(2) else 0)


def _inline_heading_is_crossref_v16(text: str, start: int) -> bool:
    # Inline headings preceded by common reference phrases are almost certainly citations,
    # not actual clause starts. Line-start headings are always retained.
    line_start=text.rfind('\n',0,start)+1
    if not text[line_start:start].strip(): return False
    before=re.sub(r'\s+',' ',text[max(0,start-90):start]).strip()
    return bool(_CROSSREF_BEFORE_V16.search(before))


def _collect_article_markers_v16(text: str):
    markers=[]
    # Strong inline headings: Article N (Title), 제N조(제목)
    for m in ENG_STRONG_INLINE_V16.finditer(text):
        if _inline_heading_is_crossref_v16(text,m.start()):
            continue
        num=m.group(1).strip(); title=(m.group(2) or '').strip()
        markers.append({'start':m.start(),'header_end':m.end(),'number':num,'title':title,
                        'header':f'Article {num}' + (f' ({title})' if title else ''),'lang':'en','strong':True})
    for m in KOR_STRONG_INLINE_V16.finditer(text):
        if _inline_heading_is_crossref_v16(text,m.start()):
            continue
        num=m.group(1)+(f'의{m.group(2)}' if m.group(2) else '')
        title=(m.group(3) or m.group(4) or '').strip()
        markers.append({'start':m.start(),'header_end':m.end(),'number':num,'title':title,
                        'header':f'제{num}조'+(f'({title})' if title else ''),'lang':'ko','strong':True})

    strong_starts=[m['start'] for m in markers]
    def overlaps_strong(pos):
        return any(abs(pos-s)<=3 for s in strong_starts)

    # Titleless / oddly formatted headings are still recognized at actual line starts.
    for m in ENG_LINE_START_V16.finditer(text):
        if overlaps_strong(m.start()): continue
        num=m.group(1).strip()
        markers.append({'start':m.start(),'header_end':m.end(),'number':num,'title':'',
                        'header':f'Article {num}','lang':'en','strong':False})
    for m in KOR_LINE_START_V16.finditer(text):
        if overlaps_strong(m.start()): continue
        num=m.group(1)+(f'의{m.group(2)}' if m.group(2) else '')
        markers.append({'start':m.start(),'header_end':m.end(),'number':num,'title':'',
                        'header':f'제{num}조','lang':'ko','strong':False})

    # De-duplicate candidates that start at the same place; prefer the strong titled form.
    by_start={}
    for m in sorted(markers,key=lambda x:(x['start'],not x['strong'])):
        old=by_start.get(m['start'])
        if old is None or (m['strong'] and not old['strong']): by_start[m['start']]=m
    markers=sorted(by_start.values(),key=lambda x:x['start'])

    # If inline strong matches accidentally caught a quoted/cross-referenced titled Article,
    # a duplicate/out-of-order number usually exposes it. Drop suspicious inline duplicates
    # when a later candidate of the same number is at a true line start.
    cleaned=[]
    for i,m in enumerate(markers):
        same_later=[x for x in markers[i+1:] if x['number'].lower()==m['number'].lower()]
        line_start=text.rfind('\n',0,m['start'])+1
        inline=bool(text[line_start:m['start']].strip())
        if inline and same_later:
            later=same_later[0]
            later_ls=text.rfind('\n',0,later['start'])+1
            if not text[later_ls:later['start']].strip():
                continue
        cleaned.append(m)

    # V1.9: reject likely line-start cross references that create an impossible local jump.
    # Example: a titleless "Article 27" captured between genuine Article 4 and Article 5.
    # Strong/titled headings are never removed by this heuristic.
    filtered=[]
    for i,m in enumerate(cleaned):
        suspicious=False
        if (not m.get('strong')) and 0 < i < len(cleaned)-1:
            pv=_article_num_value_v16(cleaned[i-1].get('number',''))
            cv=_article_num_value_v16(m.get('number',''))
            nv=_article_num_value_v16(cleaned[i+1].get('number',''))
            if pv is not None and cv is not None and nv is not None:
                # Neighbours form a normal local sequence while the middle candidate is a large
                # outlier.  This is characteristic of a cross-reference paragraph, not a heading.
                neighbours_close = 0 < (nv-pv) <= 2.1
                far_from_both = abs(cv-pv) >= 5 and abs(cv-nv) >= 5
                outside_interval = not (min(pv,nv) <= cv <= max(pv,nv))
                if neighbours_close and far_from_both and outside_interval:
                    suspicious=True
        if not suspicious:
            filtered.append(m)
    return filtered


def normalize_text(s: str) -> str:
    # Canonical Unicode normalization is required BEFORE regex/token cleanup.
    # Without this, visually identical Korean (e.g. NFC '석' vs decomposed NFD Jamo)
    # can be treated as different text and produce bogus '석 → 석' changes.
    s=unicodedata.normalize('NFC', s or '')
    # Remove only leading/heading identifiers, not ordinary cross references inside a body.
    s=re.sub(r'(?i)^\s*Article\s+(?:\d+(?:[-.]\d+)*[A-Za-z]?|[IVXLCDM]+)\b',' ',s)
    s=re.sub(r'^\s*제\s*\d+\s*조(?:\s*의\s*\d+)?',' ',s)
    s=re.sub(r'\s+','',s)
    s=re.sub(r'[^0-9A-Za-z가-힣%]','',s)
    return s.lower()


def parse_units(text: str) -> List[Unit]:
    """Robust article-first parser.

    Unlike V1.5, article headings do not have to be at a line/paragraph boundary. This is
    essential for DOCX files created by copy/paste, conversion, or soft line breaks where
    several Article headings can live inside one Word paragraph.
    """
    text=(text or '').replace('\r\n','\n').replace('\r','\n')
    markers=_collect_article_markers_v16(text)
    units=[]

    if markers:
        pre=text[:markers[0]['start']].strip()
        if pre:
            units.append(Unit(
                index=len(units),kind='preamble',number='',title='문서 머리말',header='문서 머리말',
                body=pre,text=pre,section='',norm_title='문서머리말',norm_body=normalize_text(pre),
                structure=structure_sig(pre)
            ))
        for idx,m in enumerate(markers):
            end=markers[idx+1]['start'] if idx+1<len(markers) else len(text)
            body=text[m['header_end']:end]
            body=re.sub(r'^[ \t]*(?:[:\-–—])[ \t]*','',body)
            body=body.strip()
            full=m['header'] if not body else m['header']+'\n'+body
            units.append(Unit(
                index=len(units),kind='article',number=m['number'],title=m['title'],header=m['header'],
                body=body,text=full,section='',norm_title=normalize_text(m['title']),
                norm_body=normalize_text(body),structure=structure_sig(body)
            ))
        return units

    # No formal article headings: retain paragraph fallback instead of turning the whole file
    # into one giant preamble block.
    paras=[re.sub(r'[ \t]+',' ',x).strip() for x in text.split('\n') if x.strip()]
    if not paras:
        return []
    for i,p in enumerate(paras):
        units.append(Unit(i,'paragraph',str(i+1),f'문단 {i+1}',f'문단 {i+1}',p,p,'',
                          f'문단{i+1}',normalize_text(p),structure_sig(p)))
    return units


def _title_ratio_v16(a: Unit,b: Unit) -> float:
    if not a.norm_title or not b.norm_title: return 0.0
    return seq_ratio(a.norm_title,b.norm_title)


def _body_anchor_score_v16(a: Unit,b: Unit) -> float:
    return max(seq_ratio(a.norm_body,b.norm_body), jaccard(a.body,b.body))



_TITLE_STOPWORDS_V17 = {
    'the','a','an','of','for','to','and','or','etc','etcetera','on','in','regarding','concerning',
    'article','section','clause'
}


def _title_concept_tokens_v17(title: str) -> set:
    """Return coarse legal-heading concepts rather than literal words.

    This intentionally maps inflection/nominalisation variants such as
    provide/providing/provision to the same concept so a heavily rewritten article is not
    split into deletion+addition merely because its heading wording changed.
    """
    words=re.findall(r"[A-Za-z]+|[가-힣]+", (title or '').lower())
    out=set()
    for w in words:
        if w in _TITLE_STOPWORDS_V17: continue
        # English derivational families commonly seen in legal headings.
        if w.startswith(('provid','provis')): w='provide'
        elif w.startswith('inform'): w='information'
        elif w.startswith(('obligat','duti')): w='obligation'
        elif w.startswith(('terminat','cancel')): w='termination'
        elif w.startswith(('compensat','damage')): w='damages'
        elif w.startswith(('confidenti','secret')): w='confidentiality'
        elif w.startswith(('person','privacy')): w='privacy'
        elif w.startswith(('use','usage')): w='use'
        elif w.startswith(('pay','payment')): w='payment'
        elif w.startswith(('applic','scope')): w='scope'
        elif w.startswith(('defin','meaning')): w='definition'
        elif w.startswith(('amend','revis','modif','chang')): w='amendment'
        elif w.startswith(('company','corporat')): w='company'
        elif w.startswith(('member','user')): w='user'
        # very light normalization for Korean heading variants
        w=re.sub(r'(등|관련|관한|관하여|사항)$','',w) or w
        out.add(w)
    return out


def _title_semantic_ratio_v17(a: Unit, b: Unit) -> float:
    """Heading similarity robust to wording changes and derivational variants."""
    literal=_title_ratio_v16(a,b)
    ta=_title_concept_tokens_v17(a.title); tb=_title_concept_tokens_v17(b.title)
    if not ta or not tb:
        return literal
    jac=len(ta & tb)/len(ta | tb)
    contain=len(ta & tb)/max(1,min(len(ta),len(tb)))
    # Containment matters for shortened headings: "Providing Company Information, etc."
    # -> "Provision of Information" keeps the core concepts while dropping modifiers.
    concept=max(jac,0.92*contain)
    return max(literal, concept)


def _same_number_context_support_v17(base: List[Unit], other: List[Unit], i: int, j: int) -> float:
    """Structural support for same-number matching from neighbouring article numbers.

    If surrounding clauses retain the same numbering, a same-number clause whose text was
    comprehensively rewritten is still usually the same lineage.  This is support only; it
    never overrides a clearly contradictory heading by itself.
    """
    support=0.0; checks=0
    for di in (-1,1):
        ai=i+di; bj=j+di
        if 0 <= ai < len(base) and 0 <= bj < len(other):
            aa,bb=base[ai],other[bj]
            if aa.kind=='article' and bb.kind=='article' and aa.number and bb.number:
                checks+=1
                if aa.number.lower()==bb.number.lower(): support+=1.0
                elif _title_semantic_ratio_v17(aa,bb)>=0.72: support+=0.75
    return support/checks if checks else 0.0

def match_units_fast(base: List[Unit], other: List[Unit], threshold: float = 0.38,
                     cancel_event=None, progress_cb=None) -> Tuple[Dict[int, Tuple[int,float]], set]:
    """Staged legal-clause matching.

    1) preamble↔preamble, 2) same article number when title/body supports it,
    3) same/near-identical title (captures renumbering), 4) content similarity candidates.
    This prevents an existing Article from being reported as missing simply because a greedy
    global similarity candidate consumed its counterpart first.
    """
    mapping={}; used_b=set(); used_a=set()
    number_idx=defaultdict(list); title_idx=defaultdict(list); token_idx=defaultdict(set)
    for j,b in enumerate(other):
        if b.kind=='article' and b.number: number_idx[b.number.lower()].append(j)
        if b.kind=='article' and b.norm_title: title_idx[b.norm_title].append(j)
        for tok in set(token_words((b.title or '')+' '+b.body)):
            if len(tok)>=2: token_idx[tok].add(j)

    # Preamble only matches preamble.
    for i,a in enumerate(base):
        if a.kind!='preamble': continue
        c=[j for j,b in enumerate(other) if b.kind=='preamble' and j not in used_b]
        if c:
            j=c[0]; s=max(0.75,similarity(a,other[j])); mapping[i]=(j,s); used_a.add(i); used_b.add(j)

    # Stage 1: exact article number, but require corroboration so renumbered documents do not
    # blindly match unrelated clauses that merely inherited the old number.
    for i,a in enumerate(base):
        _check_cancel(cancel_event)
        if i in used_a or a.kind!='article': continue
        cands=[j for j in number_idx.get(a.number.lower(),[]) if j not in used_b]
        if not cands: continue
        scored=[]
        for j in cands:
            b=other[j]; tr=_title_semantic_ratio_v17(a,b); body_seq=seq_ratio(a.norm_body,b.norm_body); body_jac=jaccard(a.body,b.body); s=similarity(a,b)
            context=_same_number_context_support_v17(base,other,i,j)
            # Same article number is a strong lineage clue. A substantially rewritten clause
            # may have low body similarity, so a semantically related heading OR aligned
            # neighbours is sufficient corroboration.
            corroborated=(a.norm_title and b.norm_title and tr>=0.46) or body_seq>=0.50 or body_jac>=0.20 or context>=0.75
            if not a.title and not b.title:
                corroborated = body_seq>=0.46 or body_jac>=0.18 or context>=0.90
            if corroborated:
                anchor=max(s,0.93 if tr>=0.82 else 0.86 if tr>=0.60 else 0.78 if tr>=0.46 else 0.70)
                scored.append((anchor,j))
        if scored:
            s,j=max(scored); mapping[i]=(j,min(1.0,s)); used_a.add(i); used_b.add(j)

    # Stage 2: exact normalized title -> strongest signal for moved/renumbered clauses.
    for i,a in enumerate(base):
        _check_cancel(cancel_event)
        if i in used_a or a.kind!='article' or not a.norm_title: continue
        cands=[j for j in title_idx.get(a.norm_title,[]) if j not in used_b]
        if cands:
            j=max(cands,key=lambda x: similarity(a,other[x]))
            s=max(0.88,similarity(a,other[j])); mapping[i]=(j,min(1.0,s)); used_a.add(i); used_b.add(j)

    # Stage 3: similar titles, even if punctuation/wording shifted slightly.
    title_candidates=[]
    for i,a in enumerate(base):
        if i in used_a or a.kind!='article' or not a.norm_title: continue
        for j,b in enumerate(other):
            if j in used_b or b.kind!='article' or not b.norm_title: continue
            tr=_title_semantic_ratio_v17(a,b)
            if tr>=0.64:
                s=0.78*tr+0.22*_body_anchor_score_v16(a,b)
                title_candidates.append((s,i,j))
    for s,i,j in sorted(title_candidates,reverse=True):
        if i in used_a or j in used_b: continue
        mapping[i]=(j,max(0.72,min(0.95,s))); used_a.add(i); used_b.add(j)

    # Stage 4: content-based matching for the remaining moved/retitled clauses.
    candidates=[]; total=max(1,len(base))
    # Build lookup helpers for the existing candidate-pruning function.
    title_set_idx=defaultdict(set); number_set_idx=defaultdict(set)
    for j,b in enumerate(other):
        if b.norm_title: title_set_idx[b.norm_title].add(j)
        if b.number: number_set_idx[b.number].add(j)
    for i,a in enumerate(base):
        _check_cancel(cancel_event)
        if i in used_a or a.kind=='preamble': continue
        for j in _candidate_indices(a,i,len(base),other,title_set_idx,number_set_idx,token_idx):
            if j in used_b or other[j].kind=='preamble': continue
            s=similarity(a,other[j])
            if s>=threshold:
                proximity=1.0-min(abs(i-j)/max(len(base),len(other),1),1.0)
                candidates.append((s+0.012*proximity,i,j,s))
        if progress_cb and i % max(1,total//20)==0: progress_cb(i/total)
    for _,i,j,s in sorted(candidates,reverse=True):
        if i in used_a or j in used_b: continue
        mapping[i]=(j,s); used_a.add(i); used_b.add(j)
    return mapping,used_b


def render_segments_html_v19(segments: List[Tuple[str,str]]) -> str:
    parts=[]
    for txt,sty in segments:
        esc=html.escape(txt)
        if sty=='delete':
            parts.append(f'<span style="color:#c62828;text-decoration:line-through">{esc}</span>')
        elif sty=='insert':
            parts.append(f'<span style="color:#1565c0;text-decoration:underline">{esc}</span>')
        else:
            parts.append(esc)
    return ''.join(parts).replace('\n','<br>')


def compare_documents(names: List[str], texts: List[str], base_index: int=0, progress_cb=None, cancel_event=None) -> Dict[str,Any]:
    if progress_cb: progress_cb(3,'DOCX/TXT 전체에서 조 제목을 찾아 자동 분할하는 중...')
    units=[]; parse_info=[]
    for i,t in enumerate(texts):
        _check_cancel(cancel_event); parsed=parse_units(t); units.append(parsed)
        article_count=sum(1 for u in parsed if u.kind=='article')
        pre_count=sum(1 for u in parsed if u.kind=='preamble')
        parse_info.append({'articles':article_count,'units':len(parsed),'preamble':pre_count})
        if progress_cb:
            progress_cb(5+int(15*(i+1)/len(texts)),f'문서 {chr(65+i)} · 조항 {article_count}개 / 전체 비교단위 {len(parsed)}개')
    if progress_cb: progress_cb(22,'조 번호·제목·내용을 단계적으로 대조하는 중...')
    def match_progress(x):
        if progress_cb: progress_cb(22+int(33*x),'동일/이동 조항 대응 관계를 찾는 중...')
    groups=build_groups(units,base_index=base_index,cancel_event=cancel_event,progress_cb=match_progress)
    rows=[]; counts={'total':0,'changed':0,'added':0,'deleted':0,'moved':0,'high':0}; total=max(1,len(groups))
    for ridx,g in enumerate(groups):
        _check_cancel(cancel_event); members=g['members']; scores=g['scores']
        body_segs=directional_segments_v19(members,base_index,'body')
        header_segs=directional_segments_v19(members,base_index,'header')
        full_segs=[]; htmls=[]; raw=[]
        for i,m in enumerate(members):
            if m:
                hs=header_segs[i] or [(m.header,'normal')]
                fs=hs+[('\n','normal')]+body_segs[i] if m.kind!='preamble' else [('문서 머리말\n','normal')]+body_segs[i]
                full_segs.append(fs); htmls.append(render_segments_html_v19(fs)); raw.append(m.text)
            else:
                full_segs.append([]); htmls.append('<span class="missing">[해당 조항 없음]</span>'); raw.append('')
        summary=summarize_group(members,scores,base_index=base_index)
        if g.get('axis_extra') and members[base_index] is None:
            summary['messages'].insert(0,'기준문서 외 신설 조항입니다. 기준문서 순서를 보존하기 위해 하단에 별도 표시합니다.')
            if '신설' not in summary['tags']: summary['tags'].append('신설')
        changed=summary['tags']!=['동일']
        is_preamble=any(m and m.kind=='preamble' for m in members)
        # Header/preamble metadata is displayed separately and must not inflate legal-article statistics.
        if not is_preamble:
            counts['total']+=1
            if changed: counts['changed']+=1
            if '신설' in summary['tags']: counts['added']+=1
            if '삭제' in summary['tags']: counts['deleted']+=1
            if '이동' in summary['tags']: counts['moved']+=1
            if summary['risk']=='HIGH': counts['high']+=1
        rows.append({'id':ridx,'raw':raw,'html':htmls,'segments':full_segs,'body_segments':body_segs,'header_segments':header_segs,
                     'members':[asdict(m) if m else None for m in members],'summary':summary,'changed':changed,'axis_extra':bool(g.get('axis_extra'))})
        if progress_cb and ridx % max(1,total//25)==0:
            progress_cb(58+int(40*(ridx+1)/total),f'조별 세부 변경 분석 중... {ridx+1}/{total}')
    if progress_cb: progress_cb(100,'조별 비교 완료')
    return {'names':names,'rows':rows,'counts':counts,'unit_counts':[len(x) for x in units],
            'parse_info':parse_info,'base_index':base_index}



# ---------------- V2.0: strict baseline heading parser ----------------
APP_VERSION = "2.8"

def _v20_line_bounds(text: str, pos: int):
    a=text.rfind('\n',0,pos)+1
    b=text.find('\n',pos)
    if b < 0: b=len(text)
    return a,b


def _v20_looks_like_heading_title(line: str, lang: str='en') -> bool:
    s=re.sub(r'\s+',' ',(line or '')).strip(' \t:-–—')
    if not s or len(s) > 160: return False
    if re.match(r'(?i)^(?:Article\s+\w+|Section\s+\w+|Clause\s+\w+)',s): return False
    if re.match(r'^제\s*\d+\s*조',s): return False
    if re.match(r'^(?:\(?\d+\)?[.)]|[①-⑳]|[가-하][.)])\s*',s): return False
    # A real title is normally a short noun phrase, not a complete operative sentence.
    if s.endswith(('.', ';', '?', '!')): return False
    if lang=='en':
        words=re.findall(r"[A-Za-z][A-Za-z'/-]*",s)
        if not words or len(words)>18: return False
        # Strong sentence/opening-clause signals.
        if re.match(r'(?i)^(?:the|this|these|a|an|if|when|where|unless|provided|company|member|user|service)\s+',s):
            if re.search(r'(?i)\b(?:shall|must|may|is|are|means|agrees?|provides?|applies?|will|can)\b',s): return False
        if re.search(r'(?i)\b(?:shall|must|may|means|will|hereby)\b',s): return False
    return True


def _v20_attach_split_titles(text: str, markers):
    """Fold a title stored on the same/next DOCX line into the Article header.

    Examples treated identically:
      Article 2 (Definition of Terms)
      Article 2\nDefinition of Terms
      Article 2 Definition of Terms
    The consumed title is removed from the body so it cannot be reported as newly inserted text.
    """
    out=[]
    for idx,m0 in enumerate(markers):
        m=dict(m0)
        if m.get('title'):
            out.append(m); continue
        end_limit=markers[idx+1]['start'] if idx+1<len(markers) else len(text)
        _,line_end=_v20_line_bounds(text,m['header_end'])
        line_end=min(line_end,end_limit)
        same=text[m['header_end']:line_end].strip()
        chosen=''; consume_end=m['header_end']
        if _v20_looks_like_heading_title(same,m.get('lang','en')):
            chosen=same; consume_end=line_end
        else:
            # If Article N is alone on its line, inspect exactly the next non-empty line.
            probe=line_end+1 if line_end < end_limit else line_end
            while probe < end_limit:
                nend=text.find('\n',probe,end_limit)
                if nend<0: nend=end_limit
                cand=text[probe:nend].strip()
                if cand:
                    if _v20_looks_like_heading_title(cand,m.get('lang','en')):
                        chosen=cand; consume_end=nend
                    break
                probe=nend+1
        if chosen:
            chosen=re.sub(r'^\((.*)\)$',r'\1',chosen).strip()
            m['title']=chosen
            if m.get('lang')=='ko':
                m['header']=f"제{m['number']}조({chosen})"
            else:
                m['header']=f"Article {m['number']} ({chosen})"
            m['header_end']=consume_end
            m['strong']=True
        out.append(m)
    return out


def _v20_marker_line_continuation(text: str, m) -> str:
    """Text remaining on the same source line after the captured heading."""
    _,e=_v20_line_bounds(text,m['header_end'])
    return re.sub(r'\s+',' ',text[m['header_end']:e]).strip()


def _v20_filter_false_article_refs(text: str, markers):
    """Remove Article references misread as headings without disturbing the baseline axis.

    Main failure this prevents: Article 4 -> [reference to Article 27] -> Article 5.
    A huge local number jump surrounded by a consecutive pair is treated as a cross-reference,
    even if the reference happens to contain a parenthesized title.
    """
    if len(markers)<3: return markers
    keep=[True]*len(markers)

    # Duplicate-number evidence: when an out-of-order occurrence appears before the genuine
    # occurrence of the same Article later in the document, the early one is almost certainly a citation.
    bynum=defaultdict(list)
    for i,m in enumerate(markers):
        if m.get('lang')=='en': bynum[(m.get('number') or '').lower()].append(i)

    for i in range(1,len(markers)-1):
        m=markers[i]
        if m.get('lang')!='en': continue
        pv=_article_num_value_v16(markers[i-1].get('number',''))
        cv=_article_num_value_v16(m.get('number',''))
        nv=_article_num_value_v16(markers[i+1].get('number',''))
        if pv is None or cv is None or nv is None: continue
        neighbours_consecutive=(0.5 <= (nv-pv) <= 1.5)
        outside=not (min(pv,nv) <= cv <= max(pv,nv))
        huge_jump=min(abs(cv-pv),abs(cv-nv)) >= 3
        if neighbours_consecutive and outside and huge_jump:
            key=(m.get('number') or '').lower()
            duplicate_later=any(k>i for k in bynum.get(key,[]))
            continuation=_v20_marker_line_continuation(text,m)
            sentence_cont=bool(continuation and not _v20_looks_like_heading_title(continuation,'en'))
            # The local 4 -> 27 -> 5 pattern is itself sufficiently strong evidence.  A later
            # genuine Article 27 or sentence continuation makes the decision even safer.
            if duplicate_later or sentence_cont or min(abs(cv-pv),abs(cv-nv))>=5:
                keep[i]=False

    # A second pass catches multiple citations between the same genuine consecutive headings.
    # Build a monotonic backbone from surviving numeric headings; isolated large backward/forward
    # spikes are excluded when the next surviving heading resumes the expected sequence.
    changed=True
    while changed:
        changed=False
        idxs=[i for i,k in enumerate(keep) if k]
        for z in range(1,len(idxs)-1):
            i=idxs[z]; pi=idxs[z-1]; ni=idxs[z+1]
            m=markers[i]
            if m.get('lang')!='en': continue
            pv=_article_num_value_v16(markers[pi].get('number',''))
            cv=_article_num_value_v16(m.get('number',''))
            nv=_article_num_value_v16(markers[ni].get('number',''))
            if None in (pv,cv,nv): continue
            if 0.5 <= nv-pv <= 1.5 and not (min(pv,nv)<=cv<=max(pv,nv)) and min(abs(cv-pv),abs(cv-nv))>=4:
                keep[i]=False; changed=True; break
    return [m for i,m in enumerate(markers) if keep[i]]


def _collect_article_markers_v20(text: str):
    # Start with V1.9's broad detector, then apply V2.0 structural validation.
    markers=[]
    for m in ENG_STRONG_INLINE_V16.finditer(text):
        if _inline_heading_is_crossref_v16(text,m.start()): continue
        num=m.group(1).strip(); title=(m.group(2) or '').strip()
        markers.append({'start':m.start(),'header_end':m.end(),'number':num,'title':title,
                        'header':f'Article {num}'+(f' ({title})' if title else ''),'lang':'en','strong':True})
    for m in KOR_STRONG_INLINE_V16.finditer(text):
        if _inline_heading_is_crossref_v16(text,m.start()): continue
        num=m.group(1)+(f'의{m.group(2)}' if m.group(2) else '')
        title=(m.group(3) or m.group(4) or '').strip()
        markers.append({'start':m.start(),'header_end':m.end(),'number':num,'title':title,
                        'header':f'제{num}조'+(f'({title})' if title else ''),'lang':'ko','strong':True})
    strong_starts=[x['start'] for x in markers]
    def overlap(pos): return any(abs(pos-s)<=3 for s in strong_starts)
    for m in ENG_LINE_START_V16.finditer(text):
        if overlap(m.start()): continue
        num=m.group(1).strip()
        markers.append({'start':m.start(),'header_end':m.end(),'number':num,'title':'','header':f'Article {num}','lang':'en','strong':False})
    for m in KOR_LINE_START_V16.finditer(text):
        if overlap(m.start()): continue
        num=m.group(1)+(f'의{m.group(2)}' if m.group(2) else '')
        markers.append({'start':m.start(),'header_end':m.end(),'number':num,'title':'','header':f'제{num}조','lang':'ko','strong':False})

    by_start={}
    for m in sorted(markers,key=lambda x:(x['start'],not x['strong'])):
        old=by_start.get(m['start'])
        if old is None or (m['strong'] and not old['strong']): by_start[m['start']]=m
    markers=sorted(by_start.values(),key=lambda x:x['start'])

    # Inline duplicate before a true line-start occurrence of the same number is a reference.
    cleaned=[]
    for i,m in enumerate(markers):
        later=[x for x in markers[i+1:] if x['number'].lower()==m['number'].lower()]
        ls=text.rfind('\n',0,m['start'])+1
        inline=bool(text[ls:m['start']].strip())
        if inline and later:
            l=later[0]; lls=text.rfind('\n',0,l['start'])+1
            if not text[lls:l['start']].strip(): continue
        cleaned.append(m)

    cleaned=_v20_filter_false_article_refs(text,cleaned)
    cleaned=_v20_attach_split_titles(text,cleaned)
    return cleaned


_ENG_HIER_RE_V23 = re.compile(
    r'(?im)^[ \t]*(Chapter|Part)\s+((?:\d+(?:[-.]\d+)*)|(?:[IVXLCDM]+))\s*[.\-:–—]?\s*([^\n]{0,180})$'
)
_KOR_HIER_RE_V23 = re.compile(
    r'(?m)^[ \t]*제\s*(\d+)\s*(장|절|관)\s*(?:\(([^)\n]{0,180})\)|\[([^]\n]{0,180})\]|([^\n]{0,180}))$'
)


def _collect_hierarchy_markers_v23(text: str):
    out=[]
    for m in _ENG_HIER_RE_V23.finditer(text):
        level=m.group(1).lower(); num=m.group(2).strip(); title=(m.group(3) or '').strip()
        label=f'{m.group(1).title()} {num}' + (f'. {title}' if title else '')
        out.append({'start':m.start(),'end':m.end(),'level':level,'label':label})
    for m in _KOR_HIER_RE_V23.finditer(text):
        num=m.group(1); level=m.group(2); title=(m.group(3) or m.group(4) or m.group(5) or '').strip()
        label=f'제{num}{level}' + (f' {title}' if title else '')
        out.append({'start':m.start(),'end':m.end(),'level':level,'label':label})
    return sorted(out,key=lambda x:x['start'])


def _section_for_pos_v23(hmarks, pos: int) -> str:
    # Maintain a hierarchy stack. English Part is outer to Chapter; Korean 장 > 절 > 관.
    stack={}
    for h in hmarks:
        if h['start'] >= pos: break
        lv=h['level']; stack[lv]=h['label']
        if lv in ('part','장'):
            for k in ('chapter','절','관'): stack.pop(k,None)
        elif lv in ('chapter','절'):
            stack.pop('관',None)
    vals=[]
    for k in ('part','chapter','장','절','관'):
        if k in stack and stack[k] not in vals: vals.append(stack[k])
    return ' > '.join(vals)


def _strip_hierarchy_from_body_v23(body: str) -> str:
    body=_ENG_HIER_RE_V23.sub('',body)
    body=_KOR_HIER_RE_V23.sub('',body)
    body=re.sub(r'\n{3,}','\n\n',body)
    return body.strip()


def parse_units(text: str) -> List[Unit]:
    """V2.3 parser: Article/조 + Chapter/장/절/관 hierarchy, with preamble excluded from comparison."""
    text=(text or '').replace('\r\n','\n').replace('\r','\n')
    markers=_collect_article_markers_v20(text)
    hmarks=_collect_hierarchy_markers_v23(text)
    units=[]
    if markers:
        for idx,m in enumerate(markers):
            end=markers[idx+1]['start'] if idx+1<len(markers) else len(text)
            body=text[m['header_end']:end]
            body=_strip_hierarchy_from_body_v23(body)
            body=re.sub(r'^[ \t]*(?:[:\-–—])[ \t]*','',body).strip()
            section=_section_for_pos_v23(hmarks,m['start'])
            full=m['header'] if not body else m['header']+'\n'+body
            units.append(Unit(index=len(units),kind='article',number=m['number'],title=m['title'],header=m['header'],
                              body=body,text=full,section=section,norm_title=normalize_text(m['title']),
                              norm_body=normalize_text(body),structure=structure_sig(body)))
        return units
    # Documents without formal articles still fall back to paragraphs, but hierarchy-only lines are excluded.
    raw_lines=[re.sub(r'[ \t]+',' ',x).strip() for x in text.split('\n') if x.strip()]
    paras=[]
    for line in raw_lines:
        if _ENG_HIER_RE_V23.fullmatch(line) or _KOR_HIER_RE_V23.fullmatch(line):
            continue
        paras.append(line)
    for i,p in enumerate(paras):
        units.append(Unit(i,'paragraph',str(i+1),f'문단 {i+1}',f'문단 {i+1}',p,p,'',f'문단{i+1}',normalize_text(p),structure_sig(p)))
    return units



# ---------------- V2.5: robust DOCX numbering + hierarchy parsing ----------------

def _v25_numpr_for_paragraph(paragraph):
    """Return (numId, ilvl) including numbering inherited from the paragraph style."""
    from docx.oxml.ns import qn
    def from_ppr(ppr):
        if ppr is None:
            return None
        numPr = ppr.find(qn('w:numPr'))
        if numPr is None:
            return None
        numId = numPr.find(qn('w:numId'))
        ilvl = numPr.find(qn('w:ilvl'))
        if numId is None:
            return None
        try:
            nid = int(numId.get(qn('w:val')))
            lvl = int(ilvl.get(qn('w:val'))) if ilvl is not None else 0
            return nid, lvl
        except Exception:
            return None
    got = from_ppr(paragraph._p.pPr)
    if got:
        return got
    try:
        st = paragraph.style
        seen=set()
        while st is not None and id(st) not in seen:
            seen.add(id(st))
            got = from_ppr(st.element.pPr)
            if got:
                return got
            st = st.base_style
    except Exception:
        pass
    return None


def _v25_numbering_defs(doc):
    """Read numbering.xml sufficiently to reconstruct visible automatic numbering."""
    from docx.oxml.ns import qn
    try:
        root = doc.part.numbering_part.element
    except Exception:
        return {}, {}
    abs_defs={}
    for an in root.findall(qn('w:abstractNum')):
        try: aid=int(an.get(qn('w:abstractNumId')))
        except Exception: continue
        levels={}
        for lv in an.findall(qn('w:lvl')):
            try: ilvl=int(lv.get(qn('w:ilvl')))
            except Exception: ilvl=0
            def childval(tag, default=''):
                x=lv.find(qn(tag))
                return x.get(qn('w:val')) if x is not None and x.get(qn('w:val')) is not None else default
            try: start=int(childval('w:start','1'))
            except Exception: start=1
            levels[ilvl]={'start':start,'fmt':childval('w:numFmt','decimal'),'text':childval('w:lvlText',f'%{ilvl+1}.')}
        abs_defs[aid]=levels
    num_map={}
    for num in root.findall(qn('w:num')):
        try: nid=int(num.get(qn('w:numId')))
        except Exception: continue
        a=num.find(qn('w:abstractNumId'))
        if a is None: continue
        try: aid=int(a.get(qn('w:val')))
        except Exception: continue
        overrides={}
        for ov in num.findall(qn('w:lvlOverride')):
            try: ilvl=int(ov.get(qn('w:ilvl')))
            except Exception: continue
            so=ov.find(qn('w:startOverride'))
            if so is not None:
                try: overrides[ilvl]=int(so.get(qn('w:val')))
                except Exception: pass
        num_map[nid]=(aid,overrides)
    return abs_defs,num_map


def _v25_roman(n:int)->str:
    if n<=0: return str(n)
    vals=((1000,'M'),(900,'CM'),(500,'D'),(400,'CD'),(100,'C'),(90,'XC'),(50,'L'),(40,'XL'),(10,'X'),(9,'IX'),(5,'V'),(4,'IV'),(1,'I'))
    out=[]
    for v,s in vals:
        while n>=v: out.append(s); n-=v
    return ''.join(out)


def _v25_alpha(n:int)->str:
    if n<=0: return str(n)
    out=[]
    while n:
        n-=1; out.append(chr(65+n%26)); n//=26
    return ''.join(reversed(out))


def _v25_fmt_num(n:int, fmt:str)->str:
    f=(fmt or 'decimal').lower()
    if f=='upperroman': return _v25_roman(n)
    if f=='lowerroman': return _v25_roman(n).lower()
    if f=='upperletter': return _v25_alpha(n)
    if f=='lowerletter': return _v25_alpha(n).lower()
    return str(n)


def _v25_number_label(doc, paragraph, state, defs, num_map):
    got=_v25_numpr_for_paragraph(paragraph)
    if not got: return ''
    nid,ilvl=got
    if nid==0 or nid not in num_map: return ''
    aid,overrides=num_map[nid]; levels=defs.get(aid,{})
    ld=levels.get(ilvl)
    if not ld: return ''
    counters=state.setdefault(nid,{})
    # A new item at this level invalidates all deeper counters.
    for k in list(counters):
        if k>ilvl: counters.pop(k,None)
    start=overrides.get(ilvl,ld.get('start',1))
    counters[ilvl]=counters.get(ilvl,start-1)+1
    pattern=ld.get('text') or f'%{ilvl+1}.'
    def repl(m):
        lvl=int(m.group(1))-1
        val=counters.get(lvl, levels.get(lvl,{}).get('start',1))
        fmt=levels.get(lvl,{}).get('fmt','decimal')
        return _v25_fmt_num(val,fmt)
    return re.sub(r'%(\d+)',repl,pattern).strip()


def _v25_para_xml_text(paragraph):
    """Include soft line breaks, hyperlinks and drawing/text-box text that Paragraph.text may omit."""
    from docx.oxml.ns import qn
    out=[]
    for el in paragraph._p.iter():
        if el.tag==qn('w:t') and el.text:
            out.append(el.text)
        elif el.tag==qn('w:tab'):
            out.append('\t')
        elif el.tag in (qn('w:br'),qn('w:cr')):
            out.append('\n')
    return ''.join(out).strip()


def extract_docx_text(data: bytes) -> str:
    """V2.5 DOCX extraction: preserve Word automatic numbering used for Chapter/Article headings."""
    doc=Document(io.BytesIO(data))
    defs,num_map=_v25_numbering_defs(doc); num_state={}
    lines=[]
    for block in iter_docx_blocks(doc):
        if isinstance(block, Paragraph):
            txt=_v25_para_xml_text(block) or block.text.strip()
            if not txt: continue
            label=_v25_number_label(doc,block,num_state,defs,num_map)
            if label:
                compact=lambda x: re.sub(r'\s+','',x).lower()
                if not compact(txt).startswith(compact(label)):
                    txt=(label+' '+txt).strip()
            lines.append(txt)
        else:
            for row in block.rows:
                cells=[]
                for c in row.cells:
                    cell_parts=[]
                    for p in c.paragraphs:
                        txt=_v25_para_xml_text(p) or p.text.strip()
                        if not txt: continue
                        label=_v25_number_label(doc,p,num_state,defs,num_map)
                        if label and not re.sub(r'\s+','',txt).lower().startswith(re.sub(r'\s+','',label).lower()):
                            txt=(label+' '+txt).strip()
                        cell_parts.append(txt)
                    cells.append(' '.join(cell_parts).strip())
                if any(cells): lines.append(' | '.join(cells))
    return '\n'.join(lines)


_V25_ENG_HIER = re.compile(r'(?i)^\s*(Chapter|Part)\s+((?:\d+(?:[-.]\d+)*)|(?:[IVXLCDM]+))\s*[.\-:–—]?\s*(.*?)\s*$')
_V25_KOR_HIER = re.compile(r'^\s*제\s*(\d+)\s*(장|절|관)\s*(?:\(([^)]*)\)|\[([^]]*)\]|(.*?))\s*$')


def _v25_title_like_hierarchy_tail(line:str)->bool:
    s=(line or '').strip(' \t|')
    if not s or len(s)>160: return False
    if _V25_ENG_HIER.match(s) or _V25_KOR_HIER.match(s): return False
    if re.match(r'(?i)^Article\s+\w+\b',s) or re.match(r'^제\s*\d+\s*조',s): return False
    # Do not absorb a normal legal sentence as the Chapter title.
    words=re.findall(r"[A-Za-z가-힣]+",s)
    if len(words)>14: return False
    if re.search(r'[.;!?]\s*$',s): return False
    if re.match(r'(?i)^(?:the|a|an|if|when|where|provided|notwithstanding)\b',s) and len(words)>6: return False
    return True


def _collect_hierarchy_markers_v25(text:str):
    """Recognize literal, split-line, table-cell and Word-auto-numbered Chapter/Part/장/절/관 headings."""
    out=[]; lines=[]; pos=0
    for raw in text.splitlines(True):
        content=raw.rstrip('\r\n'); lines.append((pos,content)); pos+=len(raw)
    if not lines and text: lines=[(0,text)]

    def add_marker(start,end,level,num,title,kind):
        title=(title or '').strip(' \t|.-:–—')
        if kind=='en':
            word='Chapter' if level=='chapter' else 'Part'
            label=f'{word} {num}' + (f'. {title}' if title else '')
        else:
            label=f'제{num}{level}' + (f' {title}' if title else '')
        out.append({'start':start,'end':end,'level':level,'label':label})

    for idx,(start,line) in enumerate(lines):
        # DOCX table extraction separates cells with |. Treat each cell as a possible heading.
        cursor=0
        segments=re.split(r'(\|)',line)
        for seg in segments:
            if seg=='|': cursor+=1; continue
            seg_start=start+cursor; stripped=seg.strip(); lead=len(seg)-len(seg.lstrip())
            if stripped:
                em=_V25_ENG_HIER.match(stripped)
                km=_V25_KOR_HIER.match(stripped)
                if em:
                    level=em.group(1).lower(); num=em.group(2); title=(em.group(3) or '').strip()
                    end=seg_start+len(seg)
                    # `Chapter 1` on one paragraph and `General Provisions` on the next paragraph.
                    if not title:
                        # A DOCX table may store `Chapter 1` and `General Provisions` in adjacent cells.
                        tail=line[cursor+len(seg):]
                        if '|' in tail:
                            cand=tail.split('|',1)[1].split('|',1)[0].strip(' \t|')
                            if _v25_title_like_hierarchy_tail(cand):
                                title=cand; end=start+len(line)
                        if not title:
                            for j in range(idx+1,min(len(lines),idx+3)):
                                cand=lines[j][1].strip(' \t|')
                                if not cand: continue
                                if _v25_title_like_hierarchy_tail(cand):
                                    title=cand; end=lines[j][0]+len(lines[j][1])
                                break
                    add_marker(seg_start+lead,end,level,num,title,'en')
                elif km:
                    num=km.group(1); level=km.group(2); title=(km.group(3) or km.group(4) or km.group(5) or '').strip()
                    end=seg_start+len(seg)
                    if not title:
                        tail=line[cursor+len(seg):]
                        if '|' in tail:
                            cand=tail.split('|',1)[1].split('|',1)[0].strip(' \t|')
                            if _v25_title_like_hierarchy_tail(cand):
                                title=cand; end=start+len(line)
                        if not title:
                            for j in range(idx+1,min(len(lines),idx+3)):
                                cand=lines[j][1].strip(' \t|')
                                if not cand: continue
                                if _v25_title_like_hierarchy_tail(cand):
                                    title=cand; end=lines[j][0]+len(lines[j][1])
                                break
                    add_marker(seg_start+lead,end,level,num,title,'ko')
            cursor+=len(seg)

    # Last-resort for Word content that was flattened into the preamble before Article 1.
    first_article=min([m['start'] for m in _collect_article_markers_v20(text)] or [len(text)])
    prefix=text[:first_article]
    for m in re.finditer(r'(?i)\b(Chapter|Part)\s+((?:\d+(?:[-.]\d+)*)|(?:[IVXLCDM]+))\s*[.\-:–—]\s*([A-Z][A-Za-z0-9 ,/&\-]{2,100})',prefix):
        if any(abs(m.start()-x['start'])<5 for x in out): continue
        title=m.group(3).strip()
        # Cut at obvious next heading or hard delimiter.
        title=re.split(r'\s{2,}|\||(?=\bArticle\s+\w+\b)',title)[0].strip()
        if _v25_title_like_hierarchy_tail(title):
            add_marker(m.start(),m.end(),m.group(1).lower(),m.group(2),title,'en')

    # de-duplicate while preserving earliest document position
    uniq={}
    for h in sorted(out,key=lambda x:(x['start'],x['end'])):
        key=(h['start'],h['level'])
        if key not in uniq or len(h['label'])>len(uniq[key]['label']): uniq[key]=h
    return sorted(uniq.values(),key=lambda x:x['start'])


def _strip_hierarchy_from_body_v25(body:str)->str:
    # Remove both same-line and split-line hierarchy headers from an Article body slice.
    lines=body.splitlines(); out=[]; skip_next=False
    for i,line in enumerate(lines):
        if skip_next:
            skip_next=False; continue
        s=line.strip(' \t|')
        em=_V25_ENG_HIER.match(s); km=_V25_KOR_HIER.match(s)
        if em or km:
            title=(em.group(3) if em else (km.group(3) or km.group(4) or km.group(5) or '')).strip()
            if not title and i+1<len(lines) and _v25_title_like_hierarchy_tail(lines[i+1]): skip_next=True
            continue
        out.append(line)
    return re.sub(r'\n{3,}','\n\n','\n'.join(out)).strip()


def parse_units(text: str) -> List[Unit]:
    """V2.5 parser: baseline Article order + robust Chapter/Part/장/절/관 detection."""
    text=(text or '').replace('\r\n','\n').replace('\r','\n')
    markers=_collect_article_markers_v20(text)
    hmarks=_collect_hierarchy_markers_v25(text)
    units=[]
    if markers:
        for idx,m in enumerate(markers):
            end=markers[idx+1]['start'] if idx+1<len(markers) else len(text)
            body=_strip_hierarchy_from_body_v25(text[m['header_end']:end])
            body=re.sub(r'^[ \t]*(?:[:\-–—])[ \t]*','',body).strip()
            section=_section_for_pos_v23(hmarks,m['start'])
            full=m['header'] if not body else m['header']+'\n'+body
            units.append(Unit(len(units),'article',m['number'],m['title'],m['header'],body,full,section,
                              normalize_text(m['title']),normalize_text(body),structure_sig(body)))
        return units
    raw=[x.strip() for x in text.splitlines() if x.strip()]
    for i,p in enumerate(raw):
        if _V25_ENG_HIER.match(p) or _V25_KOR_HIER.match(p): continue
        units.append(Unit(i,'paragraph',str(i+1),f'문단 {i+1}',f'문단 {i+1}',p,p,'',f'문단{i+1}',normalize_text(p),structure_sig(p)))
    return units


def _v25_clean_piece(tokens):
    s=''.join(tokens)
    s=re.sub(r'\s+',' ',s).strip()
    return s.strip(' ,;')


def _summary_tokens_v26(text: str):
    """Tokens for human-readable change summaries.

    Whitespace and standalone punctuation are intentionally excluded. Word documents
    often split quotation marks / punctuation across runs or use typographic variants,
    which previously produced nonsense such as 추가 " / 삭제 " / 추가 ".
    The visible document still preserves the original punctuation; the summary focuses
    on substantive word/number changes.
    """
    return re.findall(r"[가-힣A-Za-z]+(?:['’][A-Za-z]+)?|\d+(?:,\d{3})*(?:\.\d+)?%?", text, re.UNICODE)


def _v26_compact_ops(old: str, new: str, limit: int = 12):
    a=_summary_tokens_v26(old); b=_summary_tokens_v26(new)
    ak=[x.casefold() for x in a]; bk=[x.casefold() for x in b]
    sm=SequenceMatcher(None,ak,bk,autojunk=False)
    raw=[]
    for tag,i1,i2,j1,j2 in sm.get_opcodes():
        if tag=='equal':
            continue
        o=' '.join(a[i1:i2]).strip(); n=' '.join(b[j1:j2]).strip()
        if not o and not n:
            continue
        raw.append([tag,o,n])

    # Merge adjacent delete/insert pairs into one replacement. SequenceMatcher can
    # emit them separately around repeated words; users should see one coherent change.
    merged=[]; i=0
    while i < len(raw):
        tag,o,n=raw[i]
        if tag=='delete' and i+1 < len(raw) and raw[i+1][0]=='insert':
            merged.append(('replace',o,raw[i+1][2])); i+=2; continue
        if tag=='insert' and i+1 < len(raw) and raw[i+1][0]=='delete':
            merged.append(('replace',raw[i+1][1],n)); i+=2; continue
        merged.append((tag,o,n)); i+=1

    out=[]
    for tag,o,n in merged:
        if tag=='insert' and n:
            out.append(f'추가 “{_short_v15(n,150)}”')
        elif tag=='delete' and o:
            out.append(f'삭제 “{_short_v15(o,150)}”')
        elif tag=='replace':
            if o and n:
                out.append(f'변경 “{_short_v15(o,120)}” → “{_short_v15(n,120)}”')
            elif o:
                out.append(f'삭제 “{_short_v15(o,150)}”')
            elif n:
                out.append(f'추가 “{_short_v15(n,150)}”')
        if len(out)>=limit:
            break
    return out


def _v25_body_delta_messages(old:str,new:str,limit:int=12):
    """V2.8: concise substantive changes only; punctuation/run noise is suppressed."""
    return _v26_compact_ops(old,new,limit=limit)


_summarize_group_v24 = summarize_group

def summarize_group(members: List[Optional[Unit]], scores: List[Optional[float]], base_index: int=0) -> Dict[str,Any]:
    """V2.8 concise summary of substantive text changes; punctuation/run noise is suppressed."""
    present=[i for i,m in enumerate(members) if m]
    base=members[base_index]; msgs=[]; tags=[]; risk='LOW'
    if base is None:
        docs=', '.join(chr(65+i) for i in present)
        msgs.append(f'문서 {docs}: 기준문서에 없는 새 조항')
        tags=['신설']; risk='MEDIUM'
    elif len(present)==1:
        msgs.append('비교문서에서 이 조항이 삭제됨')
        tags=['삭제']; risk='HIGH'
    else:
        for i,m in enumerate(members):
            if i==base_index: continue
            label=chr(65+i)
            if m is None:
                msgs.append(f'문서 {label}: 조항 삭제')
                if '삭제' not in tags: tags.append('삭제')
                risk='HIGH'; continue
            if base.number!=m.number and base.kind=='article' and m.kind=='article':
                msgs.append(f'문서 {label}: {base.header} → {m.header} (이동/재번호화)')
                if '이동' not in tags: tags.append('이동')
            if base.title and m.title and base.norm_title!=m.norm_title:
                msgs.append(f'문서 {label}: 조항명 “{base.title}” → “{m.title}”')
                if '제목변경' not in tags: tags.append('제목변경')
            if base.section!=m.section and (base.section or m.section):
                msgs.append(f'문서 {label}: 편제 “{base.section or "없음"}” → “{m.section or "없음"}”')
                if '편제이동' not in tags: tags.append('편제이동')
            if base.norm_body!=m.norm_body:
                if '내용변경' not in tags: tags.append('내용변경')
                for d in _v25_body_delta_messages(base.body,m.body):
                    msgs.append(f'문서 {label}: {d}')
                risk='MEDIUM'
    if len(members)==3 and base is not None:
        others=[i for i in range(3) if i!=base_index]
        m1,m2=members[others[0]],members[others[1]]
        if m1 and m2 and m1.norm_body!=m2.norm_body:
            msgs.append(f'문서 {chr(65+others[0])}와 {chr(65+others[1])}: 서로 다른 수정 내용')
            if '비교본차이' not in tags: tags.append('비교본차이')
    if not msgs:
        msgs=['변경 없음']; tags=['동일']
    return {'messages':msgs,'tags':tags,'risk':risk,'confidence':None}



# ---------------- V2.8: hierarchy boundary + punctuation-diff correctness ----------------

_V27_INVIS_TRANS = str.maketrans({
    '\u00a0':' ', '\u2007':' ', '\u202f':' ',
    '\u200b':' ', '\u200c':' ', '\u200d':' ', '\ufeff':' ', '\u00ad':' '
})

_V27_ENG_HIER_LINE = re.compile(
    r'(?im)^[ \t]*(?:[•●▪◦·*]+[ \t]*)?(Chapter|Part)[ \t]*'
    r'((?:\d+(?:[-.]\d+)*)|(?:[IVXLCDM]+))[ \t]*[.\-:–—]?[ \t]*(.*?)[ \t]*$'
)
_V27_KOR_HIER_LINE = re.compile(
    r'(?m)^[ \t]*(?:[•●▪◦·*]+[ \t]*)?제[ \t]*(\d+)[ \t]*(장|절|관)[ \t]*'
    r'(?:\(([^)\n]*)\)|\[([^]\n]*)\]|(.*?))[ \t]*$'
)
# Strong inline form. This is intentionally conservative: a separator after the number
# is required, so prose such as "under Chapter 3 of ..." is not treated as a heading.
_V27_ENG_HIER_INLINE = re.compile(
    r'(?i)\b(Chapter|Part)[ \t]*((?:\d+(?:[-.]\d+)*)|(?:[IVXLCDM]+))[ \t]*'
    r'[.\-:–—][ \t]*([A-Z][A-Za-z0-9][A-Za-z0-9 ,/&\-()]{1,150})'
)


def _v27_norm_same_length(text: str) -> str:
    """Normalize invisible Word characters without changing string offsets."""
    return (text or '').translate(_V27_INVIS_TRANS)


def _v27_title_clean(s: str) -> str:
    s=(s or '').translate(_V27_INVIS_TRANS).strip(' \t|.-:–—')
    # A title recovered by the inline matcher must stop before the next hard heading.
    s=re.split(r'(?=\bArticle\s+\d+\b)|(?=\bChapter\s+\d+\b)|(?=\bPart\s+\d+\b)|\|',s,1,flags=re.I)[0]
    return s.strip()


def _v27_heading_tail_ok(s: str) -> bool:
    s=_v27_title_clean(s)
    if not s or len(s)>160: return False
    if re.match(r'(?i)^(?:Article|Chapter|Part)\s+\w+\b',s): return False
    if re.match(r'^제\s*\d+\s*(?:조|장|절|관)',s): return False
    words=re.findall(r'[A-Za-z가-힣]+',s)
    if not words or len(words)>16: return False
    # Typical legal chapter titles are noun phrases, not full sentences.
    if re.search(r'[!?;]\s*$',s): return False
    if len(words)>8 and re.match(r'(?i)^(?:the|a|an|if|when|where|provided|notwithstanding|company|member)\b',s):
        return False
    return True


def _collect_hierarchy_markers_v27(text: str):
    """Find Chapter/Part/장/절/관 robustly, including Word soft-break/invisible-char cases."""
    raw=text or ''
    norm=_v27_norm_same_length(raw)
    out=[]

    def add(start,end,level,num,title,kind):
        title=_v27_title_clean(title)
        if kind=='en':
            word='Chapter' if level=='chapter' else 'Part'
            label=f'{word} {num}' + (f'. {title}' if title else '')
        else:
            label=f'제{num}{level}' + (f' {title}' if title else '')
        out.append({'start':start,'end':end,'level':level,'label':label})

    # 1) Proper line / soft-line-break headings.
    for m in _V27_ENG_HIER_LINE.finditer(norm):
        level=m.group(1).lower(); num=m.group(2); title=_v27_title_clean(m.group(3))
        end=m.end()
        if not title:
            # absorb the next non-empty title line, while retaining exact offsets
            cursor=end
            for nm in re.finditer(r'[^\r\n]+',norm[end:end+400]):
                cand=nm.group(0).strip(' \t|')
                if not cand: continue
                if _v27_heading_tail_ok(cand):
                    title=_v27_title_clean(cand); end=end+nm.end()
                break
        add(m.start(),end,level,num,title,'en')

    for m in _V27_KOR_HIER_LINE.finditer(norm):
        num=m.group(1); level=m.group(2); title=_v27_title_clean(m.group(3) or m.group(4) or m.group(5))
        end=m.end()
        if not title:
            for nm in re.finditer(r'[^\r\n]+',norm[end:end+400]):
                cand=nm.group(0).strip(' \t|')
                if not cand: continue
                if _v27_heading_tail_ok(cand):
                    title=_v27_title_clean(cand); end=end+nm.end()
                break
        add(m.start(),end,level,num,title,'ko')

    # 2) Strong inline chapter heading. Handles a DOCX paragraph where the visible line break
    # was flattened or stored in an unusual run. Only accept it at a plausible boundary.
    for m in _V27_ENG_HIER_INLINE.finditer(norm):
        # already represented by a line-heading marker
        if any(abs(m.start()-h['start'])<=4 for h in out):
            continue
        prefix=norm[max(0,m.start()-12):m.start()]
        at_boundary=(m.start()==0 or '\n' in prefix or '\r' in prefix or '\x0b' in prefix or '\x0c' in prefix
                     or bool(re.search(r'[.!?][ \t]{0,4}$',prefix)))
        if not at_boundary:
            continue
        title=_v27_title_clean(m.group(3))
        # Inline regex is greedy within one visual line; stop before sentence-looking material.
        title=re.split(r'\s{2,}|(?=\bArticle\s+\d+\b)',title,1,flags=re.I)[0].strip()
        if not _v27_heading_tail_ok(title):
            continue
        add(m.start(),m.start()+len(m.group(0)),m.group(1).lower(),m.group(2),title,'en')

    # 3) Reuse V2.5 markers (automatic numbering / table cases) as an additional source.
    try:
        out.extend(_collect_hierarchy_markers_v25(raw))
    except Exception:
        pass

    # de-duplicate nearby representations of the same heading; prefer the most complete label/span.
    out=sorted(out,key=lambda h:(h['start'],h['end']))
    merged=[]
    for h in out:
        found=None
        for x in reversed(merged[-4:]):
            if x['level']==h['level'] and abs(x['start']-h['start'])<=5:
                found=x; break
        if found is None:
            merged.append(dict(h))
        else:
            if len(h['label'])>len(found['label']): found['label']=h['label']
            found['end']=max(found['end'],h['end'])
            found['start']=min(found['start'],h['start'])
    return sorted(merged,key=lambda h:h['start'])


def _v27_strip_hierarchy_by_spans(text: str, start: int, end: int, hmarks) -> str:
    """Remove recognized hierarchy headings from an Article slice using absolute spans."""
    spans=[]
    for h in hmarks:
        hs,he=h['start'],h['end']
        if hs>=start and hs<end:
            spans.append((max(start,hs),min(end,he)))
    if not spans:
        return text[start:end].strip()
    pieces=[]; cur=start
    for a,b in sorted(spans):
        if a>cur: pieces.append(text[cur:a])
        cur=max(cur,b)
    if cur<end: pieces.append(text[cur:end])
    body=''.join(pieces)
    body=re.sub(r'^[ \t]*(?:[:\-–—])[ \t]*','',body)
    body=re.sub(r'\n{3,}','\n\n',body)
    return body.strip()


def parse_units(text: str) -> List[Unit]:
    """V2.8 parser: hierarchy boundaries are first-class and never remain inside Article bodies."""
    text=(text or '').replace('\r\n','\n').replace('\r','\n')
    markers=_collect_article_markers_v20(text)
    hmarks=_collect_hierarchy_markers_v27(text)
    units=[]
    if markers:
        for idx,m in enumerate(markers):
            end=markers[idx+1]['start'] if idx+1<len(markers) else len(text)
            body=_v27_strip_hierarchy_by_spans(text,m['header_end'],end,hmarks)
            section=_section_for_pos_v23(hmarks,m['start'])
            full=m['header'] if not body else m['header']+'\n'+body
            units.append(Unit(len(units),'article',m['number'],m['title'],m['header'],body,full,section,
                              normalize_text(m['title']),normalize_text(body),structure_sig(body)))
        return units
    # Paragraph fallback, excluding all recognized hierarchy heading spans/lines.
    masked=list(text)
    for h in hmarks:
        for i in range(max(0,h['start']),min(len(masked),h['end'])):
            if masked[i] not in '\r\n': masked[i]=' '
    residual=''.join(masked)
    raw=[x.strip() for x in residual.splitlines() if x.strip()]
    for i,p in enumerate(raw):
        units.append(Unit(i,'paragraph',str(i+1),f'문단 {i+1}',f'문단 {i+1}',p,p,'',f'문단{i+1}',normalize_text(p),structure_sig(p)))
    return units


# V2.8 summary diff: lexical anchors first, punctuation by the gap between anchors.
# Repeated punctuation such as quotation marks must NOT be independently aligned by
# SequenceMatcher.  Otherwise equally-valid quote matches can create nonsense like
# add quote / delete quote / add quote when one quote merely moved to a new word boundary.
_V28_LEX_RE = re.compile(
    r"[가-힣A-Za-z]+(?:['’][A-Za-z]+)?|\\d+(?:,\\d{3})*(?:\\.\\d+)?%?",
    re.UNICODE,
)


def _v28_lex_spans(text: str):
    text=text or ''
    return [(m.group(0),m.start(),m.end()) for m in _V28_LEX_RE.finditer(text)]


def _v28_raw_lex_piece(text: str, spans, i1: int, i2: int) -> str:
    if i1>=i2: return ''
    return text[spans[i1][1]:spans[i2-1][2]].strip()


def _v28_punct_in_range(text: str, start: int, end: int, spans, si1: int, si2: int) -> str:
    """Return exact non-whitespace punctuation in a gap after masking lexical spans."""
    if end<=start: return ''
    seg=list(text[start:end])
    for k in range(si1,si2):
        _,a,b=spans[k]
        aa=max(a,start)-start; bb=min(b,end)-start
        for x in range(max(0,aa),max(0,bb)):
            if x < len(seg): seg[x]=' '
    return ''.join(ch for ch in seg if (not ch.isspace()) and (not ch.isalnum()) and ch!='_')


def _v28_anchor_label(left_tok: str|None, right_tok: str|None) -> str:
    # Kept internal for future GUI context; summaries stay concise for now.
    if left_tok and right_tok: return f'{left_tok}|{right_tok}'
    if right_tok: return f'앞|{right_tok}'
    if left_tok: return f'{left_tok}|뒤'
    return ''


def _v26_compact_ops(old: str, new: str, limit: int = 12):
    """Location-aware text diff used by the right-hand change summary.

    1) Words/numbers are aligned first.
    2) Punctuation is compared inside the exact gap between matched lexical anchors.
    3) Therefore a quote moved from before ``the`` to before ``Company`` becomes exactly
       one delete + one add, never add/delete/add from ambiguous repeated quote tokens.
    4) A visually identical but different Unicode punctuation code point is still kept as
       one delete + one add, as requested.
    """
    old=old or ''; new=new or ''
    a=_v28_lex_spans(old); b=_v28_lex_spans(new)
    ak=[t[0].casefold() for t in a]; bk=[t[0].casefold() for t in b]
    sm=SequenceMatcher(None,ak,bk,autojunk=False)

    # Build ordered equal-token anchor pairs. Sentinels delimit the whole string.
    pairs=[]
    for block in sm.get_matching_blocks():
        for n in range(block.size):
            pairs.append((block.a+n,block.b+n))
    pairs.sort()
    anchors=[(-1,-1)] + pairs + [(len(a),len(b))]

    events=[]  # (position, priority, text)
    for z in range(len(anchors)-1):
        ai,bj=anchors[z]; an,bn=anchors[z+1]
        ao1=ai+1; ao2=an; bo1=bj+1; bo2=bn

        # Lexical change between the two equal anchors.
        op=_v28_raw_lex_piece(old,a,ao1,ao2)
        np=_v28_raw_lex_piece(new,b,bo1,bo2)
        pos=(a[ai][2] if ai>=0 else 0)
        if op or np:
            if op and np:
                events.append((pos,20,f'변경 “{_short_v15(op,120)}” → “{_short_v15(np,120)}”'))
            elif op:
                events.append((pos,20,f'삭제 “{_short_v15(op,150)}”'))
            else:
                events.append((pos,20,f'추가 “{_short_v15(np,150)}”'))

        # Exact punctuation in the entire gap, masking any lexical insert/delete content.
        old_start=(a[ai][2] if ai>=0 else 0)
        old_end=(a[an][1] if an<len(a) else len(old))
        new_start=(b[bj][2] if bj>=0 else 0)
        new_end=(b[bn][1] if bn<len(b) else len(new))
        old_p=_v28_punct_in_range(old,old_start,old_end,a,ao1,ao2)
        new_p=_v28_punct_in_range(new,new_start,new_end,b,bo1,bo2)
        if old_p != new_p:
            # One gap can contribute at most one delete and one add.  This is the rule
            # that prevents add/delete/add noise for repeated quotes.
            if old_p:
                events.append((pos,10,f'삭제 “{_short_v15(old_p,80)}”'))
            if new_p:
                events.append((pos,11,f'추가 “{_short_v15(new_p,80)}”'))

    # Stable document order; punctuation delete/add at a location appears before the
    # lexical insertion/replacement that follows it.
    events.sort(key=lambda x:(x[0],x[1]))
    out=[]
    for _,_,msg in events:
        # Avoid only exact duplicate events emitted from the same logical gap.  Do not
        # globally erase repeated punctuation changes at genuinely different positions.
        if out and out[-1]==msg:
            continue
        out.append(msg)
        if len(out)>=limit: break
    return out[:limit]



# ---------------- V2.9: lineage-first article/subpart alignment ----------------
APP_VERSION = "3.0"


def _v29_title_tokens(title: str) -> set:
    words=re.findall(r"[A-Za-z]+|[가-힣]+", (title or '').lower())
    out=set()
    for w in words:
        if w in _TITLE_STOPWORDS_V17:
            continue
        if w.startswith(('provid','provis')): w='provide'
        elif w.startswith('inform'): w='information'
        elif w.startswith(('operat','policy')): w='operation'
        elif w.startswith(('protect','privacy','person')): w='privacy'
        elif w.startswith(('obligat','duti')): w='obligation'
        elif w.startswith(('terminat','cancel')): w='termination'
        elif w.startswith(('compensat','damage')): w='damages'
        elif w.startswith(('confidenti','secret')): w='confidentiality'
        elif w.startswith(('use','usage')): w='use'
        elif w.startswith(('pay','payment')): w='payment'
        elif w.startswith(('applic','scope')): w='scope'
        elif w.startswith(('defin','meaning')): w='definition'
        elif w.startswith(('amend','revis','modif','chang')): w='amendment'
        elif w.startswith(('company','corporat')): w='company'
        elif w.startswith(('member','user')): w='user'
        w=re.sub(r'(등|관련|관한|관하여|사항)$','',w) or w
        out.add(w)
    return out


def _v29_title_similarity(a: Unit, b: Unit) -> float:
    literal=_title_semantic_ratio_v17(a,b)
    ta=_v29_title_tokens(a.title); tb=_v29_title_tokens(b.title)
    if not ta or not tb:
        return literal
    jac=len(ta & tb)/max(1,len(ta | tb))
    contain=len(ta & tb)/max(1,min(len(ta),len(tb)))
    return max(literal,jac,0.94*contain)


def _v29_body_similarity(a: Unit, b: Unit) -> float:
    if a.norm_body and a.norm_body==b.norm_body:
        return 1.0
    seq=seq_ratio(a.norm_body,b.norm_body)
    wa=set(token_words(a.body)); wb=set(token_words(b.body))
    if not wa or not wb:
        return seq
    jac=len(wa & wb)/max(1,len(wa | wb))
    contain=len(wa & wb)/max(1,min(len(wa),len(wb)))
    # Containment helps when an old detailed clause was condensed or expanded.
    return max(seq,0.88*contain,0.82*jac)


def _v29_lineage_similarity(a: Unit, b: Unit) -> float:
    if a.kind!='article' or b.kind!='article':
        return similarity(a,b)
    tr=_v29_title_similarity(a,b) if (a.title and b.title) else 0.0
    br=_v29_body_similarity(a,b)
    st=struct_similarity(a.structure,b.structure)
    if a.title and b.title:
        score=0.54*tr+0.40*br+0.06*st
    else:
        score=0.90*br+0.10*st
    # Article number is deliberately only a tiny tiebreak.  A newly inserted clause may
    # renumber every following article; number equality must never overpower title/body lineage.
    if a.number and b.number and a.number.lower()==b.number.lower():
        score+=0.015
    if tr>=0.88: score=max(score,0.84)
    if br>=0.88: score=max(score,0.84)
    if tr>=0.72 and br>=0.35: score=max(score,0.74)
    return min(1.0,score)


def _v29_monotonic_article_match(base: List[Unit], other: List[Unit], cancel_event=None, progress_cb=None):
    """Global sequence alignment for legal article lineage.

    Insertions/deletions are gaps, so a new Article 5 can shift old Article 5->6, 6->7, ...
    without same-number clauses being falsely paired.  Residual high-confidence pairs are
    matched afterwards as true non-monotonic moves.
    """
    ai=[i for i,u in enumerate(base) if u.kind=='article']
    bj=[j for j,u in enumerate(other) if u.kind=='article']
    n,m=len(ai),len(bj)
    gap=0.48
    dp=[[0.0]*(m+1) for _ in range(n+1)]
    prev=[[None]*(m+1) for _ in range(n+1)]
    for i in range(1,n+1): dp[i][0]=dp[i-1][0]+gap; prev[i][0]='D'
    for j in range(1,m+1): dp[0][j]=dp[0][j-1]+gap; prev[0][j]='I'
    for i in range(1,n+1):
        _check_cancel(cancel_event)
        a=base[ai[i-1]]
        for j in range(1,m+1):
            b=other[bj[j-1]]; sim=_v29_lineage_similarity(a,b)
            # Pairing unrelated clauses is more expensive than a delete+insert (0.96).
            mc=(0.94*(1.0-sim)) if sim>=0.43 else 1.08
            choices=[(dp[i-1][j-1]+mc,'M'),(dp[i-1][j]+gap,'D'),(dp[i][j-1]+gap,'I')]
            dp[i][j],prev[i][j]=min(choices,key=lambda x:x[0])
        if progress_cb and i % max(1,n//20 or 1)==0:
            progress_cb(0.80*i/max(1,n))
    mapping={}; used=set(); modes={}
    i,j=n,m; ops=[]
    while i or j:
        op=prev[i][j]
        if op=='M': ops.append(('M',i-1,j-1)); i-=1; j-=1
        elif op=='D': ops.append(('D',i-1,None)); i-=1
        else: ops.append(('I',None,j-1)); j-=1
    ops.reverse()
    for op,x,y in ops:
        if op!='M': continue
        bi,oj=ai[x],bj[y]; sim=_v29_lineage_similarity(base[bi],other[oj])
        if sim<0.43: continue
        mapping[bi]=(oj,sim); used.add(oj)
        modes[bi]='same' if base[bi].number.lower()==other[oj].number.lower() else 'renumbered'

    # High-confidence residual matching allows an article to move to a different position.
    used_a=set(mapping); residual=[]
    for bi in ai:
        if bi in used_a: continue
        for oj in bj:
            if oj in used: continue
            a,b=base[bi],other[oj]
            sim=_v29_lineage_similarity(a,b); tr=_v29_title_similarity(a,b); br=_v29_body_similarity(a,b)
            if sim>=0.70 or tr>=0.82 or (tr>=0.68 and br>=0.48) or br>=0.84:
                residual.append((sim+0.10*tr+0.04*br,bi,oj,sim))
    for _,bi,oj,sim in sorted(residual,reverse=True):
        if bi in mapping or oj in used: continue
        mapping[bi]=(oj,sim); used.add(oj); modes[bi]='moved'
    # If the final mapping contains an inversion, those pairs are true positional moves,
    # not merely sequential renumbering caused by an insertion/deletion.
    pairs=sorted((bi,oj) for bi,(oj,_) in mapping.items())
    moved_bi=set()
    for x in range(len(pairs)):
        for y in range(x+1,len(pairs)):
            if pairs[x][1] > pairs[y][1]:
                moved_bi.add(pairs[x][0]); moved_bi.add(pairs[y][0])
    for bi in moved_bi:
        modes[bi]='moved'
    if progress_cb: progress_cb(1.0)
    return mapping,used,modes


def match_units_fast(base: List[Unit], other: List[Unit], threshold: float = 0.38,
                     cancel_event=None, progress_cb=None) -> Tuple[Dict[int, Tuple[int,float]], set]:
    mapping,used,_=_v29_monotonic_article_match(base,other,cancel_event,progress_cb)
    # Paragraph-only fallback documents still use the prior matcher.
    if not any(u.kind=='article' for u in base) or not any(u.kind=='article' for u in other):
        return match_units(base,other,max(threshold,0.48))
    return mapping,used


def _v29_addition_slot(other_index: int, reverse_map: Dict[int,int], base_len: int) -> int:
    """Place a true comparison-only addition between its mapped neighbours without reordering base."""
    prevs=[(oj,bi) for oj,bi in reverse_map.items() if oj<other_index]
    nexts=[(oj,bi) for oj,bi in reverse_map.items() if oj>other_index]
    p=max(prevs,key=lambda x:x[0])[1] if prevs else None
    q=min(nexts,key=lambda x:x[0])[1] if nexts else None
    if p is not None and q is not None and p<q:
        return q
    if p is not None:
        return min(base_len,p+1)
    if q is not None:
        return max(0,q)
    return base_len


def build_groups(docs_units: List[List[Unit]], base_index: int = 0, cancel_event=None, progress_cb=None) -> List[Dict[str,Any]]:
    base=docs_units[base_index]; others=[i for i in range(len(docs_units)) if i!=base_index]
    maps={}; used_sets={}; mode_maps={}
    for pos,di in enumerate(others):
        def sub(x,pos=pos):
            if progress_cb: progress_cb((pos+x)/max(1,len(others)))
        mp,used,modes=_v29_monotonic_article_match(base,docs_units[di],cancel_event,sub)
        maps[di]=mp; used_sets[di]=used; mode_maps[di]=modes

    base_groups=[]
    for bi,a in enumerate(base):
        members=[None]*len(docs_units); scores=[None]*len(docs_units); modes=[None]*len(docs_units)
        members[base_index]=a; scores[base_index]=1.0; modes[base_index]='base'
        for di in others:
            if bi in maps[di]:
                j,s=maps[di][bi]; members[di]=docs_units[di][j]; scores[di]=s; modes[di]=mode_maps[di].get(bi)
        base_groups.append({'members':members,'scores':scores,'modes':modes,'axis_extra':False})

    # Build addition groups, pairing B/C additions when appropriate.
    additions=[]
    if len(docs_units)==3:
        d1,d2=others
        idx1=[j for j,u in enumerate(docs_units[d1]) if u.kind=='article' and j not in used_sets[d1]]
        idx2=[j for j,u in enumerate(docs_units[d2]) if u.kind=='article' and j not in used_sets[d2]]
        u1=[docs_units[d1][j] for j in idx1]; u2=[docs_units[d2][j] for j in idx2]
        local,used2,_=_v29_monotonic_article_match(u1,u2,cancel_event,None) if u1 and u2 else ({},set(),{})
        used1=set(local)
        for i1,(i2,s) in local.items():
            members=[None]*3; scores=[None]*3; modes=[None]*3
            members[d1]=u1[i1]; members[d2]=u2[i2]; scores[d1]=1.0; scores[d2]=s
            additions.append((members,scores,modes,[(d1,idx1[i1]),(d2,idx2[i2])]))
        for k,u in enumerate(u1):
            if k not in used1:
                members=[None]*3; scores=[None]*3; modes=[None]*3; members[d1]=u; scores[d1]=1.0
                additions.append((members,scores,modes,[(d1,idx1[k])]))
        for k,u in enumerate(u2):
            if k not in used2:
                members=[None]*3; scores=[None]*3; modes=[None]*3; members[d2]=u; scores[d2]=1.0
                additions.append((members,scores,modes,[(d2,idx2[k])]))
    else:
        di=others[0]
        for j,u in enumerate(docs_units[di]):
            if u.kind=='article' and j not in used_sets[di]:
                members=[None]*2; scores=[None]*2; modes=[None]*2; members[di]=u; scores[di]=1.0
                additions.append((members,scores,modes,[(di,j)]))

    reverse={di:{oj:bi for bi,(oj,_) in maps[di].items()} for di in others}
    buckets=defaultdict(list)
    for members,scores,modes,locs in additions:
        slots=[_v29_addition_slot(oj,reverse[di],len(base)) for di,oj in locs]
        slot=round(sum(slots)/len(slots)) if slots else len(base)
        buckets[max(0,min(len(base),slot))].append({'members':members,'scores':scores,'modes':modes,'axis_extra':True})
    # Natural order within an insertion slot.
    for slot in buckets:
        buckets[slot].sort(key=lambda g:min((m.index for i,m in enumerate(g['members']) if i!=base_index and m),default=10**9))

    out=[]
    for bi in range(len(base_groups)):
        out.extend(buckets.get(bi,[])); out.append(base_groups[bi])
    out.extend(buckets.get(len(base_groups),[]))
    return out


def _v29_structural_label(label: str) -> bool:
    return bool(label and not label.startswith('문장'))


def _v29_subpart_core(label: str, text: str) -> str:
    """Compare the legal content without treating the enumerator itself as substantive text."""
    t=(text or '').strip()
    l=(label or '').strip()
    if l and not l.startswith('문장'):
        # Exact label first, then generic legal enumerators as a fallback.
        t=re.sub(r'^\s*'+re.escape(l)+r'\s*','',t,count=1)
        t=re.sub(r'^\s*(?:[①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳]|\(\d+\)|\d+[.)]|[가-하A-Za-z][.)])\s*','',t,count=1)
    return t.strip()


def _v29_subpart_similarity(a: Tuple[str,str], b: Tuple[str,str]) -> float:
    la,ta=a; lb,tb=b
    ca=_v29_subpart_core(la,ta); cb=_v29_subpart_core(lb,tb)
    na,nb=normalize_text(ca),normalize_text(cb)
    if not na or not nb: return 0.0
    seq=SequenceMatcher(None,na,nb,autojunk=False).ratio()
    wa,wb=set(token_words(ca)),set(token_words(cb))
    jac=len(wa&wb)/max(1,len(wa|wb))
    contain=len(wa&wb)/max(1,min(len(wa),len(wb)))
    s=max(0.72*seq+0.28*jac,0.86*contain)
    if la==lb and _v29_structural_label(la): s=min(1.0,s+0.08)
    return s


def _v29_subpart_changes(old: str, new: str, limit: int=14):
    """Match 항/호/문장 by content first, then report number/order changes separately."""
    ap=split_legal_subparts_v15(old); bp=split_legal_subparts_v15(new)
    used_a=set(); used_b=set(); pairs=[]
    # 1. Exact legal content, regardless of enumerator/position.
    norm_b=defaultdict(list)
    for j,(lb,tb) in enumerate(bp):
        norm_b[normalize_text(_v29_subpart_core(lb,tb))].append(j)
    for i,(la,ta) in enumerate(ap):
        n=normalize_text(_v29_subpart_core(la,ta))
        if not n: continue
        c=[j for j in norm_b.get(n,[]) if j not in used_b]
        if c:
            j=min(c,key=lambda z:abs(i-z)); pairs.append((i,j,1.0)); used_a.add(i); used_b.add(j)
    # 2. Fuzzy legal-content matches for revised subparts, allowing non-monotonic movement.
    cands=[]
    for i,a in enumerate(ap):
        if i in used_a: continue
        for j,b in enumerate(bp):
            if j in used_b: continue
            ss=_v29_subpart_similarity(a,b)
            if ss>=0.56: cands.append((ss,i,j))
    for ss,i,j in sorted(cands,reverse=True):
        if i in used_a or j in used_b: continue
        pairs.append((i,j,ss)); used_a.add(i); used_b.add(j)

    lines=[]; move_count=0; change_count=0
    for i,j,ss in sorted(pairs,key=lambda x:x[0]):
        la,ta=ap[i]; lb,tb=bp[j]
        ca=_v29_subpart_core(la,ta); cb=_v29_subpart_core(lb,tb)
        moved=(i!=j) or (la!=lb and (_v29_structural_label(la) or _v29_structural_label(lb)))
        if moved:
            move_count+=1
            if len(lines)<limit:
                kind='항/호 번호·위치 변경' if (_v29_structural_label(la) or _v29_structural_label(lb)) else '문장 위치 이동'
                lines.append(f'{kind}: [{la}] → [{lb}]')
        if normalize_text(ca)!=normalize_text(cb):
            change_count+=1
            for d in _v26_compact_ops(ca,cb,limit=4):
                if len(lines)>=limit: break
                lines.append(f'[{la}→{lb}] {d}' if la!=lb else f'[{la}] {d}')
    for i,(la,ta) in enumerate(ap):
        if i not in used_a:
            change_count+=1
            if len(lines)<limit: lines.append(f'[{la}] 삭제 “{_short_v15(_v29_subpart_core(la,ta),150)}”')
    for j,(lb,tb) in enumerate(bp):
        if j not in used_b:
            change_count+=1
            if len(lines)<limit: lines.append(f'[{lb}] 추가 “{_short_v15(_v29_subpart_core(lb,tb),150)}”')
    total=move_count+change_count
    return {'lines':lines,'count':total,'moves':move_count,'changes':change_count,'truncated':max(0,total-len(lines))}




# ---------------- V3.0: stable legal-part + punctuation diff rules ----------------
# Ordinary prose is NOT sentence-split by periods.  A punctuation edit must never change
# the comparison unit itself.  Only explicit legal enumerators ((1), 1., ①, 가.) create
# subparts.  This prevents a newly inserted period from turning one revised sentence into
# a bogus delete + add cascade.
_V30_ENUM_START_RE = re.compile(
    r'^\s*((?:[①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳]|\(\d+\)|\d+[.)]|[가-하A-Za-z][.)]))\s+',
    re.UNICODE,
)
_V30_LEX_RE = re.compile(
    r"[가-힣A-Za-z]+(?:['’][A-Za-z]+)?|\d+(?:,\d{3})*(?:\.\d+)?%?",
    re.UNICODE,
)


def _v30_split_explicit_parts(text: str) -> List[Tuple[str,str]]:
    """Split only on explicit legal enumerators, never on ordinary sentence punctuation."""
    text=(text or '').replace('\r\n','\n').replace('\r','\n').strip()
    if not text:
        return []
    # Recover a few embedded list markers that Word sometimes flattened into one line,
    # but require a strong boundary so prose references such as "Paragraph 2." are safe.
    text=re.sub(
        r'(?<=\S)[ \t]{2,}(?=(?:[①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳]|\(\d+\)|\d+[.)]|[가-하][.)])\s+)',
        '\n', text,
    )
    parts=[]; current_label='본문'; buf=[]; saw_enum=False
    for raw in text.splitlines():
        line=raw.strip()
        if not line:
            continue
        m=_V30_ENUM_START_RE.match(line)
        if m:
            if buf:
                parts.append((current_label,' '.join(buf).strip()))
            current_label=m.group(1)
            buf=[line[m.end():].strip()]
            saw_enum=True
        else:
            buf.append(line)
    if buf:
        parts.append((current_label,' '.join(buf).strip()))
    if not saw_enum:
        return [('본문',re.sub(r'\s+',' ',text).strip())]
    return [(l,t) for l,t in parts if t]


def _v30_lex_spans(text: str):
    text=text or ''
    return [(m.group(0),m.start(),m.end()) for m in _V30_LEX_RE.finditer(text)]


def _v30_punct_only(text: str, start: int, end: int, spans, si1: int, si2: int) -> str:
    if end<=start: return ''
    seg=list(text[start:end])
    for k in range(si1,si2):
        _,a,b=spans[k]
        aa=max(a,start)-start; bb=min(b,end)-start
        for x in range(max(0,aa),min(len(seg),max(0,bb))):
            seg[x]=' '
    # Ignore whitespace and invisible Word controls, but preserve exact punctuation code points.
    invis={'\u200b','\u200c','\u200d','\ufeff','\u00ad'}
    return ''.join(ch for ch in seg if (not ch.isspace()) and ch not in invis and (not ch.isalnum()) and ch!='_')


def _v30_gap_context(spans, left_idx: int, right_idx: int) -> str:
    left=spans[left_idx][0] if 0 <= left_idx < len(spans) else None
    right=spans[right_idx][0] if 0 <= right_idx < len(spans) else None
    if left and right: return f'{left} ↔ {right}'
    if right: return f'{right} 앞'
    if left: return f'{left} 뒤'
    return '문장 경계'


def _v30_cp(s: str) -> str:
    return ' '.join(f'U+{ord(ch):04X}' for ch in s)


def _v30_punct_display(s: str) -> str:
    # Do not wrap punctuation in curly quotes: doing so made a single \" look like \"\"\".
    return s.replace('\n','\\n').replace('\t','\\t')


def _v30_compact_ops(old: str, new: str, limit: int=14) -> List[str]:
    """Human-readable diff with stable lexical anchors and location-aware punctuation.

    Rules:
      * words/numbers align first;
      * punctuation is owned by the gap between aligned words;
      * same punctuation moved to another gap => one '기호 위치 변경' event;
      * different punctuation in the same gap => one '기호 변경' event (+ code points);
      * pure punctuation insert/delete stays visible;
      * Word run boundaries never create a text change.
    """
    old=old or ''; new=new or ''
    a=_v30_lex_spans(old); b=_v30_lex_spans(new)
    ak=[x[0].casefold() for x in a]; bk=[x[0].casefold() for x in b]
    sm=SequenceMatcher(None,ak,bk,autojunk=False)
    pairs=[]
    for block in sm.get_matching_blocks():
        for n in range(block.size):
            pairs.append((block.a+n,block.b+n))
    pairs.sort()
    anchors=[(-1,-1)] + pairs + [(len(a),len(b))]

    lexical=[]
    punct_same_gap=[]
    punct_del=[]; punct_add=[]
    for z in range(len(anchors)-1):
        ai,bj=anchors[z]; an,bn=anchors[z+1]
        ao1,ao2=ai+1,an; bo1,bo2=bj+1,bn
        pos=(a[ai][2] if ai>=0 and ai<len(a) else 0)

        # Exact raw lexical span between anchors, excluding surrounding punctuation.
        op=(old[a[ao1][1]:a[ao2-1][2]].strip() if ao1<ao2 else '')
        np=(new[b[bo1][1]:b[bo2-1][2]].strip() if bo1<bo2 else '')
        if op or np:
            if op and np:
                lexical.append((pos,f'변경 “{_short_v15(op,120)}” → “{_short_v15(np,120)}”'))
            elif op:
                lexical.append((pos,f'삭제 “{_short_v15(op,150)}”'))
            else:
                lexical.append((pos,f'추가 “{_short_v15(np,150)}”'))

        old_start=(a[ai][2] if ai>=0 and ai<len(a) else 0)
        old_end=(a[an][1] if an<len(a) else len(old))
        new_start=(b[bj][2] if bj>=0 and bj<len(b) else 0)
        new_end=(b[bn][1] if bn<len(b) else len(new))
        old_p=_v30_punct_only(old,old_start,old_end,a,ao1,ao2)
        new_p=_v30_punct_only(new,new_start,new_end,b,bo1,bo2)
        if old_p==new_p:
            continue
        old_ctx=_v30_gap_context(a,ai,an)
        new_ctx=_v30_gap_context(b,bj,bn)
        if old_p and new_p:
            punct_same_gap.append((pos,old_p,new_p,old_ctx,new_ctx))
        elif old_p:
            punct_del.append([pos,old_p,old_ctx,False])
        elif new_p:
            punct_add.append([pos,new_p,new_ctx,False])

    # Pair an identical punctuation sequence deleted from one gap and added at another gap.
    # This is a MOVE, not three independent add/delete records.
    punct_msgs=[]
    for d in punct_del:
        if d[3]: continue
        candidates=[x for x in punct_add if (not x[3]) and x[1]==d[1]]
        if candidates:
            aev=min(candidates,key=lambda x:abs(x[0]-d[0]))
            d[3]=True; aev[3]=True
            sym=_v30_punct_display(d[1])
            punct_msgs.append((min(d[0],aev[0]),f'기호 위치 변경: {sym} ({d[2]} → {aev[2]})'))

    for pos,op,np,oc,nc in punct_same_gap:
        # Reduce same-gap punctuation to the minimal character edit so unchanged
        # punctuation beside it (e.g. the period in ". -> ”.) is not repeated.
        psm=SequenceMatcher(None,list(op),list(np),autojunk=False)
        for tag,i1,i2,j1,j2 in psm.get_opcodes():
            if tag=='equal':
                continue
            oo=op[i1:i2]; nn=np[j1:j2]
            if tag=='replace' and oo and nn:
                punct_msgs.append((pos,f'기호 변경: {_v30_punct_display(oo)} → {_v30_punct_display(nn)} ({_v30_cp(oo)} → {_v30_cp(nn)})'))
            elif tag=='delete' and oo:
                punct_msgs.append((pos,f'기호 삭제: {_v30_punct_display(oo)} ({oc})'))
            elif tag=='insert' and nn:
                punct_msgs.append((pos,f'기호 추가: {_v30_punct_display(nn)} ({nc})'))
    for pos,p,ctx,used in punct_del:
        if not used:
            punct_msgs.append((pos,f'기호 삭제: {_v30_punct_display(p)} ({ctx})'))
    for pos,p,ctx,used in punct_add:
        if not used:
            punct_msgs.append((pos,f'기호 추가: {_v30_punct_display(p)} ({ctx})'))

    events=[(p,1,m) for p,m in punct_msgs] + [(p,2,m) for p,m in lexical]
    events.sort(key=lambda x:(x[0],x[1]))
    out=[]
    for _,_,msg in events:
        if msg not in out:
            out.append(msg)
        if len(out)>=limit: break
    return out


def _v30_structural_label(label: str) -> bool:
    return bool(label and label!='본문')


def _v30_core(label: str, text: str) -> str:
    # _v30_split_explicit_parts already strips the enumerator.
    return (text or '').strip()


def _v30_part_similarity(a: Tuple[str,str], b: Tuple[str,str]) -> float:
    la,ta=a; lb,tb=b
    na,nb=normalize_text(ta),normalize_text(tb)
    if not na or not nb: return 0.0
    seq=SequenceMatcher(None,na,nb,autojunk=False).ratio()
    wa,wb=set(token_words(ta)),set(token_words(tb))
    jac=len(wa&wb)/max(1,len(wa|wb)); contain=len(wa&wb)/max(1,min(len(wa),len(wb)))
    s=max(0.72*seq+0.28*jac,0.86*contain)
    if la==lb and _v30_structural_label(la): s=min(1.0,s+0.10)
    return s


def _v30_subpart_changes(old: str, new: str, limit: int=14):
    ap=_v30_split_explicit_parts(old); bp=_v30_split_explicit_parts(new)
    # No legal enumerators on either side: compare the entire body as one stable unit.
    if len(ap)==1 and len(bp)==1 and ap[0][0]=='본문' and bp[0][0]=='본문':
        lines=_v30_compact_ops(ap[0][1],bp[0][1],limit=limit)
        return {'lines':lines,'count':len(lines),'moves':0,'changes':len(lines),'truncated':0}

    # If only one side has legal numbering, do not invent sentence deletes/adds.  Report the
    # structural change once, then compare the whole body text.
    a_struct=any(_v30_structural_label(l) for l,_ in ap)
    b_struct=any(_v30_structural_label(l) for l,_ in bp)
    if a_struct != b_struct:
        lines=['항/호 구조 변경']
        for d in _v30_compact_ops(old,new,limit=max(0,limit-1)):
            lines.append(d)
            if len(lines)>=limit: break
        return {'lines':lines,'count':len(lines),'moves':1,'changes':max(0,len(lines)-1),'truncated':0}

    used_a=set(); used_b=set(); pairs=[]
    norm_b=defaultdict(list)
    for j,(lb,tb) in enumerate(bp):
        norm_b[normalize_text(tb)].append(j)
    for i,(la,ta) in enumerate(ap):
        n=normalize_text(ta)
        if not n: continue
        c=[j for j in norm_b.get(n,[]) if j not in used_b]
        if c:
            j=min(c,key=lambda z:abs(i-z)); pairs.append((i,j,1.0)); used_a.add(i); used_b.add(j)
    cands=[]
    for i,a0 in enumerate(ap):
        if i in used_a: continue
        for j,b0 in enumerate(bp):
            if j in used_b: continue
            ss=_v30_part_similarity(a0,b0)
            if ss>=0.54: cands.append((ss,i,j))
    for ss,i,j in sorted(cands,reverse=True):
        if i in used_a or j in used_b: continue
        pairs.append((i,j,ss)); used_a.add(i); used_b.add(j)

    lines=[]; move_count=0; change_count=0
    for i,j,ss in sorted(pairs,key=lambda x:x[0]):
        la,ta=ap[i]; lb,tb=bp[j]
        moved=(i!=j) or (la!=lb and (_v30_structural_label(la) or _v30_structural_label(lb)))
        if moved:
            move_count+=1
            if len(lines)<limit:
                lines.append(f'항/호 번호·위치 변경: [{la}] → [{lb}]')
        if normalize_text(ta)!=normalize_text(tb):
            ops=_v30_compact_ops(ta,tb,limit=4)
            if ops: change_count+=1
            for d in ops:
                if len(lines)>=limit: break
                prefix=f'[{la}→{lb}] ' if la!=lb else ('' if la=='본문' else f'[{la}] ')
                lines.append(prefix+d)
    for i,(la,ta) in enumerate(ap):
        if i not in used_a:
            change_count+=1
            if len(lines)<limit: lines.append(f'[{la}] 삭제 “{_short_v15(ta,150)}”')
    for j,(lb,tb) in enumerate(bp):
        if j not in used_b:
            change_count+=1
            if len(lines)<limit: lines.append(f'[{lb}] 추가 “{_short_v15(tb,150)}”')
    total=move_count+change_count
    return {'lines':lines,'count':total,'moves':move_count,'changes':change_count,'truncated':max(0,total-len(lines))}

def _v29_summarize_group(members: List[Optional[Unit]], scores: List[Optional[float]], modes=None, base_index:int=0):
    present=[i for i,m in enumerate(members) if m]; base=members[base_index]
    msgs=[]; tags=[]; risk='LOW'; modes=modes or [None]*len(members)
    if base is None:
        docs=', '.join(chr(65+i) for i in present); msgs.append(f'문서 {docs}: 신설 조항')
        tags=['신설']; risk='MEDIUM'
    elif len(present)==1:
        msgs.append('비교문서에서 이 조항이 삭제됨'); tags=['삭제']; risk='HIGH'
    else:
        for i,m in enumerate(members):
            if i==base_index: continue
            label=chr(65+i)
            if m is None:
                msgs.append(f'문서 {label}: 조항 삭제'); tags.append('삭제') if '삭제' not in tags else None; risk='HIGH'; continue
            mode=modes[i] if i<len(modes) else None
            if base.kind=='article' and m.kind=='article' and base.number!=m.number:
                if mode=='moved':
                    msgs.append(f'문서 {label}: 조 위치 이동/번호 변경 {base.header} → {m.header}')
                    if '이동' not in tags: tags.append('이동')
                else:
                    msgs.append(f'문서 {label}: 조 번호 변경 {base.header} → {m.header}')
                    if '재번호화' not in tags: tags.append('재번호화')
            elif mode=='moved':
                msgs.append(f'문서 {label}: 조 위치 이동 {base.header} → {m.header}')
                if '이동' not in tags: tags.append('이동')
            if base.title and m.title and base.norm_title!=m.norm_title:
                msgs.append(f'문서 {label}: 조항명 “{base.title}” → “{m.title}”')
                if '제목변경' not in tags: tags.append('제목변경')
            if base.section!=m.section and (base.section or m.section):
                msgs.append(f'문서 {label}: 편제 “{base.section or "없음"}” → “{m.section or "없음"}”')
                if '편제이동' not in tags: tags.append('편제이동')
            if base.norm_body!=m.norm_body:
                detail=_v30_subpart_changes(base.body,m.body,limit=14)
                if detail['moves'] and '항이동' not in tags: tags.append('항이동')
                if detail['changes'] and '내용변경' not in tags: tags.append('내용변경')
                for x in detail['lines']: msgs.append(f'문서 {label}: {x}')
                if detail['truncated']: msgs.append(f'문서 {label}: 그 외 변경 {detail["truncated"]}건')
                risk='MEDIUM'
    if len(members)==3 and base is not None:
        others=[i for i in range(3) if i!=base_index]; m1,m2=members[others[0]],members[others[1]]
        if m1 and m2 and m1.norm_body!=m2.norm_body:
            msgs.append(f'문서 {chr(65+others[0])}와 {chr(65+others[1])}: 서로 다른 수정 내용')
            if '비교본차이' not in tags: tags.append('비교본차이')
    if not msgs: msgs=['변경 없음']; tags=['동일']
    return {'messages':msgs,'tags':tags,'risk':risk,'confidence':None}


def compare_documents(names: List[str], texts: List[str], base_index: int=0, progress_cb=None, cancel_event=None) -> Dict[str,Any]:
    if progress_cb: progress_cb(3,'DOCX/TXT 문서 구조를 분석하는 중...')
    units=[]; parse_info=[]
    for i,t in enumerate(texts):
        _check_cancel(cancel_event); parsed=parse_units(t); units.append(parsed)
        ac=sum(1 for u in parsed if u.kind=='article'); parse_info.append({'articles':ac,'units':len(parsed),'preamble':0})
        if progress_cb: progress_cb(5+int(15*(i+1)/len(texts)),f'문서 {chr(65+i)} · 구조 블록 {len(parsed)}개')
    if progress_cb: progress_cb(22,'문서 블록의 내용·제목·위치를 기준으로 대응시키는 중...')
    def mp(x):
        if progress_cb: progress_cb(22+int(33*x),'신설·삭제·이동·재배치를 분석하는 중...')
    groups=build_groups(units,base_index=base_index,cancel_event=cancel_event,progress_cb=mp)
    rows=[]; counts={'total':0,'changed':0,'added':0,'deleted':0,'moved':0,'high':0}; total=max(1,len(groups))
    for ridx,g in enumerate(groups):
        _check_cancel(cancel_event); members=g['members']; scores=g['scores']; modes=g.get('modes')
        body_segs=directional_segments_v19(members,base_index,'body'); header_segs=directional_segments_v19(members,base_index,'header')
        full_segs=[]; htmls=[]; raw=[]
        for i,m in enumerate(members):
            if m:
                hs=header_segs[i] or [(m.header,'normal')]; fs=hs+[('\n','normal')]+body_segs[i]
                full_segs.append(fs); htmls.append(render_segments_html_v19(fs)); raw.append(m.text)
            else:
                full_segs.append([]); htmls.append('<span class="missing">[해당 조항 없음]</span>'); raw.append('')
        summary=_v29_summarize_group(members,scores,modes,base_index)
        changed=summary['tags']!=['동일']; counts['total']+=1
        if changed: counts['changed']+=1
        if '신설' in summary['tags']: counts['added']+=1
        if '삭제' in summary['tags']: counts['deleted']+=1
        if '이동' in summary['tags'] or '재번호화' in summary['tags']: counts['moved']+=1
        rows.append({'id':ridx,'raw':raw,'html':htmls,'segments':full_segs,'body_segments':body_segs,'header_segments':header_segs,
                     'members':[asdict(m) if m else None for m in members],'summary':summary,'changed':changed,'axis_extra':bool(g.get('axis_extra'))})
        if progress_cb and ridx % max(1,total//25)==0:
            progress_cb(58+int(40*(ridx+1)/total),f'세부 문구와 목록 위치까지 분석 중... {ridx+1}/{total}')
    if progress_cb: progress_cb(100,'문서 구조 비교 완료')
    return {'names':names,'rows':rows,'counts':counts,'unit_counts':[len(x) for x in units],
            'parse_info':parse_info,'base_index':base_index}


# ---------------- V3.1: numbered change-marker overlay layer ----------------
# Every visible red deletion / blue insertion gets a numbered badge positioned OVER the
# exact changed span.  The change list uses the same number and is clickable both ways.
APP_VERSION = "3.2"

_V31_PUNCT_ONLY_RE = re.compile(r'^\W+$', re.UNICODE)


def _v31_circled(n: int) -> str:
    circled = '①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳'
    return circled[n-1] if 1 <= n <= 20 else f'[{n}]'


def _v31_style_spans(tokens: List[str], marks: List[str], wanted: str):
    """Return exact character spans for contiguous styled tokens."""
    out=[]; pos=0; start=None; buf=[]
    def flush(end_pos):
        nonlocal start,buf
        if start is None: return
        raw=''.join(buf)
        l=len(raw)-len(raw.lstrip())
        r=len(raw)-len(raw.rstrip())
        st=start+l; en=end_pos-r
        txt=raw.strip()
        if txt and en>st:
            out.append((st,en,txt))
        start=None; buf=[]
    for tok,sty in zip(tokens,marks):
        if sty==wanted:
            if start is None: start=pos
            buf.append(tok)
        else:
            flush(pos)
        pos += len(tok)
    flush(pos)
    return out


def _v31_context(text: str, start: int, end: int) -> str:
    spans=_v30_lex_spans(text or '')
    left=None; right=None
    for word,a,b in spans:
        if b <= start: left=word
        elif a >= end:
            right=word; break
    if left and right: return f'{left} ↔ {right}'
    if right: return f'{right} 앞'
    if left: return f'{left} 뒤'
    return '문장 경계'


def _v31_event_message(action: str, txt: str, full_text: str, start: int, end: int) -> str:
    # Punctuation stays visible and gets its exact lexical location; words/phrases stay concise.
    punct=bool(txt) and all((not ch.isalnum()) and (not ch.isspace()) and ch!='_' for ch in txt)
    if punct:
        return f'기호 {action}: {_v30_punct_display(txt)} ({_v31_context(full_text,start,end)})'
    label='삭제' if action=='삭제' else '추가'
    return f'{label}: “{re.sub(r"\s+", " ", txt or "").strip()}”'


def _v31_split_changed_span(full_text: str, st: int, en: int):
    """Split a changed span into leading punctuation / lexical payload / trailing punctuation.

    This keeps a newly inserted period separate from the newly inserted phrase that follows it,
    while preserving exact character offsets for overlay badges.
    """
    raw=(full_text or '')[st:en]
    if not raw: return []
    # Trim only outer whitespace first.
    ltrim=len(raw)-len(raw.lstrip()); rtrim=len(raw)-len(raw.rstrip())
    st2=st+ltrim; en2=en-rtrim; core=(full_text or '')[st2:en2]
    if not core: return []
    word_positions=[i for i,ch in enumerate(core) if ch.isalnum() or ch=='_']
    if not word_positions:
        return [(st2,en2,core,'punct')]
    first=word_positions[0]; last=word_positions[-1]+1
    out=[]
    lead=core[:first]
    # Do not make whitespace a change event; preserve only punctuation in the edge piece.
    lead_clean=''.join(ch for ch in lead if not ch.isspace())
    if lead_clean:
        # Usually punctuation is contiguous at the edge. Find the first/last non-space chars exactly.
        a=0
        while a<len(lead) and lead[a].isspace(): a+=1
        b=len(lead)
        while b>a and lead[b-1].isspace(): b-=1
        out.append((st2+a,st2+b,lead[a:b],'punct'))
    lex=core[first:last]
    if lex.strip(): out.append((st2+first,st2+last,lex.strip(),'lex'))
    trail=core[last:]
    trail_clean=''.join(ch for ch in trail if not ch.isspace())
    if trail_clean:
        a=0
        while a<len(trail) and trail[a].isspace(): a+=1
        b=len(trail)
        while b>a and trail[b-1].isspace(): b-=1
        out.append((st2+last+a,st2+last+b,trail[a:b],'punct'))
    return out


def _v31_pair_marker_events(base: Unit, other: Unit, base_index: int, other_index: int):
    events=[]
    label=f'문서 {chr(65+other_index)}'
    for attr,part_order in (('header',0),('body',1)):
        old=getattr(base,attr) or ''; new=getattr(other,attr) or ''
        bt,bm,ot,om=_directional_pair_marks_v19(old,new)
        old_off=0 if attr=='header' else len(base.header)+1
        new_off=0 if attr=='header' else len(other.header)+1
        for gst,gen,_ in _v31_style_spans(bt,bm,'delete'):
            for st,en,txt,kind in _v31_split_changed_span(old,gst,gen):
                ratio=st/max(1,len(old))
                events.append({
                    'relative_doc':other_index,'target_doc':base_index,'action':'삭제','text':txt,
                    'char_start':old_off+st,'char_end':old_off+en,'part':attr,
                    # Coarse positional bucket groups a deletion with its nearby insertion;
                    # within that local change, deletion is numbered before insertion.
                    'sort_key':(other_index,part_order,int(ratio*5),0,ratio),
                    'message':_v31_event_message('삭제',txt,old,st,en),'label':label,
                })
        for gst,gen,_ in _v31_style_spans(ot,om,'insert'):
            for st,en,txt,kind in _v31_split_changed_span(new,gst,gen):
                ratio=st/max(1,len(new))
                events.append({
                    'relative_doc':other_index,'target_doc':other_index,'action':'추가','text':txt,
                    'char_start':new_off+st,'char_end':new_off+en,'part':attr,
                    'sort_key':(other_index,part_order,int(ratio*5),1,ratio),
                    'message':_v31_event_message('추가',txt,new,st,en),'label':label,
                })
    # If the same punctuation disappeared at one lexical gap and appeared at another,
    # keep the two visible operations adjacent and number DELETE before ADD.  They remain
    # two events because the user wants a badge at each physical endpoint.
    pdel=[e for e in events if e['action']=='삭제' and e['text'] and all((not ch.isalnum()) and (not ch.isspace()) for ch in e['text'])]
    padd=[e for e in events if e['action']=='추가' and e['text'] and all((not ch.isalnum()) and (not ch.isspace()) for ch in e['text'])]
    used=set()
    for d in pdel:
        cand=[a for a in padd if id(a) not in used and a['text']==d['text'] and a['part']==d['part']]
        if not cand: continue
        a=min(cand,key=lambda x:abs(x['sort_key'][-1]-d['sort_key'][-1]))
        used.add(id(a))
        logical=min(d['sort_key'][-1],a['sort_key'][-1])
        common=(other_index, 0 if d['part']=='header' else 1, int(logical*5))
        d['sort_key']=common+(0,logical)
        a['sort_key']=common+(1,logical+1e-6)
    return events


def _v31_structural_messages(messages: List[str]) -> List[str]:
    """Keep non-textual lineage/structure notes; visible text edits are represented by markers."""
    keys=(
        '조 위치 이동','조 번호 변경','편제 ','항/호 번호·위치 변경','항/호 구조 변경',
        '신설 조항','조항 삭제','비교문서에서 이 조항이 삭제','서로 다른 수정 내용',
        '기준문서 외 신설','조 위치 이동/번호 변경'
    )
    out=[]
    for m in messages or []:
        if m=='변경 없음':
            continue
        if any(k in m for k in keys):
            if m not in out: out.append(m)
    return out


_compare_documents_v30 = compare_documents

def compare_documents(names: List[str], texts: List[str], base_index: int=0, progress_cb=None, cancel_event=None) -> Dict[str,Any]:
    result=_compare_documents_v30(names,texts,base_index=base_index,progress_cb=progress_cb,cancel_event=cancel_event)
    for row in result.get('rows',[]):
        members_raw=row.get('members') or []
        members=[]
        for m in members_raw:
            members.append(Unit(**m) if m else None)
        base=members[base_index] if base_index < len(members) else None
        events=[]
        if base:
            for i,m in enumerate(members):
                if i==base_index or not m: continue
                events.extend(_v31_pair_marker_events(base,m,base_index,i))
        # Stable reading order: B then C; within a document, header then body, top to bottom;
        # deletion at a location before insertion at that location.
        events.sort(key=lambda e:e['sort_key'])
        for n,e in enumerate(events,1):
            e['num']=n
            e.pop('sort_key',None)
        row['markers']=events
        structural=_v31_structural_messages((row.get('summary') or {}).get('messages') or [])
        display=[f"{_v31_circled(e['num'])} {e['label']} · {e['message']}" for e in events]
        display.extend('• '+m for m in structural)
        if not display:
            display=['변경 없음']
        row['display_messages']=display
        # Excel/search should use the same human-visible change list rather than a second diff engine.
        row['summary']['messages']=display
    return result


def _v31_place_badge(self, text_widget, row_id: int, event: Dict[str,Any]):
    try:
        idx=f"1.0+{int(event['char_start'])}c"
        box=text_widget.bbox(idx)
        if not box:
            return
        x,y,w,h=box
        num=event['num']
        # A real overlay: Label is placed over the Text widget and does not consume a text character.
        badge=self.tk.Label(
            text_widget,text=str(num),font=('Malgun Gothic',7,'bold'),
            fg='white',bg='#245ea8',bd=0,padx=3,pady=0,cursor='hand2'
        )
        badge.place(x=max(0,x-4),y=max(0,y-7),anchor='nw')
        badge.lift()
        badge.bind('<Button-1>',lambda _e,r=row_id,n=num:self._v31_focus_summary(r,n))
        badge.bind('<Enter>',lambda _e,r=row_id,n=num:self._v31_focus_summary(r,n,flash_only=True))
        self._v31_badges.append(badge)
    except Exception:
        logging.exception('marker overlay placement failed')


def _v31_focus_marker(self, row_id: int, num: int):
    ev=self._v31_event_map.get((row_id,num))
    if not ev: return
    t=self._v31_doc_widgets.get((row_id,ev['target_doc']))
    if not t: return
    try:
        s=f"1.0+{int(ev['char_start'])}c"; e=f"1.0+{int(ev['char_end'])}c"
        t.see(s)
        t.configure(state='normal')
        t.tag_configure('marker_focus',background='#fff59d')
        t.tag_add('marker_focus',s,e)
        t.tag_raise('marker_focus')
        t.configure(state='disabled')
        def clear():
            try:
                t.configure(state='normal'); t.tag_remove('marker_focus','1.0','end'); t.configure(state='disabled')
            except Exception: pass
        t.after(1100,clear)
    except Exception:
        logging.exception('focus marker failed')


def _v31_focus_summary(self, row_id: int, num: int, flash_only: bool=False):
    t=self._v31_summary_widgets.get(row_id)
    if not t: return
    tag=f'marker_line_{num}'
    try:
        ranges=t.tag_ranges(tag)
        if ranges:
            t.see(ranges[0])
            t.configure(state='normal')
            t.tag_configure('summary_focus',background='#fff2a8')
            t.tag_add('summary_focus',ranges[0],ranges[-1])
            t.tag_raise('summary_focus')
            t.configure(state='disabled')
            def clear():
                try:
                    t.configure(state='normal'); t.tag_remove('summary_focus','1.0','end'); t.configure(state='disabled')
                except Exception: pass
            t.after(900,clear)
        if not flash_only:
            self._v31_focus_marker(row_id,num)
    except Exception:
        logging.exception('focus summary failed')


def _v31_make_doc_cell(self,parent,row,i,h):
    t=self._new_text(parent,h); m=row['members'][i]
    if not m:
        t.insert('end','[해당 조항 없음]','missing')
    else:
        hs=row.get('header_segments',[[] for _ in row['segments']])[i] if row.get('header_segments') else []
        if hs:
            for txt,sty in hs: t.insert('end',txt,sty if sty in ('delete','insert') else 'article_header')
            t.insert('end','\n')
        else:
            t.insert('end',m.get('header','')+'\n','article_header')
        for txt,sty in row.get('body_segments',row['segments'])[i]:
            t.insert('end',txt,sty if sty in ('delete','insert') else ())
    t.configure(state='disabled')
    self._v31_doc_widgets[(row['id'],i)]=t
    evs=[e for e in row.get('markers',[]) if e.get('target_doc')==i]
    if evs:
        def place_all(widget=t, rid=row['id'], events=tuple(evs)):
            for ev in events:
                _v31_place_badge(self,widget,rid,ev)
        t.after_idle(place_all)
    return t


def _v31_make_summary_cell(self,parent,row,h):
    t=self._new_text(parent,h,'#ffffff')
    self._v31_summary_widgets[row['id']]=t
    events={e['num']:e for e in row.get('markers',[])}
    structural=_v31_structural_messages((row.get('summary') or {}).get('messages') or [])
    if not events and not structural:
        t.insert('end','변경 없음','conf')
    else:
        for num in sorted(events):
            ev=events[num]
            start=t.index('end-1c')
            prefix=f"{_v31_circled(num)} {ev['label']} · "
            t.insert('end',prefix,'marker_num')
            t.insert('end',ev['message']+'\n')
            end=t.index('end-1c')
            tag=f'marker_line_{num}'
            t.tag_add(tag,start,end)
            t.tag_bind(tag,'<Button-1>',lambda _e,r=row['id'],n=num:self._v31_focus_marker(r,n))
            t.tag_bind(tag,'<Enter>',lambda _e: t.configure(cursor='hand2'))
            t.tag_bind(tag,'<Leave>',lambda _e: t.configure(cursor='arrow'))
        for msg in structural:
            t.insert('end','• '+msg+'\n')
    t.tag_configure('marker_num',foreground='#245ea8',font=('Malgun Gothic',10,'bold'))
    t.configure(state='disabled')
    return t


# Patch the existing GUI class so the rest of the mature synchronized-scroll layout is preserved.
NativeGui._v31_focus_marker = _v31_focus_marker
NativeGui._v31_focus_summary = _v31_focus_summary
NativeGui._make_doc_cell = _v31_make_doc_cell
NativeGui._make_summary_cell = _v31_make_summary_cell

# Initialize per-render marker registries without changing the established layout code.
_render_rows_v30 = NativeGui.render_rows

def _v31_render_rows(self):
    self._v31_doc_widgets={}
    self._v31_summary_widgets={}
    self._v31_event_map={}
    self._v31_badges=[]
    if self.result:
        for row in self.result.get('rows',[]):
            for ev in row.get('markers',[]):
                self._v31_event_map[(row['id'],ev['num'])]=ev
    return _render_rows_v30(self)

NativeGui.render_rows = _v31_render_rows



def _word_compare_source(path: Path, temp_dir: Path) -> Path:
    """Return a DOCX path suitable for Word CompareDocuments.

    DOCX inputs are used as-is. TXT inputs are converted to a temporary DOCX so
    the GUI's existing TXT support can also produce a Word redline.
    """
    path=Path(path)
    if path.suffix.lower()=='.docx':
        return path.resolve()
    if path.suffix.lower()!='.txt':
        raise ValueError(f'Word 변경추적에서 지원하지 않는 파일 형식입니다: {path.suffix}')
    text=read_document(path.name,path.read_bytes())
    doc=Document()
    # Preserve explicit line/paragraph boundaries rather than collapsing the file.
    lines=text.splitlines()
    if not lines:
        doc.add_paragraph('')
    else:
        for line in lines:
            doc.add_paragraph(line)
    out=temp_dir/(path.stem+'_converted.docx')
    doc.save(out)
    return out.resolve()


def create_word_tracked_compare(original_path: Path, revised_path: Path, output_path: Path,
                                revised_author: str='Revised Document', word_factory=None):
    """Create a native Word CompareDocuments redline.

    On Windows this uses Microsoft Word's COM automation API. The resulting DOCX
    contains real tracked revisions (<w:ins>/<w:del> and, when Word recognizes
    them, move revisions) and preserves Word's own compare semantics/formatting.

    ``word_factory`` is injectable for automated tests; normal callers leave it None.
    """
    import tempfile
    original_path=Path(original_path)
    revised_path=Path(revised_path)
    output_path=Path(output_path)
    output_path.parent.mkdir(parents=True,exist_ok=True)
    if output_path.exists():
        output_path.unlink()

    if os.name!='nt' and word_factory is None:
        raise RuntimeError('Word 변경추적 내보내기는 Windows의 Microsoft Word 데스크톱 앱에서 실행해야 합니다.')

    pythoncom=None
    if word_factory is None:
        try:
            import pythoncom as _pythoncom
            import win32com.client
        except ImportError as exc:
            raise RuntimeError('pywin32가 설치되지 않았습니다. START_HERE.cmd를 다시 실행해 설치를 완료하세요.') from exc
        pythoncom=_pythoncom
        pythoncom.CoInitialize()
        word_factory=lambda: win32com.client.DispatchEx('Word.Application')

    word=None; original_doc=None; revised_doc=None; compared_doc=None
    try:
        with tempfile.TemporaryDirectory(prefix='legal_compare_word_') as td:
            td=Path(td)
            original_docx=_word_compare_source(original_path,td)
            revised_docx=_word_compare_source(revised_path,td)
            word=word_factory()
            try: word.Visible=False
            except Exception: pass
            try: word.DisplayAlerts=0
            except Exception: pass

            # Open read-only so the user's originals can never be modified.
            original_doc=word.Documents.Open(str(original_docx),False,True,False)
            revised_doc=word.Documents.Open(str(revised_docx),False,True,False)

            # Word constants: wdCompareDestinationNew=2, wdGranularityWordLevel=1,
            # wdFormatXMLDocument=12.  Positional args avoid dependency on generated
            # COM type libraries and work with late-bound DispatchEx as well.
            compared_doc=word.CompareDocuments(
                original_doc, revised_doc,
                2, 1,
                True,   # formatting
                True,   # case changes
                True,   # whitespace
                True,   # tables
                True,   # headers/footers
                True,   # footnotes/endnotes
                True,   # text boxes
                True,   # fields
                True,   # comments
                True,   # moves
                str(revised_author or revised_path.stem),
                True    # ignore comparison warnings
            )
            try: compared_doc.TrackRevisions=True
            except Exception: pass
            try: compared_doc.ShowRevisions=True
            except Exception: pass
            compared_doc.SaveAs2(str(output_path.resolve()),12)
    finally:
        for doc in (compared_doc,revised_doc,original_doc):
            if doc is not None:
                try: doc.Close(False)
                except Exception: pass
        if word is not None:
            try: word.Quit()
            except Exception: pass
        if pythoncom is not None:
            try: pythoncom.CoUninitialize()
            except Exception: pass
    if not output_path.exists() or output_path.stat().st_size==0:
        raise RuntimeError('Word가 비교 문서를 저장하지 못했습니다.')
    return output_path


# ---------------- V3.3: Word tracked-change export for any selected pair ----------------
APP_VERSION = "3.3"


def _v33_begin_word_export(self, original_index: int, revised_index: int):
    """Start native Word CompareDocuments export for one ordered document pair."""
    paths=list(getattr(self,'_last_compare_paths',[]) or [])
    names=list((self.result or {}).get('names',[]) or [])
    n=len(paths)
    if not self.result or n not in (2,3) or len(names)!=n:
        self.messagebox.showwarning('Word 변경추적','비교에 사용한 원본 파일 정보를 찾을 수 없습니다. 다시 비교한 뒤 시도하세요.')
        return
    if original_index==revised_index or not (0<=original_index<n and 0<=revised_index<n):
        self.messagebox.showwarning('Word 변경추적','변경 전 문서와 변경 후 문서는 서로 달라야 합니다.')
        return

    original_path=Path(paths[original_index])
    revised_path=Path(paths[revised_index])
    initial=f'{original_path.stem}_vs_{revised_path.stem}_변경추적.docx'
    out=self.filedialog.asksaveasfilename(
        title='Word 변경 내용 추적 문서 저장',
        defaultextension='.docx',
        initialfile=initial,
        filetypes=[('Word 문서','*.docx')]
    )
    if not out:
        return
    self.word_export_btn.configure(state='disabled')
    self.status_var.set(
        f'Word 변경추적 생성 중 · 변경 전: {original_path.name} → 변경 후: {revised_path.name}'
    )
    self.progress_var.set(15)
    threading.Thread(
        target=self._word_export_worker,
        args=(original_path,revised_path,Path(out),revised_path.stem),
        daemon=True
    ).start()


def _v33_choose_word_pair(self):
    """For a three-document comparison, choose the ordered Original/Revised pair.

    The user can choose 1→2, 2→3, 1→3 or any reverse direction. This is
    deliberately independent from the GUI's three-way base axis: Word's native
    CompareDocuments always compares exactly two documents.
    """
    if not self.result:
        return
    names=list(self.result.get('names',[]) or [])
    paths=list(getattr(self,'_last_compare_paths',[]) or [])
    if len(names)!=3 or len(paths)!=3:
        self.messagebox.showwarning('Word 변경추적','3개 문서 비교 결과를 다시 만든 뒤 시도하세요.')
        return

    tk,ttk=self.tk,self.ttk
    dlg=tk.Toplevel(self.root)
    dlg.title('Word 변경추적 비교 문서 선택')
    dlg.transient(self.root)
    dlg.grab_set()
    dlg.resizable(False,False)

    frm=ttk.Frame(dlg,padding=(18,16)); frm.pack(fill='both',expand=True)
    ttk.Label(frm,text='변경 내용 추적을 만들 두 문서를 선택하세요.',font=('Malgun Gothic',11,'bold')).grid(row=0,column=0,columnspan=3,sticky='w',pady=(0,10))
    ttk.Label(frm,text='앞 문서는 변경 전(원본), 뒤 문서는 변경 후(수정본)으로 Word에 전달됩니다.',foreground='#5f6b76').grid(row=1,column=0,columnspan=3,sticky='w',pady=(0,12))

    labels=[f'{i+1}열 · {names[i]}' for i in range(3)]
    original_var=tk.StringVar()
    revised_var=tk.StringVar()
    base=int(self.result.get('base_index',0) or 0)
    other=next(i for i in range(3) if i!=base)
    original_var.set(labels[base]); revised_var.set(labels[other])

    ttk.Label(frm,text='변경 전 (원본)').grid(row=2,column=0,sticky='w',pady=4)
    original_box=ttk.Combobox(frm,textvariable=original_var,values=labels,state='readonly',width=54)
    original_box.grid(row=2,column=1,columnspan=2,sticky='ew',pady=4)
    ttk.Label(frm,text='변경 후 (수정본)').grid(row=3,column=0,sticky='w',pady=4)
    revised_box=ttk.Combobox(frm,textvariable=revised_var,values=labels,state='readonly',width=54)
    revised_box.grid(row=3,column=1,columnspan=2,sticky='ew',pady=4)

    presets=ttk.Frame(frm); presets.grid(row=4,column=0,columnspan=3,sticky='w',pady=(12,8))
    ttk.Label(presets,text='빠른 선택:').pack(side='left',padx=(0,6))
    def preset(a,b):
        original_var.set(labels[a]); revised_var.set(labels[b])
    ttk.Button(presets,text='1열 → 2열',command=lambda:preset(0,1)).pack(side='left',padx=3)
    ttk.Button(presets,text='2열 → 3열',command=lambda:preset(1,2)).pack(side='left',padx=3)
    ttk.Button(presets,text='1열 → 3열',command=lambda:preset(0,2)).pack(side='left',padx=3)
    ttk.Button(presets,text='방향 뒤집기',command=lambda: (original_var.set(revised_var.get()), revised_var.set(original_var.get()))).pack(side='left',padx=(10,3))

    # The direction-swap button needs to preserve both values before assignment.
    def swap_direction():
        a=original_var.get(); b=revised_var.get(); original_var.set(b); revised_var.set(a)
    for child in presets.winfo_children():
        try:
            if child.cget('text')=='방향 뒤집기': child.configure(command=swap_direction)
        except Exception:
            pass

    btns=ttk.Frame(frm); btns.grid(row=5,column=0,columnspan=3,sticky='e',pady=(12,0))
    def create_selected():
        try:
            oi=labels.index(original_var.get()); ri=labels.index(revised_var.get())
        except ValueError:
            self.messagebox.showwarning('Word 변경추적','비교할 문서를 다시 선택하세요.',parent=dlg); return
        if oi==ri:
            self.messagebox.showwarning('Word 변경추적','변경 전 문서와 변경 후 문서는 서로 달라야 합니다.',parent=dlg); return
        dlg.grab_release(); dlg.destroy(); _v33_begin_word_export(self,oi,ri)
    ttk.Button(btns,text='취소',command=dlg.destroy).pack(side='right',padx=(6,0))
    ttk.Button(btns,text='변경추적 문서 만들기',style='Primary.TButton',command=create_selected).pack(side='right')

    dlg.update_idletasks()
    try:
        x=self.root.winfo_rootx()+(self.root.winfo_width()-dlg.winfo_width())//2
        y=self.root.winfo_rooty()+(self.root.winfo_height()-dlg.winfo_height())//2
        dlg.geometry(f'+{max(0,x)}+{max(0,y)}')
    except Exception:
        pass
    original_box.focus_set()


def _v33_export_word_tracked(self):
    """2 docs: export immediately using selected base. 3 docs: choose any pair."""
    if not self.result:
        return
    n=len(self.result.get('names',[]) or [])
    paths=list(getattr(self,'_last_compare_paths',[]) or [])
    if n not in (2,3) or len(paths)!=n:
        self.messagebox.showwarning('Word 변경추적','TXT/DOCX 2개 또는 3개를 비교한 뒤 사용할 수 있습니다.')
        return
    if n==2:
        bi=int(self.result.get('base_index',0) or 0)
        ri=1-bi
        _v33_begin_word_export(self,bi,ri)
    else:
        _v33_choose_word_pair(self)


# Patch completion/error handlers so Word export remains available for 2 or 3 docs.
_v32_compare_done = NativeGui._compare_done

def _v33_compare_done(self,result):
    _v32_compare_done(self,result)
    n=len(result.get('names',[]) or [])
    self.word_export_btn.configure(state=('normal' if n in (2,3) else 'disabled'))

_v32_word_export_done = NativeGui._word_export_done

def _v33_word_export_done(self,out_path):
    _v32_word_export_done(self,out_path)
    if self.result and len(self.result.get('names',[]) or []) in (2,3):
        self.word_export_btn.configure(state='normal')

_v32_word_export_failed = NativeGui._word_export_failed

def _v33_word_export_failed(self,msg):
    _v32_word_export_failed(self,msg)
    if self.result and len(self.result.get('names',[]) or []) in (2,3):
        self.word_export_btn.configure(state='normal')

NativeGui.export_word_tracked = _v33_export_word_tracked
NativeGui._compare_done = _v33_compare_done
NativeGui._word_export_done = _v33_word_export_done
NativeGui._word_export_failed = _v33_word_export_failed


# ---------------- V3.4: phrase-level change hunks ----------------
# V3.3 numbered every small SequenceMatcher island.  In heavily rewritten clauses,
# common glue words (the/of/and/...) split one logical edit into dozens of markers.
# V3.4 keeps long/meaningful equal runs as anchors and merges weak equal bridges into
# one phrase-level hunk.  The overlay and the visible red/blue styling share this rule.
APP_VERSION = "3.4"

_V34_STOPWORDS = {
    'a','an','the','and','or','of','to','in','on','at','by','for','from','with','as','is','are','was','were','be','been','being',
    'this','that','these','those','it','its','their','there','here','such','any','all','other','between','within','through','under',
    '및','또는','그','이','저','의','에','에서','으로','로','를','을','은','는','가','이'
}


def _v34_lexical_token(tok: str) -> bool:
    return bool(tok and not tok.isspace() and any(ch.isalnum() for ch in tok))


def _v34_strong_equal(tokens: List[str], i1: int, i2: int) -> bool:
    words=[t for t in tokens[i1:i2] if _v34_lexical_token(t)]
    if not words:
        return False
    content=[w for w in words if w.casefold() not in _V34_STOPWORDS]
    chars=sum(len(w) for w in words)
    content_chars=sum(len(w) for w in content)
    # Four words are enough to be a stable anchor even if some are stop words.
    if len(words) >= 4:
        return True
    # Two meaningful words with reasonable substance also form an anchor.
    if len(content) >= 2 and content_chars >= 13:
        return True
    # One distinctive long term plus another lexical token is useful in legal clauses.
    if len(words) >= 2 and any(len(w) >= 11 for w in content):
        return True
    return False


def _v34_change_hunks(old: str, new: str):
    """Return phrase-level changed token ranges (i1,i2,j1,j2).

    Adjacent non-equal opcodes are merged across weak equal bridges.  Strong equal
    runs remain anchors.  This avoids 30+ markers when a sentence is substantially
    rewritten but still shares many short words.
    """
    at=display_tokens(old or ''); bt=display_tokens(new or '')
    ak=[token_key(x) for x in at]; bk=[token_key(x) for x in bt]
    ops=SequenceMatcher(None,ak,bk,autojunk=False).get_opcodes()
    hunks=[]; i=0
    while i < len(ops):
        while i < len(ops) and ops[i][0]=='equal':
            i += 1
        if i >= len(ops):
            break
        tag,a1,a2,b1,b2=ops[i]
        hs=[a1,a2,b1,b2]
        i += 1
        while i < len(ops):
            if ops[i][0] != 'equal':
                _,x1,x2,y1,y2=ops[i]
                hs[1]=x2; hs[3]=y2; i += 1
                continue
            _,e1,e2,f1,f2=ops[i]
            if _v34_strong_equal(at,e1,e2) or _v34_strong_equal(bt,f1,f2):
                break
            # Weak equal bridge: merge it only if another change follows before a
            # strong anchor/end.  Otherwise leave trailing context unchanged.
            j=i+1
            while j < len(ops) and ops[j][0]=='equal' and not (_v34_strong_equal(at,ops[j][1],ops[j][2]) or _v34_strong_equal(bt,ops[j][3],ops[j][4])):
                j += 1
            if j < len(ops) and ops[j][0] != 'equal':
                # Keep a lexical token as an anchor when punctuation moved across it.
                # Example: as "the Company" -> as the "Company" must remain
                # DELETE quote + ADD quote, not DELETE '"the' + ADD 'the "'.
                prev_op=ops[i-1] if i-1 >= 0 else None
                next_op=ops[j]
                def punct_only_opcode(op):
                    if not op or op[0]=='equal': return False
                    _,x1,x2,y1,y2=op
                    raw=''.join(at[x1:x2])+''.join(bt[y1:y2])
                    chars=[ch for ch in raw if not ch.isspace()]
                    return bool(chars) and all(not ch.isalnum() and ch!='_' for ch in chars)
                if punct_only_opcode(prev_op) and punct_only_opcode(next_op):
                    break
                hs[1]=e2; hs[3]=f2
                i += 1
                continue
            break
        hunks.append(tuple(hs))
    return at,bt,hunks


def _v34_directional_pair_marks(base_text: str, other_text: str):
    at,bt,hunks=_v34_change_hunks(base_text,other_text)
    am=['normal']*len(at); bm=['normal']*len(bt)
    for a1,a2,b1,b2 in hunks:
        for k in range(a1,a2): am[k]='delete'
        for k in range(b1,b2): bm[k]='insert'
    return at,am,bt,bm

# All body/header red/blue highlighting now uses the same phrase-level grouping.
_directional_pair_marks_v19 = _v34_directional_pair_marks


def _v34_tok_char_span(tokens: List[str], i1: int, i2: int):
    starts=[]; pos=0
    for t in tokens:
        starts.append(pos); pos += len(t)
    if i1 < i2:
        st=starts[i1]; en=starts[i2-1]+len(tokens[i2-1])
    else:
        st=starts[i1] if i1 < len(starts) else pos; en=st
    return st,en


def _v34_trim_span(text: str, st: int, en: int):
    raw=(text or '')[st:en]
    l=len(raw)-len(raw.lstrip()); r=len(raw)-len(raw.rstrip())
    return st+l, en-r


def _v34_header_equivalent(a: Unit, b: Unit) -> bool:
    if not a or not b: return False
    if (a.number or '').casefold() != (b.number or '').casefold(): return False
    na=(a.norm_title or normalize_text(a.title or '') or '').strip()
    nb=(b.norm_title or normalize_text(b.title or '') or '').strip()
    return bool(na or nb) and na==nb


def _v34_event_parts(full_text: str, st: int, en: int):
    """Keep one logical phrase as one marker; split only a detached edge punctuation run.

    Example '. These Terms ...' -> '.' and 'These Terms ...'.  Quotes attached to a
    phrase stay with the phrase unless the whole hunk is punctuation-only.
    """
    st,en=_v34_trim_span(full_text,st,en)
    if en<=st: return []
    raw=full_text[st:en]
    if not any(ch.isalnum() for ch in raw):
        return [(st,en,raw,'punct')]
    # Split a leading sentence delimiter only when whitespace separates it from text.
    m=re.match(r'^([.!?;:]+)(\s+)(.+)$',raw,flags=re.S)
    if m:
        p1=st; p2=st+len(m.group(1)); q1=p2+len(m.group(2)); q2=en
        return [(p1,p2,m.group(1),'punct'),(q1,q2,full_text[q1:q2],'lex')]
    return [(st,en,raw,'lex')]


def _v34_pair_marker_events(base: Unit, other: Unit, base_index: int, other_index: int):
    events=[]; label=f'문서 {chr(65+other_index)}'
    for attr,part_order in (('header',0),('body',1)):
        if attr=='header' and _v34_header_equivalent(base,other):
            continue
        old=getattr(base,attr) or ''; new=getattr(other,attr) or ''
        at,bt,hunks=_v34_change_hunks(old,new)
        old_off=0 if attr=='header' else len(base.header)+1
        new_off=0 if attr=='header' else len(other.header)+1
        for hi,(a1,a2,b1,b2) in enumerate(hunks):
            ast,aen=_v34_tok_char_span(at,a1,a2); bst,ben=_v34_tok_char_span(bt,b1,b2)
            # old endpoint first, then new endpoint: matches the user's physical-marker rule.
            for st,en,txt,kind in _v34_event_parts(old,ast,aen):
                if not txt.strip(): continue
                events.append({
                    'relative_doc':other_index,'target_doc':base_index,'action':'삭제','text':txt.strip(),
                    'char_start':old_off+st,'char_end':old_off+en,'part':attr,
                    'sort_key':(other_index,part_order,hi,0,st/max(1,len(old))),
                    'message':_v31_event_message('삭제',txt.strip(),old,st,en),'label':label,
                })
            for st,en,txt,kind in _v34_event_parts(new,bst,ben):
                if not txt.strip(): continue
                events.append({
                    'relative_doc':other_index,'target_doc':other_index,'action':'추가','text':txt.strip(),
                    'char_start':new_off+st,'char_end':new_off+en,'part':attr,
                    'sort_key':(other_index,part_order,hi,1,st/max(1,len(new))),
                    'message':_v31_event_message('추가',txt.strip(),new,st,en),'label':label,
                })
    return events

# Marker generation uses phrase-level hunks; overlay plumbing from V3.1 is unchanged.
_v31_pair_marker_events = _v34_pair_marker_events

# V3.1's compare wrapper resolves _v31_pair_marker_events at call time, so replacing the
# global above is sufficient.  Keep version visible in the window title.


# ---------------- V3.5: subpart-aware minimal diff ----------------
# Explicit legal items ((1), (2), 1., 2., ①...) are matched first.  Diff is then
# performed INSIDE the matched item, so a small edit such as
#   "use restriction" -> "the usage restriction measure"
# is reported as "use -> the usage" + "measure added" instead of replacing the whole item.
APP_VERSION = "3.5"

_V35_ENUM_LINE_RE = re.compile(
    r'(?m)^[ \t]*(?P<label>(?:[①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳]|\(\d+\)|\d+[.)]|[가-하A-Za-z][.)]))(?P<ws>[ \t]+)'
)


def _v35_parts_with_spans(text: str):
    """Return explicit legal parts with exact body offsets.

    Each dict contains label, whole span and core (text after enumerator) span.
    A leading unnumbered body becomes one '본문' part.  No sentence-period splitting occurs.
    """
    text=text or ''
    ms=list(_V35_ENUM_LINE_RE.finditer(text))
    if not ms:
        st=0; en=len(text)
        while st<en and text[st].isspace(): st+=1
        while en>st and text[en-1].isspace(): en-=1
        return [] if en<=st else [{'label':'본문','start':st,'end':en,'core_start':st,'core_end':en,'core':text[st:en]}]
    out=[]
    first=ms[0]
    lead=text[:first.start()]
    if lead.strip():
        st=0; en=first.start()
        while st<en and text[st].isspace(): st+=1
        while en>st and text[en-1].isspace(): en-=1
        if en>st:
            out.append({'label':'본문','start':st,'end':en,'core_start':st,'core_end':en,'core':text[st:en]})
    for i,m in enumerate(ms):
        seg_start=m.start(); seg_end=ms[i+1].start() if i+1<len(ms) else len(text)
        core_start=m.end(); core_end=seg_end
        while core_start<core_end and text[core_start].isspace(): core_start+=1
        while core_end>core_start and text[core_end-1].isspace(): core_end-=1
        out.append({
            'label':m.group('label'),'start':seg_start,'end':seg_end,
            'core_start':core_start,'core_end':core_end,'core':text[core_start:core_end],
        })
    return out


def _v35_part_similarity(a, b):
    ca=(a.get('core') or '').strip(); cb=(b.get('core') or '').strip()
    if not ca or not cb: return 0.0
    na=normalize_text(ca); nb=normalize_text(cb)
    if na==nb: return 1.0
    seq=SequenceMatcher(None,na,nb,autojunk=False).ratio()
    wa=set(token_words(ca)); wb=set(token_words(cb))
    jac=len(wa&wb)/max(1,len(wa|wb))
    contain=len(wa&wb)/max(1,min(len(wa),len(wb)))
    s=max(0.68*seq+0.32*jac,0.90*contain)
    if a.get('label')==b.get('label') and a.get('label')!='본문':
        s=min(1.0,s+0.08)
    return s


def _v35_match_parts(ap, bp):
    """Content-first legal-item matching, with same-label fallback for light rewrites."""
    used_a=set(); used_b=set(); pairs=[]
    # 1) exact normalized content, regardless of number/location (handles moved items).
    idx=defaultdict(list)
    for j,b in enumerate(bp):
        n=normalize_text(b.get('core',''))
        if n: idx[n].append(j)
    for i,a in enumerate(ap):
        n=normalize_text(a.get('core',''))
        if not n: continue
        cand=[j for j in idx.get(n,[]) if j not in used_b]
        if cand:
            j=min(cand,key=lambda x:abs(i-x)); pairs.append((i,j,1.0)); used_a.add(i); used_b.add(j)
    # 2) strong fuzzy content match.
    c=[]
    for i,a in enumerate(ap):
        if i in used_a: continue
        for j,b in enumerate(bp):
            if j in used_b: continue
            ss=_v35_part_similarity(a,b)
            if ss>=0.50: c.append((ss,i,j))
    for ss,i,j in sorted(c,reverse=True):
        if i in used_a or j in used_b: continue
        pairs.append((i,j,ss)); used_a.add(i); used_b.add(j)
    # 3) same enumerator fallback, only if still plausibly related.
    for i,a in enumerate(ap):
        if i in used_a or a.get('label')=='본문': continue
        cand=[]
        for j,b in enumerate(bp):
            if j in used_b or b.get('label')!=a.get('label'): continue
            ss=_v35_part_similarity(a,b)
            if ss>=0.28: cand.append((ss,j))
        if cand:
            ss,j=max(cand); pairs.append((i,j,ss)); used_a.add(i); used_b.add(j)
    return pairs,used_a,used_b


def _v35_local_hunks(old: str, new: str):
    """Fine diff inside one already-matched legal item.

    Unlike V3.4 clause-level grouping, one meaningful equal word is allowed as an anchor.
    This is intentional inside a short enumerated item: 'restriction' should separate
    'use -> the usage' from the trailing insertion 'measure'.
    """
    at=display_tokens(old or ''); bt=display_tokens(new or '')
    ak=[token_key(x) for x in at]; bk=[token_key(x) for x in bt]
    ops=SequenceMatcher(None,ak,bk,autojunk=False).get_opcodes()
    return at,bt,[(a1,a2,b1,b2) for tag,a1,a2,b1,b2 in ops if tag!='equal']


def _v35_mark_range(tokens, marks, i1, i2, style):
    for k in range(i1,i2): marks[k]=style


def _v35_directional_pair_marks(base_text: str, other_text: str):
    """Subpart-aware red/blue body highlighting."""
    at=display_tokens(base_text or ''); bt=display_tokens(other_text or '')
    am=['normal']*len(at); bm=['normal']*len(bt)
    # Fast token->character starts for whole-body style projection.
    astarts=[]; p=0
    for t in at: astarts.append(p); p+=len(t)
    bstarts=[]; p=0
    for t in bt: bstarts.append(p); p+=len(t)
    def tok_range(st,en,starts,toks):
        if not toks: return (0,0)
        i1=0
        while i1<len(toks) and starts[i1]+len(toks[i1])<=st: i1+=1
        i2=i1
        while i2<len(toks) and starts[i2]<en: i2+=1
        return i1,i2
    ap=_v35_parts_with_spans(base_text); bp=_v35_parts_with_spans(other_text)
    # Preserve V3.4/V3.0 punctuation-safe behavior for ordinary unnumbered prose.
    if len(ap)==1 and len(bp)==1 and ap[0].get('label')=='본문' and bp[0].get('label')=='본문':
        return _v34_directional_pair_marks(base_text,other_text)
    pairs,ua,ub=_v35_match_parts(ap,bp)
    for i,j,ss in pairs:
        a=ap[i]; b=bp[j]
        # Fine word-level diff is used only for clearly corresponding legal items.
        # Heavily rewritten items retain V3.4 phrase-level grouping to avoid a marker explosion.
        fine=(a.get('label')!='본문' and b.get('label')!='본문' and ss>=0.72)
        lat,lbt,hunks=(_v35_local_hunks(a['core'],b['core']) if fine else _v34_change_hunks(a['core'],b['core']))
        for x1,x2,y1,y2 in hunks:
            ast,aen=_v34_tok_char_span(lat,x1,x2); bst,ben=_v34_tok_char_span(lbt,y1,y2)
            if x1<x2:
                g1,g2=tok_range(a['core_start']+ast,a['core_start']+aen,astarts,at); _v35_mark_range(at,am,g1,g2,'delete')
            if y1<y2:
                g1,g2=tok_range(b['core_start']+bst,b['core_start']+ben,bstarts,bt); _v35_mark_range(bt,bm,g1,g2,'insert')
    for i,a in enumerate(ap):
        if i not in ua:
            g1,g2=tok_range(a['core_start'],a['core_end'],astarts,at); _v35_mark_range(at,am,g1,g2,'delete')
    for j,b in enumerate(bp):
        if j not in ub:
            g1,g2=tok_range(b['core_start'],b['core_end'],bstarts,bt); _v35_mark_range(bt,bm,g1,g2,'insert')
    return at,am,bt,bm


def _v35_span_text(text, toks, i1, i2):
    st,en=_v34_tok_char_span(toks,i1,i2)
    st,en=_v34_trim_span(text,st,en)
    return st,en,text[st:en]


def _v35_pair_marker_events(base: Unit, other: Unit, base_index: int, other_index: int):
    events=[]; label=f'문서 {chr(65+other_index)}'
    # Header keeps mature V3.4 behavior.
    if not _v34_header_equivalent(base,other):
        old=base.header or ''; new=other.header or ''
        at,bt,hunks=_v34_change_hunks(old,new)
        for hi,(a1,a2,b1,b2) in enumerate(hunks):
            ast,aen=_v34_tok_char_span(at,a1,a2); bst,ben=_v34_tok_char_span(bt,b1,b2)
            for st,en,txt,kind in _v34_event_parts(old,ast,aen):
                if txt.strip():
                    events.append({'relative_doc':other_index,'target_doc':base_index,'action':'삭제','text':txt.strip(),
                        'char_start':st,'char_end':en,'part':'header','sort_key':(other_index,0,hi,0,st/max(1,len(old))),
                        'message':_v31_event_message('삭제',txt.strip(),old,st,en),'label':label})
            for st,en,txt,kind in _v34_event_parts(new,bst,ben):
                if txt.strip():
                    events.append({'relative_doc':other_index,'target_doc':other_index,'action':'추가','text':txt.strip(),
                        'char_start':st,'char_end':en,'part':'header','sort_key':(other_index,0,hi,1,st/max(1,len(new))),
                        'message':_v31_event_message('추가',txt.strip(),new,st,en),'label':label})

    old=base.body or ''; new=other.body or ''
    ap=_v35_parts_with_spans(old); bp=_v35_parts_with_spans(new)
    # No explicit legal enumerators: keep the mature punctuation-safe V3.4 marker engine.
    if len(ap)==1 and len(bp)==1 and ap[0].get('label')=='본문' and bp[0].get('label')=='본문':
        events.extend(e for e in _v34_pair_marker_events(base,other,base_index,other_index) if e.get('part')=='body')
        return events
    pairs,ua,ub=_v35_match_parts(ap,bp)
    old_off=len(base.header)+1; new_off=len(other.header)+1
    seq=0
    for i,j,ss in sorted(pairs,key=lambda x:(x[0],x[1])):
        a=ap[i]; b=bp[j]; seq+=1
        fine=(a.get('label')!='본문' and b.get('label')!='본문' and ss>=0.72)
        lat,lbt,hunks=(_v35_local_hunks(a['core'],b['core']) if fine else _v34_change_hunks(a['core'],b['core']))
        for hi,(x1,x2,y1,y2) in enumerate(hunks):
            ast,aen,otxt=_v35_span_text(a['core'],lat,x1,x2)
            bst,ben,ntxt=_v35_span_text(b['core'],lbt,y1,y2)
            oa1=a['core_start']+ast; oa2=a['core_start']+aen
            nb1=b['core_start']+bst; nb2=b['core_start']+ben
            # A replacement is one logical change with one number and two physical endpoints.
            if x1<x2 and y1<y2 and otxt.strip() and ntxt.strip():
                events.append({
                    'relative_doc':other_index,'target_doc':other_index,'action':'변경','text':ntxt.strip(),
                    'char_start':new_off+nb1,'char_end':new_off+nb2,'part':'body',
                    'endpoints':[
                        {'target_doc':base_index,'char_start':old_off+oa1,'char_end':old_off+oa2},
                        {'target_doc':other_index,'char_start':new_off+nb1,'char_end':new_off+nb2},
                    ],
                    'sort_key':(other_index,1,i,hi,0),
                    'message':f'[{a.get("label")}] 변경: “{re.sub(r"\s+", " ", otxt.strip())}” → “{re.sub(r"\s+", " ", ntxt.strip())}”',
                    'label':label,
                })
            elif x1<x2 and otxt.strip():
                events.append({'relative_doc':other_index,'target_doc':base_index,'action':'삭제','text':otxt.strip(),
                    'char_start':old_off+oa1,'char_end':old_off+oa2,'part':'body','sort_key':(other_index,1,i,hi,0),
                    'message':f'[{a.get("label")}] '+_v31_event_message('삭제',otxt.strip(),old,oa1,oa2),'label':label})
            elif y1<y2 and ntxt.strip():
                events.append({'relative_doc':other_index,'target_doc':other_index,'action':'추가','text':ntxt.strip(),
                    'char_start':new_off+nb1,'char_end':new_off+nb2,'part':'body','sort_key':(other_index,1,i,hi,1),
                    'message':f'[{b.get("label")}] '+_v31_event_message('추가',ntxt.strip(),new,nb1,nb2),'label':label})
        # Structural item-number/order changes are kept separately by existing summary engine.
    for i,a in enumerate(ap):
        if i in ua: continue
        txt=a['core'].strip()
        if txt:
            events.append({'relative_doc':other_index,'target_doc':base_index,'action':'삭제','text':txt,
                'char_start':old_off+a['core_start'],'char_end':old_off+a['core_end'],'part':'body',
                'sort_key':(other_index,1,i,999,0),'message':f'[{a.get("label")}] 삭제: “{re.sub(r"\s+", " ", txt).strip()}”','label':label})
    for j,b in enumerate(bp):
        if j in ub: continue
        txt=b['core'].strip()
        if txt:
            events.append({'relative_doc':other_index,'target_doc':other_index,'action':'추가','text':txt,
                'char_start':new_off+b['core_start'],'char_end':new_off+b['core_end'],'part':'body',
                'sort_key':(other_index,1,j,999,1),'message':f'[{b.get("label")}] 추가: “{re.sub(r"\s+", " ", txt).strip()}”','label':label})
    return events


# Use the same subpart-aware algorithm for colors and numbered marker events.
_directional_pair_marks_v19 = _v35_directional_pair_marks
_v31_pair_marker_events = _v35_pair_marker_events


def _v35_make_doc_cell(self,parent,row,i,h):
    t=self._new_text(parent,h); m=row['members'][i]
    if not m:
        t.insert('end','[해당 조항 없음]','missing')
    else:
        hs=row.get('header_segments',[[] for _ in row['segments']])[i] if row.get('header_segments') else []
        if hs:
            for txt,sty in hs: t.insert('end',txt,sty if sty in ('delete','insert') else 'article_header')
            t.insert('end','\n')
        else:
            t.insert('end',m.get('header','')+'\n','article_header')
        for txt,sty in row.get('body_segments',row['segments'])[i]:
            t.insert('end',txt,sty if sty in ('delete','insert') else ())
    t.configure(state='disabled')
    self._v31_doc_widgets[(row['id'],i)]=t
    placements=[]
    for ev in row.get('markers',[]):
        eps=ev.get('endpoints')
        if eps:
            for ep in eps:
                if ep.get('target_doc')==i:
                    clone=dict(ev); clone['target_doc']=i; clone['char_start']=ep['char_start']; clone['char_end']=ep['char_end']; placements.append(clone)
        elif ev.get('target_doc')==i:
            placements.append(ev)
    if placements:
        def place_all(widget=t, rid=row['id'], events=tuple(placements)):
            for ev in events: _v31_place_badge(self,widget,rid,ev)
        t.after_idle(place_all)
    return t


def _v35_focus_marker(self,row_id:int,num:int):
    ev=self._v31_event_map.get((row_id,num))
    if not ev: return
    eps=ev.get('endpoints') or [{'target_doc':ev.get('target_doc'),'char_start':ev.get('char_start'),'char_end':ev.get('char_end')}]
    for ep in eps:
        t=self._v31_doc_widgets.get((row_id,ep.get('target_doc')))
        if not t: continue
        try:
            s=f"1.0+{int(ep['char_start'])}c"; e=f"1.0+{int(ep['char_end'])}c"
            t.see(s); t.configure(state='normal'); t.tag_configure('marker_focus',background='#fff59d'); t.tag_add('marker_focus',s,e); t.tag_raise('marker_focus'); t.configure(state='disabled')
            def clear(widget=t):
                try: widget.configure(state='normal'); widget.tag_remove('marker_focus','1.0','end'); widget.configure(state='disabled')
                except Exception: pass
            t.after(1100,clear)
        except Exception:
            logging.exception('focus marker failed')

NativeGui._make_doc_cell = _v35_make_doc_cell
NativeGui._v31_focus_marker = _v35_focus_marker




# ---------------- V3.6: GUI-engine-native Word tracked changes ----------------
# Word CompareDocuments is deliberately NOT used here.  It can align clauses differently
# from the GUI (e.g. merge the next Article into the current revision).  V3.6 rebuilds a
# native tracked-changes DOCX directly from this application's own clause/item lineage and
# diff engine, so Word shows the same logical pairings as the GUI.
APP_VERSION = "3.6"


def _v36_chunks_from_hunks(old: str, new: str, fine: bool=False):
    """Return ordered (normal|delete|insert, text) chunks using the same V3.4/V3.5 diff rules."""
    old=old or ''; new=new or ''
    if old==new:
        return [('normal',old)] if old else []
    at,bt,hunks=(_v35_local_hunks(old,new) if fine else _v34_change_hunks(old,new))
    out=[]; ai=bi=0
    for a1,a2,b1,b2 in hunks:
        # Text between hunks is an equal anchor under the GUI's token-key rules.
        common=''.join(at[ai:a1])
        if common: out.append(('normal',common))
        if a1<a2:
            txt=''.join(at[a1:a2])
            if txt: out.append(('delete',txt))
        if b1<b2:
            txt=''.join(bt[b1:b2])
            if txt: out.append(('insert',txt))
        ai,bi=a2,b2
    tail=''.join(at[ai:])
    if tail: out.append(('normal',tail))
    # Coalesce adjacent chunks of the same type.
    merged=[]
    for kind,txt in out:
        if not txt: continue
        if merged and merged[-1][0]==kind:
            merged[-1]=(kind,merged[-1][1]+txt)
        else:
            merged.append((kind,txt))
    return merged


def _v36_body_chunks(old: str, new: str):
    """Subpart-aware redline chunks; moved items become delete+insert so accept/reject are reversible."""
    old=old or ''; new=new or ''
    if old==new:
        return [('normal',old)] if old else []
    ap=_v35_parts_with_spans(old); bp=_v35_parts_with_spans(new)
    if len(ap)==1 and len(bp)==1 and ap[0].get('label')=='본문' and bp[0].get('label')=='본문':
        return _v36_chunks_from_hunks(old,new,False)

    pairs,ua,ub=_v35_match_parts(ap,bp)
    pair_info={}
    ida={}; idb={}
    for k,(i,j,ss) in enumerate(sorted(pairs,key=lambda x:(x[0],x[1]))):
        ident=('M',k); ida[i]=ident; idb[j]=ident; pair_info[ident]=(i,j,ss)
    for i in range(len(ap)):
        if i not in ida: ida[i]=('A',i)
    for j in range(len(bp)):
        if j not in idb: idb[j]=('B',j)
    aseq=[ida[i] for i in range(len(ap))]
    bseq=[idb[j] for j in range(len(bp))]
    sm=SequenceMatcher(None,aseq,bseq,autojunk=False)
    out=[]

    def atext(i):
        st=ap[i]['start']; en=(ap[i+1]['start'] if i+1<len(ap) else len(old))
        return old[st:en]
    def btext(j):
        st=bp[j]['start']; en=(bp[j+1]['start'] if j+1<len(bp) else len(new))
        return new[st:en]

    for tag,a1,a2,b1,b2 in sm.get_opcodes():
        if tag=='equal':
            for off in range(a2-a1):
                ident=aseq[a1+off]
                if ident and ident[0]=='M':
                    i,j,ss=pair_info[ident]
                    ao=atext(i); bn=btext(j)
                    fine=(ap[i].get('label')!='본문' and bp[j].get('label')!='본문' and ss>=0.72)
                    out.extend(_v36_chunks_from_hunks(ao,bn,fine))
                else:
                    # Should not happen for equal unmatched identities, but keep it safe.
                    txt=atext(a1+off)
                    if txt: out.append(('normal',txt))
        elif tag=='delete':
            txt=''.join(atext(i) for i in range(a1,a2))
            if txt: out.append(('delete',txt))
        elif tag=='insert':
            txt=''.join(btext(j) for j in range(b1,b2))
            if txt: out.append(('insert',txt))
        else:  # replace: includes moved / renumbered parts that crossed order
            oldtxt=''.join(atext(i) for i in range(a1,a2))
            newtxt=''.join(btext(j) for j in range(b1,b2))
            if oldtxt: out.append(('delete',oldtxt))
            if newtxt: out.append(('insert',newtxt))
    merged=[]
    for kind,txt in out:
        if not txt: continue
        if merged and merged[-1][0]==kind:
            merged[-1]=(kind,merged[-1][1]+txt)
        else:
            merged.append((kind,txt))
    return merged


def _v36_section_key(s: str) -> str:
    s=(s or '').strip()
    s=re.sub(r'(?i)\b(?:chapter|part|section)\s+[IVXLCDM\d]+\s*[.:\-]?\s*','',s)
    s=re.sub(r'제\s*\d+\s*(?:장|절|관)\s*','',s)
    return normalize_text(s) or normalize_text((s or ''))


def _v36_extract_preamble(text: str) -> str:
    text=(text or '').replace('\r\n','\n').replace('\r','\n')
    pat=re.compile(r'(?im)^\s*(?:chapter\s+[IVXLCDM\d]+\b|part\s+[IVXLCDM\d]+\b|article\s+[IVXLCDM\d]+\b|제\s*\d+\s*(?:장|절|관|조)\b)')
    m=pat.search(text)
    return (text[:m.start()] if m else '').strip()


def _v36_pair_nodes(result: Dict[str,Any], old_text: str, new_text: str):
    """Build reversible document-order node sequences from the pairwise GUI lineage result."""
    rows=result.get('rows',[])
    row_by_id={r['id']:r for r in rows}
    old_units=[]; new_units=[]; old_id={}; new_id={}
    for r in rows:
        ms=r.get('members') or [None,None]
        a=ms[0] if len(ms)>0 else None; b=ms[1] if len(ms)>1 else None
        if a and b: ident=('M',r['id'])
        elif a: ident=('D',r['id'])
        else: ident=('I',r['id'])
        if a:
            old_units.append((int(a.get('index',0)),a,ident)); old_id[int(a.get('index',0))]=ident
        if b:
            new_units.append((int(b.get('index',0)),b,ident)); new_id[int(b.get('index',0))]=ident
    old_units.sort(); new_units.sort()

    # Match section/chapter headings by semantic title (number may change).
    def build(side_units, side):
        seq=[]; content={}; sec_occ={}; prev_sec=None
        pre=_v36_extract_preamble(old_text if side==0 else new_text)
        if pre:
            seq.append(('P',0)); content[('P',0)]=pre
        for _,u,aid in side_units:
            sec=(u.get('section') or '').strip()
            if sec and sec!=prev_sec:
                key=_v36_section_key(sec); occ=sec_occ.get(key,0); sec_occ[key]=occ+1
                sid=('S',key,occ)
                seq.append(sid); content[sid]=sec; prev_sec=sec
            seq.append(aid); content[aid]=u
        return seq,content
    aseq,acont=build(old_units,0); bseq,bcont=build(new_units,1)
    return aseq,bseq,acont,bcont,row_by_id


def _v36_append_run_xml(parent, text: str, *, deleted=False, bold=False):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    if not text: return
    r=OxmlElement('w:r')
    if bold:
        rpr=OxmlElement('w:rPr'); b=OxmlElement('w:b'); rpr.append(b); r.append(rpr)
    pieces=text.split('\n')
    for i,piece in enumerate(pieces):
        if piece:
            t=OxmlElement('w:delText' if deleted else 'w:t')
            t.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
            t.text=piece; r.append(t)
        if i < len(pieces)-1:
            br=OxmlElement('w:br'); r.append(br)
    parent.append(r)


def _v36_append_chunk(paragraph, kind: str, text: str, revstate: Dict[str,Any], bold=False):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    if not text: return
    if kind=='normal':
        _v36_append_run_xml(paragraph._p,text,deleted=False,bold=bold); return
    tag='w:del' if kind=='delete' else 'w:ins'
    wrap=OxmlElement(tag)
    wrap.set(qn('w:id'),str(revstate['id'])); revstate['id']+=1
    wrap.set(qn('w:author'),revstate['author'])
    wrap.set(qn('w:date'),revstate['date'])
    _v36_append_run_xml(wrap,text,deleted=(kind=='delete'),bold=bold)
    paragraph._p.append(wrap)


def _v36_add_chunks_paragraph(doc: Document, chunks, revstate, bold=False):
    p=doc.add_paragraph()
    for kind,text in chunks:
        _v36_append_chunk(p,kind,text,revstate,bold=bold)
    return p


def _v36_add_article(doc: Document, old_u, new_u, revstate, mode='matched'):
    if mode=='delete':
        _v36_add_chunks_paragraph(doc,[('delete',old_u.get('header',''))],revstate,bold=True)
        if old_u.get('body'): _v36_add_chunks_paragraph(doc,[('delete',old_u.get('body',''))],revstate)
        return
    if mode=='insert':
        _v36_add_chunks_paragraph(doc,[('insert',new_u.get('header',''))],revstate,bold=True)
        if new_u.get('body'): _v36_add_chunks_paragraph(doc,[('insert',new_u.get('body',''))],revstate)
        return
    oh=old_u.get('header',''); nh=new_u.get('header','')
    if _v34_header_equivalent(Unit(**old_u),Unit(**new_u)):
        hc=[('normal',oh)]
    else:
        hc=_v36_chunks_from_hunks(oh,nh,False)
    _v36_add_chunks_paragraph(doc,hc,revstate,bold=True)
    bc=_v36_body_chunks(old_u.get('body',''),new_u.get('body',''))
    if bc or old_u.get('body') or new_u.get('body'):
        _v36_add_chunks_paragraph(doc,bc,revstate)


def _v36_enable_track_revisions(doc: Document):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    settings=doc.settings.element
    if settings.find(qn('w:trackRevisions')) is None:
        settings.insert(0,OxmlElement('w:trackRevisions'))


def _v36_build_tracked_doc(original_path: Path, revised_path: Path, output_path: Path, revised_author: str):
    from datetime import datetime, timezone
    original_path=Path(original_path); revised_path=Path(revised_path); output_path=Path(output_path)
    old_text=read_document(original_path.name,original_path.read_bytes())
    new_text=read_document(revised_path.name,revised_path.read_bytes())
    # IMPORTANT: pairwise compare with the requested Original as base.  This is the same
    # lineage/item/diff engine used by the GUI; no second Word-side alignment occurs.
    pair=compare_documents([original_path.name,revised_path.name],[old_text,new_text],base_index=0)
    aseq,bseq,acont,bcont,row_by_id=_v36_pair_nodes(pair,old_text,new_text)

    doc=Document(); _v36_enable_track_revisions(doc)
    # Remove python-docx's initial empty paragraph only if it exists after we add content later.
    revstate={'id':1,'author':str(revised_author or revised_path.stem),
              'date':datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace('+00:00','Z')}
    sm=SequenceMatcher(None,aseq,bseq,autojunk=False)

    def emit_ident(ident, mode):
        kind=ident[0]
        if kind=='P':
            if mode=='matched': chunks=_v36_chunks_from_hunks(acont.get(ident,''),bcont.get(ident,''),False)
            elif mode=='delete': chunks=[('delete',acont.get(ident,''))]
            else: chunks=[('insert',bcont.get(ident,''))]
            if any(t for _,t in chunks): _v36_add_chunks_paragraph(doc,chunks,revstate)
        elif kind=='S':
            if mode=='matched': chunks=_v36_chunks_from_hunks(acont.get(ident,''),bcont.get(ident,''),False)
            elif mode=='delete': chunks=[('delete',acont.get(ident,''))]
            else: chunks=[('insert',bcont.get(ident,''))]
            if any(t for _,t in chunks): _v36_add_chunks_paragraph(doc,chunks,revstate,bold=True)
        else:
            row=row_by_id[ident[1]]; ms=row.get('members') or [None,None]
            if mode=='matched': _v36_add_article(doc,ms[0],ms[1],revstate,'matched')
            elif mode=='delete': _v36_add_article(doc,ms[0],None,revstate,'delete')
            else: _v36_add_article(doc,None,ms[1],revstate,'insert')

    for tag,a1,a2,b1,b2 in sm.get_opcodes():
        if tag=='equal':
            for ident in aseq[a1:a2]: emit_ident(ident,'matched')
        elif tag=='delete':
            for ident in aseq[a1:a2]: emit_ident(ident,'delete')
        elif tag=='insert':
            for ident in bseq[b1:b2]: emit_ident(ident,'insert')
        else:
            for ident in aseq[a1:a2]: emit_ident(ident,'delete')
            for ident in bseq[b1:b2]: emit_ident(ident,'insert')

    # Remove a leading/trailing empty paragraph if python-docx happened to create one.
    for p in list(doc.paragraphs):
        if not p.text and len(p._p)==1 and p._p.pPr is not None:
            try: p._element.getparent().remove(p._element)
            except Exception: pass
    output_path.parent.mkdir(parents=True,exist_ok=True)
    doc.save(str(output_path))
    return output_path


def create_word_tracked_compare(original_path: Path, revised_path: Path, output_path: Path,
                                revised_author: str='Revised', word_factory=None) -> Path:
    """Create a native tracked-changes DOCX from the GUI's own comparison engine.

    Reject All Changes reconstructs the original-side logical text; Accept All Changes
    reconstructs the revised-side logical text.  Word CompareDocuments is intentionally
    bypassed so Article/item lineage never diverges from the GUI.
    """
    return _v36_build_tracked_doc(Path(original_path),Path(revised_path),Path(output_path),revised_author)


def _v36_word_export_failed(self,msg):
    self.progress_var.set(0)
    self.status_var.set('Word 변경추적 문서 생성에 실패했습니다.')
    if self.result and len(self.result.get('names',[])) in (2,3):
        self.word_export_btn.configure(state='normal')
    self.messagebox.showerror('Word 변경추적 오류',msg)

NativeGui._word_export_failed=_v36_word_export_failed


# ---------------- V3.7: robust on-text numbered overlay badges ----------------
# The visible change number is a true overlay child of each Text widget.  It does not
# consume a character, so line wrapping and the source text remain untouched.  Badge
# coordinates are recalculated after Text resize/reflow so markers follow the changed
# character even when the window width/DPI changes.
APP_VERSION = "3.7"


def _v37_marker_role(self, event: Dict[str,Any], target_doc: int) -> str:
    action=(event.get('action') or '').strip()
    if action=='삭제':
        return 'delete'
    if action=='추가':
        return 'insert'
    # A logical replacement has one number but two physical endpoints.  The baseline
    # endpoint is a deletion and the revised endpoint is an insertion.
    if action=='변경':
        bi=int((self.result or {}).get('base_index',0))
        return 'delete' if int(target_doc)==bi else 'insert'
    return 'change'


def _v37_badge_palette(role: str):
    if role=='delete':
        return '#c62828', '#ffffff'
    if role=='insert':
        return '#1565c0', '#ffffff'
    return '#6a1b9a', '#ffffff'


def _v37_bbox_for_offset(widget, char_start: int):
    """Get a usable bbox even for a marker anchored at the logical end of text."""
    try:
        idx=f"1.0+{max(0,int(char_start))}c"
        box=widget.bbox(idx)
        if box:
            return box, False
        # Insertion exactly at end has no glyph bbox.  Anchor after previous glyph.
        if int(char_start)>0:
            pidx=f"1.0+{int(char_start)-1}c"
            pbox=widget.bbox(pidx)
            if pbox:
                x,y,w,h=pbox
                return (x+w,y,1,h), True
    except Exception:
        pass
    return None, False


def _v37_destroy_widget_badges(self, widget):
    badges=getattr(self,'_v37_widget_badges',{}).pop(widget,[])
    for b in badges:
        try: b.destroy()
        except Exception: pass


def _v37_refresh_widget_badges(self, widget):
    if not widget or not widget.winfo_exists():
        return
    _v37_destroy_widget_badges(self,widget)
    events=list(getattr(self,'_v37_widget_events',{}).get(widget,[]) or [])
    if not events:
        return
    try:
        widget.update_idletasks()
    except Exception:
        return
    made=[]
    occupied=[]
    for ev in events:
        res=_v37_bbox_for_offset(widget,ev.get('char_start',0))
        if not res:
            continue
        (x,y,w,h),at_end=res
        role=ev.get('_overlay_role') or _v37_marker_role(self,ev,ev.get('target_doc',0))
        bg,fg=_v37_badge_palette(role)
        num=int(ev.get('num',0) or 0)
        label=_v31_circled(num)
        # Keep badges over the changed location, but stagger collisions so two nearby
        # changes remain individually clickable.  This is especially important for
        # punctuation delete/insert pairs around the same word.
        bx=max(1,x-3)
        by=max(1,y-6)
        for ox,oy in occupied:
            if abs(bx-ox)<18 and abs(by-oy)<14:
                by += 12
        occupied.append((bx,by))
        badge=self.tk.Label(
            widget,text=label,font=('Malgun Gothic',8,'bold'),
            fg=fg,bg=bg,bd=1,relief='solid',padx=2,pady=0,
            cursor='hand2',takefocus=0
        )
        badge.place(x=bx,y=by,anchor='nw')
        badge.lift()
        rid=int(ev.get('_row_id',-1)); n=num
        badge.bind('<Button-1>',lambda _e,r=rid,k=n:self._v31_focus_summary(r,k))
        badge.bind('<Enter>',lambda _e,r=rid,k=n:self._v31_focus_summary(r,k,flash_only=True))
        badge.bind('<MouseWheel>',self._child_wheel)
        made.append(badge)
    self._v37_widget_badges[widget]=made


def _v37_schedule_badge_refresh(self, widget, delay=55):
    if not widget or not widget.winfo_exists():
        return
    jobs=getattr(self,'_v37_refresh_jobs',{})
    old=jobs.get(widget)
    if old:
        try: widget.after_cancel(old)
        except Exception: pass
    try:
        jobs[widget]=widget.after(delay,lambda w=widget:_v37_refresh_widget_badges(self,w))
    except Exception:
        return
    self._v37_refresh_jobs=jobs


def _v37_make_doc_cell(self,parent,row,i,h):
    t=self._new_text(parent,h); m=row['members'][i]
    if not m:
        t.insert('end','[해당 조항 없음]','missing')
    else:
        hs=row.get('header_segments',[[] for _ in row['segments']])[i] if row.get('header_segments') else []
        if hs:
            for txt,sty in hs:
                t.insert('end',txt,sty if sty in ('delete','insert') else 'article_header')
            t.insert('end','\n')
        else:
            t.insert('end',m.get('header','')+'\n','article_header')
        for txt,sty in row.get('body_segments',row['segments'])[i]:
            t.insert('end',txt,sty if sty in ('delete','insert') else ())
    t.configure(state='disabled')
    self._v31_doc_widgets[(row['id'],i)]=t

    placements=[]
    for ev in row.get('markers',[]):
        eps=ev.get('endpoints')
        if eps:
            for ep in eps:
                if int(ep.get('target_doc',-1))==i:
                    clone=dict(ev)
                    clone['target_doc']=i
                    clone['char_start']=int(ep['char_start'])
                    clone['char_end']=int(ep['char_end'])
                    clone['_overlay_role']=_v37_marker_role(self,ev,i)
                    clone['_row_id']=row['id']
                    placements.append(clone)
        elif int(ev.get('target_doc',-1))==i:
            clone=dict(ev)
            clone['_overlay_role']=_v37_marker_role(self,ev,i)
            clone['_row_id']=row['id']
            placements.append(clone)

    if placements:
        self._v37_widget_events[t]=placements
        # Recalculate after initial mapping and every width/height change.  Text wrapping
        # changes bbox coordinates, so a one-time placement is not sufficient.
        t.bind('<Configure>',lambda _e,w=t:_v37_schedule_badge_refresh(self,w),add='+')
        t.bind('<Map>',lambda _e,w=t:_v37_schedule_badge_refresh(self,w,20),add='+')
        _v37_schedule_badge_refresh(self,t,90)
    return t


_v37_render_rows_base = NativeGui.render_rows

def _v37_render_rows(self):
    # Destroy any overlays from a previous rendering before the row widgets are rebuilt.
    for badges in list(getattr(self,'_v37_widget_badges',{}).values()):
        for b in badges:
            try: b.destroy()
            except Exception: pass
    self._v37_widget_events={}
    self._v37_widget_badges={}
    self._v37_refresh_jobs={}
    result=_v37_render_rows_base(self)
    # Some platforms finish geometry propagation one idle turn later than Text <Map>.
    # A second pass makes the first screen deterministic.
    def refresh_all():
        for w in list(getattr(self,'_v37_widget_events',{})):
            _v37_schedule_badge_refresh(self,w,0)
    try: self.root.after(120,refresh_all)
    except Exception: pass
    return result


NativeGui._make_doc_cell = _v37_make_doc_cell
NativeGui.render_rows = _v37_render_rows
NativeGui._v37_refresh_widget_badges = _v37_refresh_widget_badges
NativeGui._v37_schedule_badge_refresh = _v37_schedule_badge_refresh

def main():
    tk,ttk,filedialog,messagebox=_safe_import_tk(); dnd=False
    try:
        from tkinterdnd2 import TkinterDnD
        root=TkinterDnD.Tk(); dnd=True
    except Exception:
        root=tk.Tk()
    try:
        if os.name=='nt':
            import ctypes
            try: ctypes.windll.shcore.SetProcessDpiAwareness(1)
            except Exception: pass
    except Exception: pass
    NativeGui(root,dnd_available=dnd); root.mainloop()

# ---------------- V3.8: non-obscuring margin marker rail ----------------
# V3.7 placed numbered badges directly over changed glyphs.  That made the exact mapping
# obvious, but the badge could hide the very text being reviewed.  V3.8 moves markers into
# a reserved left-side annotation rail inside each document Text widget.  The source text
# starts after the rail, so markers never cover characters.  Markers on the same visual
# line are grouped into one readable pill (e.g. "1·2·3").  Clicking a right-side summary
# number still jumps to and highlights the exact changed span; clicking a rail pill
# highlights every change represented by that pill.
APP_VERSION = "3.8"

_V38_GUTTER_PX = 48
_V38_BADGE_X = 5
_V38_TICK_X = 39


def _v38_role_color(role: str) -> str:
    if role == 'delete':
        return '#c62828'
    if role == 'insert':
        return '#1565c0'
    return '#6a1b9a'


def _v38_apply_marker_rail(widget):
    """Reserve a left annotation rail without inserting any characters."""
    try:
        state=str(widget.cget('state'))
        if state == 'disabled':
            widget.configure(state='normal')
        # This tag only controls paragraph margins.  Existing delete/insert/article tags
        # continue to control text colour/underline/strike-through.
        widget.tag_configure('marker_rail', lmargin1=_V38_GUTTER_PX, lmargin2=_V38_GUTTER_PX)
        widget.tag_add('marker_rail', '1.0', 'end')
        # Keep the rail tag low so style tags remain visually authoritative.
        try: widget.tag_lower('marker_rail')
        except Exception: pass
        if state == 'disabled':
            widget.configure(state='disabled')
    except Exception:
        logging.exception('failed to apply marker rail')


def _v38_destroy_widget_badges(self, widget):
    items=getattr(self,'_v37_widget_badges',{}).pop(widget,[])
    for obj in items:
        try: obj.destroy()
        except Exception: pass


def _v38_refresh_widget_badges(self, widget):
    if not widget or not widget.winfo_exists():
        return
    _v38_destroy_widget_badges(self,widget)
    events=list(getattr(self,'_v37_widget_events',{}).get(widget,[]) or [])
    if not events:
        return
    try:
        widget.update_idletasks()
    except Exception:
        return

    # Build visual-line groups after the Text widget has wrapped/reflowed.
    raw=[]
    for ev in events:
        res=_v37_bbox_for_offset(widget,ev.get('char_start',0))
        if not res:
            continue
        (x,y,w,h),at_end=res
        role=ev.get('_overlay_role') or _v37_marker_role(self,ev,ev.get('target_doc',0))
        raw.append((y,h,ev,(x,y,w,h),role))
    raw.sort(key=lambda z:(z[0], int(z[2].get('num',0))))
    groups=[]
    for y,h,ev,box,role in raw:
        if groups and abs(groups[-1]['y']-y) <= 3:
            groups[-1]['items'].append((ev,box,role))
            groups[-1]['h']=max(groups[-1]['h'],h)
        else:
            groups.append({'y':y,'h':h,'items':[(ev,box,role)]})

    made=[]
    for grp in groups:
        items=grp['items']
        nums=[]
        roles=[]
        for ev,_box,role in items:
            n=int(ev.get('num',0) or 0)
            if n and n not in nums:
                nums.append(n)
            roles.append(role)
        if not nums:
            continue
        nums.sort()
        # Show all numbers when reasonably compact; otherwise show a range/count while
        # retaining exact one-to-one mapping in the right summary pane.
        joined='·'.join(str(n) for n in nums)
        if len(joined) > 8:
            if len(nums) >= 2:
                joined=f'{nums[0]}–{nums[-1]}'
            else:
                joined=str(nums[0])
        role=roles[0] if roles and all(r==roles[0] for r in roles) else 'change'
        color=_v38_role_color(role)
        by=max(1,grp['y'] + max(0,(grp['h']-18)//2))

        badge=self.tk.Label(
            widget,text=joined,font=('Malgun Gothic',9,'bold'),
            fg=color,bg='#ffffff',bd=1,relief='solid',padx=3,pady=0,
            cursor='hand2',takefocus=0
        )
        badge.place(x=_V38_BADGE_X,y=by,anchor='nw')
        badge.lift()

        # Short coloured tick points from the rail toward the corresponding text line
        # without crossing over or obscuring any characters.
        tick=self.tk.Frame(widget,bg=color,width=7,height=2,takefocus=0)
        tick.place(x=_V38_TICK_X,y=by+8,anchor='nw')
        tick.lift()

        rid=int(items[0][0].get('_row_id',-1))
        def focus_group(_e=None, r=rid, ns=tuple(nums)):
            for n in ns:
                try: self._v31_focus_marker(r,n)
                except Exception: pass
            if ns:
                try: self._v31_focus_summary(r,ns[0],flash_only=True)
                except Exception: pass
        badge.bind('<Button-1>',focus_group)
        badge.bind('<Enter>',focus_group)
        badge.bind('<MouseWheel>',self._child_wheel)
        tick.bind('<Button-1>',focus_group)
        tick.bind('<Enter>',focus_group)
        tick.bind('<MouseWheel>',self._child_wheel)
        made.extend([badge,tick])
    self._v37_widget_badges[widget]=made


_v38_make_doc_cell_base = NativeGui._make_doc_cell

def _v38_make_doc_cell(self,parent,row,i,h):
    t=_v38_make_doc_cell_base(self,parent,row,i,h)
    _v38_apply_marker_rail(t)
    # Reflow after the margin is applied, then recalculate rail-marker y coordinates.
    if getattr(self,'_v37_widget_events',{}).get(t):
        _v37_schedule_badge_refresh(self,t,40)
    return t


# Replace V3.7's on-text placement with the non-obscuring annotation rail.  Existing
# schedule callbacks resolve the global function name at runtime, so resizing/reflow keeps
# using this V3.8 implementation as well.
_v37_refresh_widget_badges = _v38_refresh_widget_badges
NativeGui._make_doc_cell = _v38_make_doc_cell
NativeGui._v37_refresh_widget_badges = _v38_refresh_widget_badges



# ---------------- V3.9: true side marker rail + text-hover linkage ----------------
# IMPORTANT: V3.8's patch block used to be located after the __main__ call, so it was
# not executed until the Tk mainloop closed.  V3.9 moves the __main__ call to the very
# end of the file and replaces on-text overlays with a real sibling gutter.
APP_VERSION = "3.9"

_V39_GUTTER_W = 46
_V39_GUTTER_BG = '#f7f9fc'
_V39_HOVER_BG = '#fff3b0'
_V39_DOC_HOVER_BG = '#fff7b2'


def _v39_collect_placements(self, row, i):
    placements=[]
    for ev in row.get('markers',[]) or []:
        eps=ev.get('endpoints')
        if eps:
            for ep in eps:
                try: target=int(ep.get('target_doc',-1))
                except Exception: target=-1
                if target==i:
                    clone=dict(ev)
                    clone['target_doc']=i
                    clone['char_start']=int(ep.get('char_start',0) or 0)
                    clone['char_end']=int(ep.get('char_end',clone['char_start']) or clone['char_start'])
                    clone['_overlay_role']=_v37_marker_role(self,ev,i)
                    clone['_row_id']=row['id']
                    placements.append(clone)
        else:
            try: target=int(ev.get('target_doc',-1))
            except Exception: target=-1
            if target==i:
                clone=dict(ev)
                clone['_overlay_role']=_v37_marker_role(self,ev,i)
                clone['_row_id']=row['id']
                clone['char_start']=int(clone.get('char_start',0) or 0)
                clone['char_end']=int(clone.get('char_end',clone['char_start']) or clone['char_start'])
                placements.append(clone)
    placements.sort(key=lambda e:(int(e.get('char_start',0)),int(e.get('num',0))))
    return placements


def _v39_safe_range(widget, start, end):
    try:
        start=max(0,int(start)); end=max(start,int(end))
        s=f'1.0+{start}c'; e=f'1.0+{end}c'
        # Position-only marker: give hover a one-character target when possible.
        if end<=start:
            last=int(widget.count('1.0','end-1c','chars')[0])
            if start < last:
                e=f'1.0+{start+1}c'
            elif start>0:
                s=f'1.0+{start-1}c'
        return s,e
    except Exception:
        return None,None


def _v39_summary_hover(self,row_id:int,num:int,on:bool=True):
    t=getattr(self,'_v31_summary_widgets',{}).get(row_id)
    if not t: return
    tag=f'marker_line_{num}'
    try:
        ranges=t.tag_ranges(tag)
        if not ranges: return
        t.configure(state='normal')
        if on:
            t.tag_configure('summary_hover',background=_V39_HOVER_BG,underline=1)
            t.tag_add('summary_hover',ranges[0],ranges[-1])
            t.tag_raise('summary_hover')
            t.see(ranges[0])
        else:
            t.tag_remove('summary_hover',ranges[0],ranges[-1])
        t.configure(state='disabled')
    except Exception:
        logging.exception('summary hover failed')


def _v39_doc_hover(self,row_id:int,num:int,on:bool=True):
    ev=getattr(self,'_v31_event_map',{}).get((row_id,num))
    if not ev: return
    eps=ev.get('endpoints') or [{'target_doc':ev.get('target_doc'),'char_start':ev.get('char_start'),'char_end':ev.get('char_end')}]
    for ep in eps:
        t=getattr(self,'_v31_doc_widgets',{}).get((row_id,ep.get('target_doc')))
        if not t: continue
        s,e=_v39_safe_range(t,ep.get('char_start',0),ep.get('char_end',0))
        if not s: continue
        try:
            t.configure(state='normal')
            if on:
                t.tag_configure('marker_hover',background=_V39_DOC_HOVER_BG)
                t.tag_add('marker_hover',s,e)
                t.tag_raise('marker_hover')
                t.see(s)
            else:
                t.tag_remove('marker_hover',s,e)
            t.configure(state='disabled')
        except Exception:
            logging.exception('document hover failed')


def _v39_hover_from_doc(self,row_id:int,num:int,on:bool=True):
    _v39_summary_hover(self,row_id,num,on)


def _v39_hover_from_summary(self,row_id:int,num:int,on:bool=True):
    _v39_doc_hover(self,row_id,num,on)


def _v39_bind_text_event(self,t,row_id:int,ev,serial:int):
    num=int(ev.get('num',0) or 0)
    if num<=0: return
    s,e=_v39_safe_range(t,ev.get('char_start',0),ev.get('char_end',0))
    if not s: return
    tag=f'v39_change_{row_id}_{num}_{serial}'
    try:
        t.tag_add(tag,s,e)
        # No visual options on this tag: red strike / blue underline remains authoritative.
        t.tag_bind(tag,'<Enter>',lambda _e,r=row_id,n=num:_v39_hover_from_doc(self,r,n,True))
        t.tag_bind(tag,'<Leave>',lambda _e,r=row_id,n=num:_v39_hover_from_doc(self,r,n,False))
        t.tag_bind(tag,'<Button-1>',lambda _e,r=row_id,n=num:self._v31_focus_summary(r,n))
    except Exception:
        logging.exception('failed to bind hover tag')


def _v39_role_color(role:str):
    if role=='delete': return '#c62828'
    if role=='insert': return '#1565c0'
    return '#6a1b9a'


def _v39_draw_gutter(self,t,gutter,row_id:int,placements):
    try:
        if not t.winfo_exists() or not gutter.winfo_exists(): return
        t.update_idletasks(); gutter.update_idletasks(); gutter.delete('all')
    except Exception:
        return
    raw=[]
    for ev in placements:
        res=_v37_bbox_for_offset(t,ev.get('char_start',0))
        if not res: continue
        (x,y,w,h),_at_end=res
        role=ev.get('_overlay_role') or _v37_marker_role(self,ev,ev.get('target_doc',0))
        raw.append((int(y),int(h),ev,role))
    raw.sort(key=lambda z:(z[0],int(z[2].get('num',0))))
    groups=[]
    for y,h,ev,role in raw:
        if groups and abs(groups[-1]['y']-y)<=3:
            groups[-1]['items'].append((ev,role)); groups[-1]['h']=max(groups[-1]['h'],h)
        else:
            groups.append({'y':y,'h':h,'items':[(ev,role)]})
    for gi,grp in enumerate(groups):
        nums=[]; roles=[]
        for ev,role in grp['items']:
            n=int(ev.get('num',0) or 0)
            if n and n not in nums: nums.append(n)
            roles.append(role)
        if not nums: continue
        nums.sort()
        display='·'.join(str(n) for n in nums)
        if len(display)>7:
            display=f'{nums[0]}–{nums[-1]}' if len(nums)>1 else str(nums[0])
        role=roles[0] if roles and all(r==roles[0] for r in roles) else 'change'
        color=_v39_role_color(role)
        cy=max(9,grp['y']+max(9,grp['h']//2))
        tag=f'v39g_{row_id}_{gi}'
        # A real rail: all graphics live on a sibling Canvas, never over the Text widget.
        gutter.create_rectangle(3,cy-8,_V39_GUTTER_W-7,cy+8,fill='#ffffff',outline=color,width=1,tags=(tag,))
        gutter.create_text((_V39_GUTTER_W-4)//2,cy,text=display,fill=color,font=('Malgun Gothic',9,'bold'),tags=(tag,))
        gutter.create_line(_V39_GUTTER_W-7,cy,_V39_GUTTER_W-1,cy,fill=color,width=2,tags=(tag,))
        ns=tuple(nums)
        def enter(_e=None,r=row_id,vals=ns):
            for n in vals:
                _v39_summary_hover(self,r,n,True); _v39_doc_hover(self,r,n,True)
        def leave(_e=None,r=row_id,vals=ns):
            for n in vals:
                _v39_summary_hover(self,r,n,False); _v39_doc_hover(self,r,n,False)
        def click(_e=None,r=row_id,vals=ns):
            if vals:
                self._v31_focus_summary(r,vals[0])
        gutter.tag_bind(tag,'<Enter>',enter)
        gutter.tag_bind(tag,'<Leave>',leave)
        gutter.tag_bind(tag,'<Button-1>',click)
    try:
        gutter.configure(scrollregion=gutter.bbox('all'))
    except Exception: pass


def _v39_schedule_gutter(self,t,gutter,row_id,placements,delay=45):
    jobs=getattr(self,'_v39_gutter_jobs',{})
    key=(t,gutter)
    old=jobs.get(key)
    if old:
        try: t.after_cancel(old)
        except Exception: pass
    try:
        jobs[key]=t.after(delay,lambda:_v39_draw_gutter(self,t,gutter,row_id,placements))
    except Exception:
        return
    self._v39_gutter_jobs=jobs


def _v39_make_doc_cell(self,parent,row,i,h):
    frame=self.tk.Frame(parent,bg='#ffffff',bd=0,highlightthickness=0)
    frame.grid_columnconfigure(0,minsize=_V39_GUTTER_W,weight=0)
    frame.grid_columnconfigure(1,weight=1)
    frame.grid_rowconfigure(0,weight=1)
    gutter=self.tk.Canvas(frame,width=_V39_GUTTER_W,bg=_V39_GUTTER_BG,highlightthickness=0,bd=0,takefocus=0,cursor='arrow')
    gutter.grid(row=0,column=0,sticky='ns')
    gutter.bind('<MouseWheel>',self._child_wheel)
    t=self._new_text(frame,h)
    t.grid(row=0,column=1,sticky='nsew')
    m=row['members'][i]
    if not m:
        t.insert('end','[해당 조항 없음]','missing')
    else:
        hs=row.get('header_segments',[[] for _ in row['segments']])[i] if row.get('header_segments') else []
        if hs:
            for txt,sty in hs:
                t.insert('end',txt,sty if sty in ('delete','insert') else 'article_header')
            t.insert('end','\n')
        else:
            t.insert('end',m.get('header','')+'\n','article_header')
        for txt,sty in row.get('body_segments',row['segments'])[i]:
            t.insert('end',txt,sty if sty in ('delete','insert') else ())
    self._v31_doc_widgets[(row['id'],i)]=t
    placements=_v39_collect_placements(self,row,i)
    for serial,ev in enumerate(placements):
        _v39_bind_text_event(self,t,row['id'],ev,serial)
    t.configure(state='disabled')
    if placements:
        self._v39_gutters[(row['id'],i)]=(gutter,t,placements)
        t.bind('<Configure>',lambda _e,tw=t,g=gutter,r=row['id'],ps=tuple(placements):_v39_schedule_gutter(self,tw,g,r,ps),add='+')
        t.bind('<Map>',lambda _e,tw=t,g=gutter,r=row['id'],ps=tuple(placements):_v39_schedule_gutter(self,tw,g,r,ps,20),add='+')
        frame.bind('<Configure>',lambda _e,tw=t,g=gutter,r=row['id'],ps=tuple(placements):_v39_schedule_gutter(self,tw,g,r,ps),add='+')
        _v39_schedule_gutter(self,t,gutter,row['id'],tuple(placements),80)
    return frame


def _v39_make_summary_cell(self,parent,row,h):
    t=self._new_text(parent,h,'#ffffff')
    self._v31_summary_widgets[row['id']]=t
    events={e['num']:e for e in row.get('markers',[]) or []}
    structural=_v31_structural_messages((row.get('summary') or {}).get('messages') or [])
    if not events and not structural:
        t.insert('end','변경 없음','conf')
    else:
        for num in sorted(events):
            ev=events[num]
            start=t.index('end-1c')
            prefix=f"{_v31_circled(num)} {ev['label']} · "
            t.insert('end',prefix,'marker_num')
            t.insert('end',ev['message']+'\n')
            end=t.index('end-1c')
            tag=f'marker_line_{num}'
            t.tag_add(tag,start,end)
            t.tag_bind(tag,'<Button-1>',lambda _e,r=row['id'],n=num:self._v31_focus_marker(r,n))
            t.tag_bind(tag,'<Enter>',lambda _e,r=row['id'],n=num,tw=t:(_v39_hover_from_summary(self,r,n,True),tw.configure(cursor='hand2')))
            t.tag_bind(tag,'<Leave>',lambda _e,r=row['id'],n=num,tw=t:(_v39_hover_from_summary(self,r,n,False),tw.configure(cursor='arrow')))
        for msg in structural:
            t.insert('end','• '+msg+'\n')
    t.tag_configure('marker_num',foreground='#245ea8',font=('Malgun Gothic',10,'bold'))
    t.configure(state='disabled')
    return t


_v39_render_rows_base = NativeGui.render_rows

def _v39_render_rows(self):
    self._v39_gutters={}
    self._v39_gutter_jobs={}
    result=_v39_render_rows_base(self)
    def redraw():
        for (_rid,_i),(g,t,ps) in list(getattr(self,'_v39_gutters',{}).items()):
            try: _v39_draw_gutter(self,t,g,_rid,ps)
            except Exception: pass
    try: self.root.after(140,redraw)
    except Exception: pass
    return result


NativeGui._make_doc_cell = _v39_make_doc_cell
NativeGui._make_summary_cell = _v39_make_summary_cell
NativeGui.render_rows = _v39_render_rows
NativeGui._v39_summary_hover = _v39_summary_hover
NativeGui._v39_doc_hover = _v39_doc_hover



# ---------------- V4.0: inline change-number buttons ----------------
# The separate gutter used in V3.9 is removed.  Each change number is now an embedded
# Tk button placed at the exact text boundary.  Embedded windows participate in Text
# layout, so the source text moves aside instead of being obscured.  Hovering the
# red/blue changed text itself still highlights the matching summary entry.
APP_VERSION = "4.0"

_V40_DELETE_BG = '#c62828'
_V40_INSERT_BG = '#1565c0'
_V40_CHANGE_BG = '#6a1b9a'
_V40_BUTTON_FG = '#ffffff'


def _v40_role_color(role: str) -> str:
    if role == 'delete': return _V40_DELETE_BG
    if role == 'insert': return _V40_INSERT_BG
    return _V40_CHANGE_BG


def _v40_group_placements(placements):
    """Group multiple marker numbers that start at the same source character offset.

    One embedded window is inserted per source offset.  If several logical changes share
    that exact boundary, their small buttons are packed side-by-side inside one frame.
    This keeps original-offset -> rendered-index mapping deterministic (one window char
    per boundary, not one per number).
    """
    groups={}
    for ev in placements or []:
        pos=max(0,int(ev.get('char_start',0) or 0))
        groups.setdefault(pos,[]).append(ev)
    out=[]
    for pos,evs in groups.items():
        evs=sorted(evs,key=lambda e:int(e.get('num',0) or 0))
        out.append((pos,evs))
    return sorted(out,key=lambda x:x[0])


def _v40_display_offset(self,row_id:int,doc_index:int,original_offset:int,for_end:bool=False) -> int:
    """Translate a source-text character offset after inline marker windows were inserted."""
    original_offset=max(0,int(original_offset or 0))
    positions=(getattr(self,'_v40_insert_positions',{}) or {}).get((row_id,doc_index),())
    if for_end:
        # A marker inserted exactly at the end boundary belongs to the following text and
        # should not be included in this range.
        extra=sum(1 for p in positions if p < original_offset)
    else:
        # A marker at the start boundary sits immediately BEFORE the changed text.
        extra=sum(1 for p in positions if p <= original_offset)
    return original_offset + extra


def _v40_safe_range(self,widget,row_id:int,doc_index:int,start:int,end:int):
    try:
        start=max(0,int(start or 0)); end=max(start,int(end or start))
        ds=_v40_display_offset(self,row_id,doc_index,start,False)
        de=_v40_display_offset(self,row_id,doc_index,end,True)
        # Position-only changes still need a one-character hover target where possible.
        if de<=ds:
            last=int(widget.count('1.0','end-1c','chars')[0])
            if ds < last: de=ds+1
            elif ds>0: ds-=1
        return f'1.0+{ds}c',f'1.0+{de}c'
    except Exception:
        return None,None


def _v40_inline_button_frame(self,t,row_id:int,doc_index:int,evs):
    frame=self.tk.Frame(t,bg=t.cget('bg'),bd=0,highlightthickness=0,takefocus=0)
    for ev in evs:
        num=int(ev.get('num',0) or 0)
        role=ev.get('_overlay_role') or _v37_marker_role(self,ev,doc_index)
        bg=_v40_role_color(role)
        # A readable 18-22px-ish button that sits inside the line flow.  Text is never
        # covered because Text.window_create allocates real horizontal space for it.
        b=self.tk.Button(
            frame,text=str(num),font=('Malgun Gothic',9,'bold'),fg=_V40_BUTTON_FG,bg=bg,
            activeforeground=_V40_BUTTON_FG,activebackground=bg,bd=1,relief='raised',
            padx=3,pady=0,cursor='hand2',takefocus=0,highlightthickness=0
        )
        b.pack(side='left',padx=(0,2))
        b.bind('<Enter>',lambda _e,r=row_id,n=num:(_v39_summary_hover(self,r,n,True),_v40_doc_hover(self,r,n,True)))
        b.bind('<Leave>',lambda _e,r=row_id,n=num:(_v39_summary_hover(self,r,n,False),_v40_doc_hover(self,r,n,False)))
        b.bind('<Button-1>',lambda _e,r=row_id,n=num:self._v31_focus_summary(r,n))
        b.bind('<MouseWheel>',self._child_wheel)
    return frame


def _v40_insert_inline_buttons(self,t,row_id:int,doc_index:int,placements):
    groups=_v40_group_placements(placements)
    self._v40_insert_positions[(row_id,doc_index)]=tuple(pos for pos,_ in groups)
    # Insert from the end toward the beginning so every source offset remains valid while
    # windows are being added.  The window itself occupies one Text index.
    for pos,evs in reversed(groups):
        try:
            frame=_v40_inline_button_frame(self,t,row_id,doc_index,evs)
            t.window_create(f'1.0+{pos}c',window=frame,padx=2,pady=1,align='center')
            self._v40_inline_widgets.append(frame)
        except Exception:
            logging.exception('inline marker button insertion failed')


def _v40_bind_text_event(self,t,row_id:int,doc_index:int,ev,serial:int):
    num=int(ev.get('num',0) or 0)
    if num<=0: return
    s,e=_v40_safe_range(self,t,row_id,doc_index,ev.get('char_start',0),ev.get('char_end',0))
    if not s: return
    tag=f'v40_change_{row_id}_{doc_index}_{num}_{serial}'
    try:
        t.tag_add(tag,s,e)
        # The style remains red strike / blue underline.  This tag only supplies linkage.
        t.tag_bind(tag,'<Enter>',lambda _e,r=row_id,n=num:_v39_summary_hover(self,r,n,True))
        t.tag_bind(tag,'<Leave>',lambda _e,r=row_id,n=num:_v39_summary_hover(self,r,n,False))
        t.tag_bind(tag,'<Button-1>',lambda _e,r=row_id,n=num:self._v31_focus_summary(r,n))
    except Exception:
        logging.exception('inline change hover binding failed')


def _v40_doc_hover(self,row_id:int,num:int,on:bool=True):
    ev=getattr(self,'_v31_event_map',{}).get((row_id,num))
    if not ev: return
    eps=ev.get('endpoints') or [{'target_doc':ev.get('target_doc'),'char_start':ev.get('char_start'),'char_end':ev.get('char_end')}]
    for ep in eps:
        try: di=int(ep.get('target_doc'))
        except Exception: continue
        t=getattr(self,'_v31_doc_widgets',{}).get((row_id,di))
        if not t: continue
        s,e=_v40_safe_range(self,t,row_id,di,ep.get('char_start',0),ep.get('char_end',0))
        if not s: continue
        try:
            t.configure(state='normal')
            if on:
                t.tag_configure('marker_hover',background=_V39_DOC_HOVER_BG)
                t.tag_add('marker_hover',s,e); t.tag_raise('marker_hover'); t.see(s)
            else:
                t.tag_remove('marker_hover',s,e)
            t.configure(state='disabled')
        except Exception:
            logging.exception('inline document hover failed')


def _v40_focus_marker(self,row_id:int,num:int):
    ev=getattr(self,'_v31_event_map',{}).get((row_id,num))
    if not ev: return
    eps=ev.get('endpoints') or [{'target_doc':ev.get('target_doc'),'char_start':ev.get('char_start'),'char_end':ev.get('char_end')}]
    for ep in eps:
        try: di=int(ep.get('target_doc'))
        except Exception: continue
        t=getattr(self,'_v31_doc_widgets',{}).get((row_id,di))
        if not t: continue
        s,e=_v40_safe_range(self,t,row_id,di,ep.get('char_start',0),ep.get('char_end',0))
        if not s: continue
        try:
            t.see(s); t.configure(state='normal')
            t.tag_configure('marker_focus',background='#fff59d')
            t.tag_add('marker_focus',s,e); t.tag_raise('marker_focus')
            t.configure(state='disabled')
            def clear(widget=t,ss=s,ee=e):
                try:
                    widget.configure(state='normal'); widget.tag_remove('marker_focus',ss,ee); widget.configure(state='disabled')
                except Exception: pass
            t.after(1100,clear)
        except Exception:
            logging.exception('inline focus marker failed')


def _v40_make_doc_cell(self,parent,row,i,h):
    # No gutter frame: the Text widget gets the full document-column width again.
    t=self._new_text(parent,h)
    m=row['members'][i]
    if not m:
        t.insert('end','[해당 조항 없음]','missing')
    else:
        hs=row.get('header_segments',[[] for _ in row['segments']])[i] if row.get('header_segments') else []
        if hs:
            for txt,sty in hs:
                t.insert('end',txt,sty if sty in ('delete','insert') else 'article_header')
            t.insert('end','\n')
        else:
            t.insert('end',m.get('header','')+'\n','article_header')
        for txt,sty in row.get('body_segments',row['segments'])[i]:
            t.insert('end',txt,sty if sty in ('delete','insert') else ())

    self._v31_doc_widgets[(row['id'],i)]=t
    placements=_v39_collect_placements(self,row,i)
    # Windows must be inserted while Text is editable.  Then translated ranges are bound.
    if placements:
        _v40_insert_inline_buttons(self,t,row['id'],i,placements)
        for serial,ev in enumerate(placements):
            _v40_bind_text_event(self,t,row['id'],i,ev,serial)
    else:
        self._v40_insert_positions[(row['id'],i)]=()
    t.configure(state='disabled')
    return t


def _v40_hover_from_summary(self,row_id:int,num:int,on:bool=True):
    _v40_doc_hover(self,row_id,num,on)


def _v40_make_summary_cell(self,parent,row,h):
    t=self._new_text(parent,h,'#ffffff')
    self._v31_summary_widgets[row['id']]=t
    events={e['num']:e for e in row.get('markers',[]) or []}
    structural=_v31_structural_messages((row.get('summary') or {}).get('messages') or [])
    if not events and not structural:
        t.insert('end','변경 없음','conf')
    else:
        for num in sorted(events):
            ev=events[num]; start=t.index('end-1c')
            prefix=f"{_v31_circled(num)} {ev['label']} · "
            t.insert('end',prefix,'marker_num'); t.insert('end',ev['message']+'\n')
            end=t.index('end-1c'); tag=f'marker_line_{num}'
            t.tag_add(tag,start,end)
            t.tag_bind(tag,'<Button-1>',lambda _e,r=row['id'],n=num:_v40_focus_marker(self,r,n))
            t.tag_bind(tag,'<Enter>',lambda _e,r=row['id'],n=num,tw=t:(_v40_hover_from_summary(self,r,n,True),tw.configure(cursor='hand2')))
            t.tag_bind(tag,'<Leave>',lambda _e,r=row['id'],n=num,tw=t:(_v40_hover_from_summary(self,r,n,False),tw.configure(cursor='arrow')))
        for msg in structural: t.insert('end','• '+msg+'\n')
    t.tag_configure('marker_num',foreground='#245ea8',font=('Malgun Gothic',10,'bold'))
    t.configure(state='disabled')
    return t


_v40_render_rows_base = NativeGui.render_rows

def _v40_render_rows(self):
    # Dispose references from the previous render and initialise offset maps BEFORE cells
    # are created.  Embedded window widgets are children of the Text and are destroyed with it.
    self._v40_insert_positions={}
    self._v40_inline_widgets=[]
    self._v39_gutters={}
    self._v39_gutter_jobs={}
    return _v40_render_rows_base(self)


NativeGui._make_doc_cell = _v40_make_doc_cell
NativeGui._make_summary_cell = _v40_make_summary_cell
NativeGui.render_rows = _v40_render_rows
NativeGui._v31_focus_marker = _v40_focus_marker
NativeGui._v39_doc_hover = _v40_doc_hover



# ---------------- V4.1: marker buttons at the visual start of the changed word ----------------
# V4.0 inserted an inline button at the exact character offset returned by the token diff.
# For suffix-level edits in Korean/English that can land in the middle of a word, e.g.
#   설[3]치된 / 정상[4]적인
# which is technically exact but visually misleading.  Keep the true red/blue diff span
# unchanged, but move only the inline marker button to the beginning of the containing word:
#   [3]설치된 / [4]정상적인
# Punctuation-only edits remain anchored at the exact punctuation boundary.
APP_VERSION = "4.5"


def _v41_is_word_char(ch: str) -> bool:
    return bool(ch) and (ch.isalnum() or ch == '_')


def _v41_marker_anchor(text: str, pos: int) -> int:
    """Return a reader-friendly inline-button anchor without changing the actual diff range."""
    text=text or ''
    pos=max(0,min(int(pos or 0),len(text)))
    # An end-of-text insertion has no following word to expand into.
    if pos>=len(text):
        return pos
    # Only expand when the changed position itself is inside a word.  Quotes, commas,
    # periods, brackets, etc. retain their exact location as requested for symbol edits.
    if not _v41_is_word_char(text[pos]):
        return pos
    start=pos
    while start>0 and _v41_is_word_char(text[start-1]):
        start-=1
    return start


def _v41_group_placements(placements):
    groups={}
    for ev in placements or []:
        pos=max(0,int(ev.get('_inline_anchor',ev.get('char_start',0)) or 0))
        groups.setdefault(pos,[]).append(ev)
    out=[]
    for pos,evs in groups.items():
        evs=sorted(evs,key=lambda e:int(e.get('num',0) or 0))
        out.append((pos,evs))
    return sorted(out,key=lambda x:x[0])


def _v41_insert_inline_buttons(self,t,row_id:int,doc_index:int,placements):
    # Compute anchors from the pristine Text contents before any embedded windows exist.
    try:
        raw=t.get('1.0','end-1c')
    except Exception:
        raw=''
    prepared=[]
    for ev0 in placements or []:
        ev=dict(ev0)
        raw_pos=max(0,int(ev.get('char_start',0) or 0))
        ev['_inline_anchor']=_v41_marker_anchor(raw,raw_pos)
        prepared.append(ev)
    groups=_v41_group_placements(prepared)
    self._v40_insert_positions[(row_id,doc_index)]=tuple(pos for pos,_ in groups)
    for pos,evs in reversed(groups):
        try:
            frame=_v40_inline_button_frame(self,t,row_id,doc_index,evs)
            t.window_create(f'1.0+{pos}c',window=frame,padx=2,pady=1,align='center')
            self._v40_inline_widgets.append(frame)
        except Exception:
            logging.exception('V4.1 inline marker button insertion failed')


# Only the visual button anchor changes.  Hover/focus ranges still use the original
# char_start/char_end values, translated through the embedded-window positions.
_v40_insert_inline_buttons = _v41_insert_inline_buttons



# ---------------- V4.3: general-purpose document profiles ----------------
# The legal parser remains a specialized profile.  Auto mode selects it only when the
# document contains strong Article/제N조 structure; otherwise a generic heading/paragraph
# parser is used.  This keeps the current legal-document accuracy while allowing the same
# comparison engine to evolve into a general document comparer.
ACTIVE_COMPARE_MODE = 'auto'
_v42_legal_parse_units = parse_units

_V42_GENERIC_NUMBERED_HEADING = re.compile(
    r'^\s*((?:\d+(?:\.\d+){0,5})|(?:[IVXLCDM]+)|(?:[A-Z]))[.)]?\s+(.{1,140})$', re.I
)

def _v42_generic_heading(line: str):
    s=(line or '').strip()
    if not s or len(s)>180:
        return None
    m=_V42_GENERIC_NUMBERED_HEADING.match(s)
    if m:
        title=m.group(2).strip()
        # Long sentence-like numbered list items are not promoted to headings.
        if len(title)<=110 and not re.search(r'[.!?;]\s*$',title):
            return (m.group(1),s)
    # Conservative unnumbered heading heuristic for ordinary reports/memos.
    words=re.findall(r'[A-Za-z가-힣0-9]+',s)
    if 1 <= len(words) <= 10 and len(s)<=80 and not re.search(r'[.!?;:]\s*$',s):
        if s.isupper() or (re.search(r'[A-Za-z]',s) and sum(1 for w in words if w[:1].isupper()) >= max(1,len(words)//2)):
            return ('',s)
    return None

def _v42_generic_parse_units(text: str) -> List[Unit]:
    text=(text or '').replace('\r\n','\n').replace('\r','\n')
    paras=[x.strip() for x in re.split(r'\n\s*\n|\n',text) if x.strip()]
    if not paras:
        return []
    units=[]; current_header=None; body=[]; seq=0
    def flush():
        nonlocal seq,current_header,body
        if current_header is None and not body:
            return
        seq+=1
        if current_header:
            num,header=current_header; b='\n'.join(body).strip(); title=re.sub(r'^\s*'+re.escape(num)+r'[.)]?\s*','',header,1) if num else header
            full=header if not b else header+'\n'+b
            units.append(Unit(len(units),'article',num or f'g{seq}',title,header,b,full,'',normalize_text(title),normalize_text(b),structure_sig(b)))
        else:
            # A run of ordinary paragraphs is kept as one logical block so paragraph
            # insertion/deletion can be aligned without inventing legal article numbers.
            b='\n'.join(body).strip(); header=''
            units.append(Unit(len(units),'article',f'g{seq}','',header,b,b,'','',normalize_text(b),structure_sig(b)))
        current_header=None; body=[]
    for p in paras:
        h=_v42_generic_heading(p)
        if h:
            flush(); current_header=h
        else:
            if current_header is None and body:
                # Keep ordinary prose paragraphs individually addressable.
                flush()
            body.append(p)
    flush()
    return units

def _v42_looks_legal(text: str) -> bool:
    t=text or ''
    try:
        if len(_collect_article_markers_v20(t)) >= 2:
            return True
    except Exception:
        pass
    return len(re.findall(r'(?im)^\s*(?:Article\s+\d+(?:[-.]\d+)*|제\s*\d+\s*조)',t)) >= 2

def parse_units(text: str) -> List[Unit]:
    mode=globals().get('ACTIVE_COMPARE_MODE','auto')
    if mode=='legal':
        return _v42_legal_parse_units(text)
    if mode=='general':
        return _v42_generic_parse_units(text)
    return _v42_legal_parse_units(text) if _v42_looks_legal(text) else _v42_generic_parse_units(text)


# ---------------- V4.4: Excel export by paragraph/item instead of whole article ----------------
# GUI keeps article-level synchronized rows, but Excel is a review/report artifact.  A whole
# article in one cell becomes unreadable, so the exporter expands each matched article into a
# slim article-title row followed by one row per explicit item/paragraph ((1), 1., ①, etc.).
APP_VERSION = "4.5"

def _v44_excel_parts(member):
    if not member:
        return []
    body=(member.get('body') if isinstance(member,dict) else getattr(member,'body','')) or ''
    body=body.replace('\r\n','\n').replace('\r','\n').strip()
    if not body:
        return []
    # Legal/regulatory numbered items have exact offsets and are the preferred split.
    try:
        parts=_v35_parts_with_spans(body)
    except Exception:
        parts=[]
    explicit=bool(parts) and (len(parts)>1 or any((x.get('label') or '')!='본문' for x in parts))
    if explicit:
        out=[]
        for x in parts:
            label=(x.get('label') or '').strip(); core=(x.get('core') or '').strip()
            if not core:
                continue
            text=core if label=='본문' or not label else f'{label} {core}'
            out.append({'label':label or '본문','core':core,'text':text})
        return out
    # Generic documents: preserve real paragraph/line boundaries rather than inventing legal items.
    lines=[re.sub(r'[ \t]+',' ',x).strip() for x in body.split('\n') if x.strip()]
    if len(lines)>1:
        return [{'label':'문단','core':x,'text':x} for x in lines]
    return [{'label':'본문','core':body,'text':body}]

def _v44_part_unit(text, label=''):
    text=(text or '').strip()
    return Unit(0,'article',label or '','', '', text, text, '', '', normalize_text(text), structure_sig(text))

def _v44_match_parts(ap, bp):
    used_a=set(); used_b=set(); pairs=[]
    # 1) Exact normalized content, independent of numbering/location.
    idx=defaultdict(list)
    for j,b in enumerate(bp):
        n=normalize_text(b.get('core',''))
        if n: idx[n].append(j)
    for i,a in enumerate(ap):
        n=normalize_text(a.get('core',''))
        if not n: continue
        cand=[j for j in idx.get(n,[]) if j not in used_b]
        if cand:
            j=min(cand,key=lambda x:abs(i-x)); pairs.append((i,j,1.0)); used_a.add(i); used_b.add(j)
    # 2) Same item number/label when the texts remain recognizably related.  This prevents
    # a heavily edited (1) from being paired with an unrelated (4) merely because both
    # contain generic words such as Company/Service.
    same=[]
    for i,a in enumerate(ap):
        if i in used_a or a.get('label') in ('','본문','문단'): continue
        for j,b in enumerate(bp):
            if j in used_b or b.get('label')!=a.get('label'): continue
            ss=_v35_part_similarity(a,b)
            if ss>=0.38: same.append((ss,i,j))
    for ss,i,j in sorted(same,reverse=True):
        if i in used_a or j in used_b: continue
        pairs.append((i,j,ss)); used_a.add(i); used_b.add(j)
    # 3) Strong content match for moved or renumbered items.
    fuzzy=[]
    for i,a in enumerate(ap):
        if i in used_a: continue
        for j,b in enumerate(bp):
            if j in used_b: continue
            ss=_v35_part_similarity(a,b)
            if ss>=0.45: fuzzy.append((ss,i,j))
    for ss,i,j in sorted(fuzzy,reverse=True):
        if i in used_a or j in used_b: continue
        pairs.append((i,j,ss)); used_a.add(i); used_b.add(j)
    return pairs,used_a,used_b

def _v44_align_parts(members, base_index):
    """Return Excel child rows, preserving the selected base document's item order.

    Each returned item is a list with one part (or None) per document.  Unmatched additions are
    inserted near the closest matched neighbour instead of being appended as one giant article.
    """
    n=len(members); all_parts=[_v44_excel_parts(m) for m in members]
    anchor=base_index if base_index<n and all_parts[base_index] else next((i for i,p in enumerate(all_parts) if p),0)
    ap=all_parts[anchor]
    if not ap:
        return []
    entries=[{'parts':[None]*n,'anchor_pos':i,'insert':False} for i in range(len(ap))]
    for i,p in enumerate(ap): entries[i]['parts'][anchor]=p
    mappings={anchor:{i:i for i in range(len(ap))}}
    unmatched_by_doc={}
    for d in range(n):
        if d==anchor: continue
        bp=all_parts[d]
        if not bp:
            mappings[d]={}; unmatched_by_doc[d]=list(range(0)); continue
        # Excel alignment is slightly stricter than the GUI marker matcher:
        # 1) exact content may move/renumber, 2) same enumerator wins when still plausibly
        # related, 3) strong content similarity can then connect moved/renumbered items.
        pairs,ua,ub=_v44_match_parts(ap,bp)
        mp={}
        for ai,bj,score in pairs:
            if 0<=ai<len(entries) and 0<=bj<len(bp):
                entries[ai]['parts'][d]=bp[bj]; mp[bj]=ai
        mappings[d]=mp
        unmatched_by_doc[d]=[j for j in range(len(bp)) if j not in ub]
    # Add items that exist only in a comparison document.  Position them after the nearest
    # preceding matched base item (or before the first base item when appropriate).
    buckets=defaultdict(list)
    for d,unmatched in unmatched_by_doc.items():
        bp=all_parts[d]; mp=mappings.get(d,{})
        matched_js=sorted(mp)
        for j in unmatched:
            prev=[x for x in matched_js if x<j]
            after=mp[max(prev)] if prev else -1
            buckets[after].append((d,j,bp[j]))
    # Merge same/near-identical inserted items from B/C into one row when possible.
    def merged_bucket(items):
        rows=[]
        for d,j,p in sorted(items,key=lambda x:(x[1],x[0])):
            placed=False
            for rr in rows:
                exemplar=next((q for q in rr['parts'] if q),None)
                if not exemplar: continue
                na=normalize_text(exemplar.get('core','')); nb=normalize_text(p.get('core',''))
                sim=SequenceMatcher(None,na,nb,autojunk=False).ratio() if na and nb else 0
                if (na and na==nb) or sim>=0.78:
                    rr['parts'][d]=p; placed=True; break
            if not placed:
                rr={'parts':[None]*n,'insert':True}; rr['parts'][d]=p; rows.append(rr)
        return rows
    out=[]
    out.extend(merged_bucket(buckets.get(-1,[])))
    for i,e in enumerate(entries):
        out.append(e)
        out.extend(merged_bucket(buckets.get(i,[])))
    return out

def _v44_segments_for_parts(parts, base_index):
    units=[_v44_part_unit(p.get('text',''),p.get('label','')) if p else None for p in parts]
    return directional_segments_v19(units,base_index,'body')

def _v44_part_messages(parts, base_index):
    n=len(parts); base=parts[base_index] if base_index<n else None; lines=[]
    if base:
        bu=_v44_part_unit(base.get('text',''),base.get('label',''))
        for d,p in enumerate(parts):
            if d==base_index: continue
            label=f'문서 {chr(65+d)}'
            if not p:
                lines.append(f'{label} · 삭제: “{base.get("text","")}”')
                continue
            ou=_v44_part_unit(p.get('text',''),p.get('label',''))
            evs=_v31_pair_marker_events(bu,ou,base_index,d)
            for e in evs:
                msg=e.get('message') or f'{e.get("action")}: “{e.get("text","")}”'
                line=f'{label} · {msg}'
                if line not in lines: lines.append(line)
    else:
        for d,p in enumerate(parts):
            if p:
                lines.append(f'문서 {chr(65+d)} · 추가: “{p.get("text","")}”')
    return lines or ['변경 없음']

def make_xlsx(result: Dict[str,Any]) -> bytes:
    out=io.BytesIO(); wb=xlsxwriter.Workbook(out, {'in_memory':True})
    ws=wb.add_worksheet('문서비교'); sw=wb.add_worksheet('요약')
    hf=wb.add_format({'bold':True,'bg_color':'#1F4E78','font_color':'#FFFFFF','align':'center','valign':'vcenter','border':1})
    cf=wb.add_format({'text_wrap':True,'valign':'top','border':1,'font_size':10})
    mf=wb.add_format({'text_wrap':True,'valign':'top','border':1,'font_color':'#777777','italic':True,'bg_color':'#F2F2F2'})
    af=wb.add_format({'bold':True,'bg_color':'#D9EAF7','font_color':'#17365D','valign':'vcenter','border':1,'text_wrap':True})
    acf=wb.add_format({'bg_color':'#D9EAF7','valign':'vcenter','border':1,'text_wrap':True})
    delf=wb.add_format({'font_color':'#C00000','font_strikeout':True})
    insf=wb.add_format({'font_color':'#1565C0','underline':True})
    delcell=wb.add_format({'text_wrap':True,'valign':'top','border':1,'font_color':'#C00000','font_strikeout':True})
    inscell=wb.add_format({'text_wrap':True,'valign':'top','border':1,'font_color':'#1565C0','underline':True})
    title=wb.add_format({'bold':True,'font_size':16,'font_color':'#1F4E78'})
    sl=wb.add_format({'bold':True,'bg_color':'#D9EAF7','border':1}); sv=wb.add_format({'border':1,'align':'center','text_wrap':True})
    bi=int(result.get('base_index',0) or 0); n=len(result.get('names',[])); include_ac=bool(result.get('include_ac',False))
    headers=[]
    for i,name in enumerate(result.get('names',[])):
        headers.append(f'{chr(65+i)} · {name}' + (' · 기준' if i==bi else ''))
    headers.append('변경사항')
    ws.write_row(0,0,headers,hf); ws.freeze_panes(1,0)
    if n==3:
        for c in range(3): ws.set_column(c,c,30)
        ws.set_column(3,3,48)
    else:
        for c in range(max(1,n)): ws.set_column(c,c,40)
        ws.set_column(n,n,50)

    def write_segments(r,c,segs):
        if not segs:
            ws.write(r,c,'[해당 항/문단 없음]',mf); return
        nonempty=[(txt,sty) for txt,sty in segs if txt]
        if not nonempty:
            ws.write(r,c,'',cf); return
        styles={sty for txt,sty in nonempty if txt.strip()}
        if styles=={'delete'}:
            ws.write(r,c,''.join(t for t,_ in nonempty),delcell); return
        if styles=={'insert'}:
            ws.write(r,c,''.join(t for t,_ in nonempty),inscell); return
        args=[]; rich=False
        for txt,sty in nonempty:
            if sty=='delete' and txt.strip(): args.extend([delf,txt]); rich=True
            elif sty=='insert' and txt.strip(): args.extend([insf,txt]); rich=True
            else: args.append(txt)
        if rich:
            try:
                ws.write_rich_string(r,c,*args,cf); return
            except Exception:
                pass
        ws.write(r,c,''.join(t for t,_ in nonempty),cf)

    xr=1
    for row in result.get('rows',[]):
        members=row.get('members') or [None]*n
        # 1) slim article/block title row.  This preserves context but does not swallow the body.
        any_header=any((m or {}).get('header','').strip() for m in members if m)
        if any_header:
            header_units=[_v44_part_unit((m or {}).get('header',''),'header') if m else None for m in members]
            hsegs=_v55_pairwise_segments(header_units,bi,'body',include_ac)
            for c in range(n):
                if members[c]: write_segments(xr,c,hsegs[c])
                else: ws.write(xr,c,'[해당 조/블록 없음]',mf)
            # Structural/article-level notes belong on the title row only.
            structural=_v31_structural_messages((row.get('summary') or {}).get('messages') or [])
            ws.write(xr,n,'\n'.join(structural) if structural else '',acf)
            ws.set_row(xr,22); xr+=1
        # 2) one Excel row per 항/호/paragraph.
        child_rows=_v44_align_parts(members,bi)
        if not child_rows and any(members):
            # Header-only unit.
            continue
        for child in child_rows:
            parts=child['parts']; segs=_v55_excel_segments_for_parts(parts,bi,include_ac)
            for c in range(n): write_segments(xr,c,segs[c])
            msgs=_v55_excel_part_messages(parts,bi,include_ac)
            ws.write(xr,n,'\n'.join(msgs),cf)
            maxchars=max([len((p or {}).get('text','')) for p in parts]+[len(' '.join(msgs))])
            ws.set_row(xr,min(180,max(30,18+12*(maxchars//75))))
            xr+=1
    ws.autofilter(0,0,max(1,xr-1),n)

    sw.write('A1','문서 비교 요약',title); sw.set_column('A:A',22); sw.set_column('B:B',72)
    sw.write('A3','기준 문서',sl); sw.write('B3',result['names'][bi] if result.get('names') else '',sv)
    sw.write('A4','비교 문서',sl); sw.write('B4',' / '.join(name for i,name in enumerate(result.get('names',[])) if i!=bi),sv)
    sw.write('A6','Excel 행 구성',sl); sw.write('B6','조/블록 제목은 얇은 구분행으로 표시하고, 본문은 항·호·번호목록·문단 단위로 한 행씩 분리합니다.',cf)
    sw.write('A7','표시 규칙',sl); sw.write('B7','삭제/대체된 기존 문구 = 붉은색 취소선\n신설/대체된 새 문구 = 파란색 밑줄',cf)
    sw.write('A8','정렬 규칙',sl); sw.write('B8','선택한 기준문서의 항/문단 순서를 우선 유지하며, 이동·재번호화된 항은 내용 계보로 같은 행에 대응시킵니다.',cf)
    wb.close(); out.seek(0); return out.read()

# Generic product naming for exported workbooks.
_v44_export_excel = NativeGui.export_excel
def _v44_export_excel_method(self):
    if not self.result: return
    p=self.filedialog.asksaveasfilename(title='Excel 비교 결과 저장',defaultextension='.xlsx',initialfile='문서_비교결과.xlsx',filetypes=[('Excel 통합 문서','*.xlsx')])
    if not p: return
    try:
        Path(p).write_bytes(make_xlsx(self.result)); self.status_var.set(f'Excel 저장 완료: {Path(p).name}'); self.messagebox.showinfo('저장 완료',f'Excel 파일을 저장했습니다.\n\n{p}')
    except Exception as e:
        logging.exception('excel export failed'); self.messagebox.showerror('저장 오류',str(e))
NativeGui.export_excel = _v44_export_excel_method



# ---------------- V4.5: performance pass + Avalonia bridge readiness ----------------
# 1) The expensive article alignment now uses RapidFuzz's native Indel similarity through
#    seq_ratio() above.  This leaves opcode-producing difflib calls intact, so the visual
#    red/blue diff behavior is unchanged.
# 2) DOCX/TXT extraction is cached by (absolute path, size, mtime_ns). Re-comparing the same
#    files (for example after changing the base document or output choice) no longer reparses
#    unchanged Word XML.
_V45_TEXT_CACHE: Dict[Tuple[str, int, int], str] = {}
_V45_TEXT_CACHE_ORDER: List[Tuple[str, int, int]] = []
_V45_TEXT_CACHE_MAX = 12


def _v45_read_path_cached(path: Path) -> Tuple[str, bool]:
    st = path.stat()
    key = (str(path.resolve()), int(st.st_size), int(st.st_mtime_ns))
    cached = _V45_TEXT_CACHE.get(key)
    if cached is not None:
        return cached, True
    raw = path.read_bytes()
    if len(raw) > 40 * 1024 * 1024:
        raise ValueError(f'{path.name}: 파일당 40MB까지 지원합니다.')
    text = read_document(path.name, raw)
    # Remove stale versions of the same path before inserting the new one.
    resolved = key[0]
    stale = [k for k in _V45_TEXT_CACHE_ORDER if k[0] == resolved and k != key]
    for k in stale:
        _V45_TEXT_CACHE.pop(k, None)
        try: _V45_TEXT_CACHE_ORDER.remove(k)
        except ValueError: pass
    _V45_TEXT_CACHE[key] = text
    _V45_TEXT_CACHE_ORDER.append(key)
    while len(_V45_TEXT_CACHE_ORDER) > _V45_TEXT_CACHE_MAX:
        old_key = _V45_TEXT_CACHE_ORDER.pop(0)
        _V45_TEXT_CACHE.pop(old_key, None)
    return text, False


def _v45_compare_worker(self, paths, base_index):
    try:
        names=[]; texts=[]; cache_hits=0
        for k,p in enumerate(paths):
            _check_cancel(self.cancel_event)
            path=Path(p)
            text,hit=_v45_read_path_cached(path)
            cache_hits += int(hit)
            names.append(path.name); texts.append(text)
            suffix=' · 캐시' if hit else ''
            self._progress_from_worker(2+int(8*(k+1)/len(paths)),f'{path.name} 읽기 완료{suffix}')
        result=compare_documents(names,texts,base_index=base_index,progress_cb=self._progress_from_worker,cancel_event=self.cancel_event)
        result['perf_info']={'text_cache_hits':cache_hits,'text_cache_total':len(paths),'fast_similarity':_RapidFuzzIndel is not None}
        self.root.after(0,lambda:self._compare_done(result))
    except ComparisonCancelled:
        self.root.after(0,self._compare_cancelled)
    except Exception as e:
        logging.exception('comparison failed'); self.root.after(0,lambda:self._compare_failed(str(e)))

NativeGui._compare_worker = _v45_compare_worker



# ---------------- V5.5: sequential three-document comparison + revised-style Word export ----------------
# Three-document comparison is chronological by default: A↔B and B↔C.  A↔C is an
# explicit optional overlay.  The selected baseline still controls the physical row axis,
# but it no longer changes which revision pairs are visualized.


def _v55_pair_plan(doc_count: int, base_index: int=0, include_ac: bool=False):
    if doc_count <= 1:
        return []
    if doc_count == 2:
        old_i = 0 if base_index == 0 else 1
        new_i = 1 - old_i
        return [(old_i, new_i, 'A↔B', 0)]
    plan=[(0,1,'A↔B',0),(1,2,'B↔C',1)]
    if include_ac:
        plan.append((0,2,'A↔C',2))
    return plan


def _v55_pairwise_segments(members: List[Optional[Unit]], base_index: int, attr: str='body', include_ac: bool=False):
    """Build document-local red/blue spans from the selected pair plan.

    In three-document mode B is both the revised side of A↔B and the original side of B↔C.
    Text that is introduced in B and removed again in C is marked ``both`` so the UI can
    render that intermediate-only wording without silently losing either relationship.
    """
    n=len(members)
    tokens=[]; flags=[]
    for m in members:
        txt=(getattr(m,attr) or '') if m else ''
        tt=display_tokens(txt)
        tokens.append(tt)
        flags.append([set() for _ in tt])

    for old_i,new_i,pair,_order in _v55_pair_plan(n,base_index,include_ac):
        old=members[old_i] if old_i<n else None
        new=members[new_i] if new_i<n else None
        if old is None and new is None:
            continue
        if old is None:
            for f in flags[new_i]: f.add('insert')
            continue
        if new is None:
            for f in flags[old_i]: f.add('delete')
            continue
        if attr=='header' and _v34_header_equivalent(old,new):
            continue
        old_text=getattr(old,attr) or ''; new_text=getattr(new,attr) or ''
        ot,om,nt,nm=_directional_pair_marks_v19(old_text,new_text)
        # display_tokens() is deterministic; nevertheless clamp to protect malformed inputs.
        for k,sty in enumerate(om[:len(flags[old_i])]):
            if sty=='delete': flags[old_i][k].add('delete')
        for k,sty in enumerate(nm[:len(flags[new_i])]):
            if sty=='insert': flags[new_i][k].add('insert')

    result=[]
    for d,tt in enumerate(tokens):
        styles=[]
        for f in flags[d]:
            if 'delete' in f and 'insert' in f: styles.append('both')
            elif 'delete' in f: styles.append('delete')
            elif 'insert' in f: styles.append('insert')
            else: styles.append('normal')
        result.append(_merge_styled_tokens_v19(tt,styles))
    return result


def _v55_pair_structural_messages(members: List[Optional[Unit]], base_index: int, include_ac: bool=False):
    out=[]
    for old_i,new_i,pair,_order in _v55_pair_plan(len(members),base_index,include_ac):
        old=members[old_i] if old_i<len(members) else None
        new=members[new_i] if new_i<len(members) else None
        if old is None and new is None:
            continue
        if old is not None and new is None:
            out.append(f'{pair} · 조/블록 삭제: {old.header or old.title or old.number or "본문"}')
            continue
        if old is None and new is not None:
            out.append(f'{pair} · 조/블록 추가: {new.header or new.title or new.number or "본문"}')
            continue
        assert old is not None and new is not None
        if old.kind=='article' and new.kind=='article' and old.number!=new.number:
            out.append(f'{pair} · 조 번호/위치 변경: {old.header} → {new.header}')
        elif old.header and new.header and normalize_text(old.header)!=normalize_text(new.header) and _v34_header_equivalent(old,new):
            out.append(f'{pair} · 제목/번호 표기 변경: {old.header} → {new.header}')
        if old.section!=new.section and (old.section or new.section):
            out.append(f'{pair} · 편제 변경: {old.section or "없음"} → {new.section or "없음"}')
        if old.norm_body!=new.norm_body:
            detail=_v30_subpart_changes(old.body,new.body,limit=30)
            for line in detail.get('lines',[]):
                if ('항/호 번호·위치 변경' in line or '항/호 구조 변경' in line or '문장 위치 이동' in line):
                    msg=f'{pair} · {line}'
                    if msg not in out: out.append(msg)
    return out


_V518_STRUCTURAL_NUMBER_RE = re.compile(
    r'^(?:[①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳]|\(\d+\)|\d+[.)]|[가-하A-Za-z][.)])$',
    re.UNICODE,
)


def _v518_is_structural_number_event(event: Dict[str, Any]) -> bool:
    """True when a numbered marker is itself an item/enumerator change such as ``(8)``."""
    txt = re.sub(r'\s+', '', str(event.get('text') or ''))
    if _V518_STRUCTURAL_NUMBER_RE.fullmatch(txt):
        return True
    msg = str(event.get('message') or '')
    # Messages use curly quotes in the active engine path.  Keep ASCII quote support for
    # older cached results/tests as well.
    m = re.search(r'(?:추가|삭제|변경):\s*[“"]([^”"]+)[”"]\s*(?:→\s*[“"]([^”"]+)[”"])?', msg)
    if not m:
        return False
    vals = [re.sub(r'\s+', '', g) for g in m.groups() if g]
    return bool(vals) and all(_V518_STRUCTURAL_NUMBER_RE.fullmatch(v) for v in vals)


def _v55_pair_events(members: List[Optional[Unit]], base_index: int, include_ac: bool=False):
    events=[]
    for old_i,new_i,pair,pair_order in _v55_pair_plan(len(members),base_index,include_ac):
        old=members[old_i] if old_i<len(members) else None
        new=members[new_i] if new_i<len(members) else None
        if old is None or new is None:
            continue
        pair_events=_v31_pair_marker_events(old,new,old_i,new_i)
        for e in pair_events:
            e['pair']=pair
            e['pair_order']=pair_order
            e['label']=pair
            sk=tuple(e.get('sort_key') or ())
            # Older marker engines begin sort_key with relative document index.  Pair order is
            # now the stable first key, so A↔B always precedes B↔C, then optional A↔C.
            e['sort_key']=(pair_order,) + (sk[1:] if len(sk)>1 else sk)
            e['structural_number']=_v518_is_structural_number_event(e)
            events.append(e)
    events.sort(key=lambda e:e.get('sort_key',()))
    for n,e in enumerate(events,1):
        e['num']=n
        e.pop('sort_key',None)
    return events


def _v55_marker_label(n: int) -> str:
    return f'[{int(n)}]'


def _v55_excel_segments_for_parts(parts, base_index: int, include_ac: bool=False):
    units=[_v44_part_unit(p.get('text',''),p.get('label','')) if p else None for p in parts]
    return _v55_pairwise_segments(units,base_index,'body',include_ac)


def _v55_excel_part_messages(parts, base_index: int, include_ac: bool=False):
    units=[_v44_part_unit(p.get('text',''),p.get('label','')) if p else None for p in parts]
    events=_v55_pair_events(units,base_index,include_ac)
    lines=[f"{_v55_marker_label(e['num'])} {e.get('pair','')} · {e.get('message','')}" for e in events]
    lines.extend('• '+m for m in _v55_pair_structural_messages(units,base_index,include_ac))
    return lines or ['변경 없음']


_compare_documents_v54 = compare_documents

def compare_documents(names: List[str], texts: List[str], base_index: int=0, progress_cb=None, cancel_event=None,
                      include_ac: bool=False) -> Dict[str,Any]:
    result=_compare_documents_v54(names,texts,base_index=base_index,progress_cb=progress_cb,cancel_event=cancel_event)
    n=len(names)
    if n not in (2,3):
        return result
    # V5.15: marker/change numbers are LOCAL to each aligned paragraph/block.
    # _v55_pair_events already numbers one row [1]..[N], so the next row intentionally
    # restarts from [1].  UI routing combines row id + local number to avoid collisions.
    for row in result.get('rows',[]):
        members=[Unit(**m) if m else None for m in (row.get('members') or [])]
        while len(members)<n: members.append(None)

        # Replace baseline-to-all coloring with the same pair plan used by the markers.
        row['header_segments']=_v55_pairwise_segments(members,base_index,'header',include_ac)
        row['body_segments']=_v55_pairwise_segments(members,base_index,'body',include_ac)
        full=[]
        for i,m in enumerate(members):
            if not m:
                full.append([]); continue
            hs=row['header_segments'][i] or [(m.header,'normal')]
            full.append(hs+[('\n','normal')]+row['body_segments'][i])
        row['segments']=full

        events=_v55_pair_events(members,base_index,include_ac)
        row['markers']=events
        structural=_v55_pair_structural_messages(members,base_index,include_ac)
        display=[f"{_v55_marker_label(e['num'])} {e.get('pair','')} · {e.get('message','')}" for e in events]
        display.extend('• '+m for m in structural)
        if not display:
            display=['변경 없음']
        row['display_messages']=display
        if isinstance(row.get('summary'),dict):
            row['summary']['messages']=display
            row['summary']['tags']=['동일'] if display==['변경 없음'] else ['변경']
        row['changed']=display!=['변경 없음']

    result['include_ac']=bool(include_ac and n==3)
    result['comparison_pairs']=[p[2] for p in _v55_pair_plan(n,base_index,bool(include_ac and n==3))]
    # Recount changed rows after replacing the old baseline-to-all comparison presentation.
    if isinstance(result.get('counts'),dict):
        result['counts']['changed']=sum(1 for r in result.get('rows',[]) if r.get('changed'))
    return result


# Word export: use the final/revised DOCX package as the template.  This preserves its
# styles.xml, theme, margins/sections, headers/footers and custom style definitions.  The
# body is rebuilt with GUI-engine tracked changes, so Accept All yields the revised text.
def _v55_capture_revised_styles(doc: Document):
    exact={}; paragraphs=[]
    for p in doc.paragraphs:
        txt=(p.text or '').strip()
        if not txt: continue
        style_name=''
        try: style_name=p.style.name if p.style is not None else ''
        except Exception: style_name=''
        key=normalize_text(txt)
        if key and key not in exact: exact[key]=style_name
        paragraphs.append((txt,key,style_name))
    return exact,paragraphs


def _v55_pick_style(text: str, exact, paragraphs, fallback: str='Normal') -> str:
    txt=(text or '').strip(); key=normalize_text(txt)
    if key in exact and exact[key]: return exact[key]
    if txt:
        # Body units can contain several source paragraphs.  Prefer a revised paragraph whose
        # full text occurs inside the emitted unit, or whose opening phrase matches it.
        best=None
        for ptxt,pkey,sty in paragraphs:
            if not sty: continue
            if ptxt in txt or txt.startswith(ptxt) or (len(txt)>=24 and ptxt.startswith(txt[:24])):
                score=min(len(ptxt),len(txt))
                if best is None or score>best[0]: best=(score,sty)
        if best: return best[1]
    return fallback


def _v55_clear_body_keep_section(doc: Document):
    body=doc._body._element
    for child in list(body):
        # Preserve final section properties so page setup and header/footer references from
        # the revised document survive the generated redline.
        if child.tag.endswith('}sectPr'):
            continue
        body.remove(child)


def _v55_add_chunks_paragraph(doc: Document, chunks, revstate, bold=False, style_name: str|None=None):
    p=doc.add_paragraph()
    if style_name:
        try: p.style=style_name
        except Exception: pass
    for kind,text in chunks:
        _v36_append_chunk(p,kind,text,revstate,bold=bold)
    return p


def _v55_build_tracked_doc(original_path: Path, revised_path: Path, output_path: Path, revised_author: str):
    from datetime import datetime, timezone
    original_path=Path(original_path); revised_path=Path(revised_path); output_path=Path(output_path)
    old_text=read_document(original_path.name,original_path.read_bytes())
    new_text=read_document(revised_path.name,revised_path.read_bytes())
    pair=compare_documents([original_path.name,revised_path.name],[old_text,new_text],base_index=0)
    aseq,bseq,acont,bcont,row_by_id=_v36_pair_nodes(pair,old_text,new_text)

    if revised_path.suffix.lower()=='.docx':
        doc=Document(str(revised_path))
        exact_styles,revised_paragraphs=_v55_capture_revised_styles(doc)
        _v55_clear_body_keep_section(doc)
    else:
        doc=Document(); exact_styles={}; revised_paragraphs=[]
    _v36_enable_track_revisions(doc)
    revstate={'id':1,'author':str(revised_author or revised_path.stem),
              'date':datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace('+00:00','Z')}
    sm=SequenceMatcher(None,aseq,bseq,autojunk=False)

    def add_para(chunks, revised_text='', bold=False):
        style=_v55_pick_style(revised_text,exact_styles,revised_paragraphs,'Normal') if revised_text else 'Normal'
        return _v55_add_chunks_paragraph(doc,chunks,revstate,bold=bold,style_name=style)

    def emit_ident(ident, mode):
        kind=ident[0]
        if kind=='P':
            oldv=acont.get(ident,''); newv=bcont.get(ident,'')
            chunks=_v36_chunks_from_hunks(oldv,newv,False) if mode=='matched' else ([('delete',oldv)] if mode=='delete' else [('insert',newv)])
            if any(t for _,t in chunks): add_para(chunks,newv if mode!='delete' else '',False)
        elif kind=='S':
            oldv=acont.get(ident,''); newv=bcont.get(ident,'')
            chunks=_v36_chunks_from_hunks(oldv,newv,False) if mode=='matched' else ([('delete',oldv)] if mode=='delete' else [('insert',newv)])
            if any(t for _,t in chunks): add_para(chunks,newv if mode!='delete' else '',True)
        else:
            row=row_by_id[ident[1]]; ms=row.get('members') or [None,None]
            old_u=ms[0] if len(ms)>0 else None; new_u=ms[1] if len(ms)>1 else None
            if mode=='delete':
                if old_u:
                    add_para([('delete',old_u.get('header',''))],'',True)
                    if old_u.get('body'): add_para([('delete',old_u.get('body',''))],'',False)
                return
            if mode=='insert':
                if new_u:
                    add_para([('insert',new_u.get('header',''))],new_u.get('header',''),True)
                    if new_u.get('body'): add_para([('insert',new_u.get('body',''))],new_u.get('body',''),False)
                return
            if not old_u or not new_u: return
            oh=old_u.get('header',''); nh=new_u.get('header','')
            hc=[('normal',oh)] if _v34_header_equivalent(Unit(**old_u),Unit(**new_u)) else _v36_chunks_from_hunks(oh,nh,False)
            add_para(hc,nh,True)
            bc=_v36_body_chunks(old_u.get('body',''),new_u.get('body',''))
            if bc or old_u.get('body') or new_u.get('body'):
                add_para(bc,new_u.get('body',''),False)

    for tag,a1,a2,b1,b2 in sm.get_opcodes():
        if tag=='equal':
            for ident in aseq[a1:a2]: emit_ident(ident,'matched')
        elif tag=='delete':
            for ident in aseq[a1:a2]: emit_ident(ident,'delete')
        elif tag=='insert':
            for ident in bseq[b1:b2]: emit_ident(ident,'insert')
        else:
            for ident in aseq[a1:a2]: emit_ident(ident,'delete')
            for ident in bseq[b1:b2]: emit_ident(ident,'insert')

    output_path.parent.mkdir(parents=True,exist_ok=True)
    doc.save(str(output_path))
    return output_path


def create_word_tracked_compare(original_path: Path, revised_path: Path, output_path: Path,
                                revised_author: str='Revised', word_factory=None) -> Path:
    """Create GUI-engine tracked changes using the revised document as the style template."""
    return _v55_build_tracked_doc(Path(original_path),Path(revised_path),Path(output_path),revised_author)



# ---------------- V5.19.0: semantic no-op guard ----------------
# A visible text comparator must not report canonically identical Unicode as a change.
# Keep punctuation differences meaningful, but normalize canonical Unicode, ordinary whitespace
# classes and non-printing format controls that have no review value.
_V5189_IGNORABLE_FORMAT = {
    '\u200b',  # ZERO WIDTH SPACE
    '\u200c',  # ZERO WIDTH NON-JOINER
    '\u200d',  # ZERO WIDTH JOINER
    '\u2060',  # WORD JOINER
    '\ufeff',  # BOM / ZERO WIDTH NO-BREAK SPACE
}


def _v5189_canonical_visible_text(value: str) -> str:
    value=unicodedata.normalize('NFC', value or '')
    value=''.join(ch for ch in value if ch not in _V5189_IGNORABLE_FORMAT and unicodedata.category(ch) not in ('Cf',))
    # Review text treats all whitespace families as ordinary spacing.  Punctuation itself is
    # intentionally NOT compatibility-normalized, so visually similar punctuation code points
    # remain reportable by the existing punctuation-diff logic.
    value=re.sub(r'\s+', ' ', value).strip()
    return value


def _v5189_semantic_equal(a: str, b: str) -> bool:
    return _v5189_canonical_visible_text(a) == _v5189_canonical_visible_text(b)


# ---------------- V5.6: paired replacement markers ----------------
# Phrase replacements must be one logical event/number.  In prose without explicit
# legal item markers the mature V3.4 engine can emit one delete event followed by one
# insert event for the same SequenceMatcher hunk.  Coalesce those two physical edges
# here so A/B (or B/C) share the same marker number and the summary says old -> new.
_v56_previous_pair_marker_events = _v31_pair_marker_events


def _v56_same_hunk(a: Dict[str,Any], b: Dict[str,Any]) -> bool:
    if a.get('part') != b.get('part'):
        return False
    ka=tuple(a.get('sort_key') or ()); kb=tuple(b.get('sort_key') or ())
    if len(ka)!=5 or len(kb)!=5:
        return False
    # V3.4: (doc, part, hunk, role, normalized_position)
    if isinstance(ka[4], float) or isinstance(kb[4], float):
        return ka[:3]==kb[:3] and ka[3]==0 and kb[3]==1
    # V3.5 fine diff: (doc, part, item, hunk, role)
    return ka[:4]==kb[:4] and ka[4]==0 and kb[4]==1


def _v56_coalesce_replacements(events: List[Dict[str,Any]], base_index: int, other_index: int):
    out=[]; i=0
    while i < len(events):
        old=events[i]
        if (i+1 < len(events) and old.get('action')=='삭제' and
                events[i+1].get('action')=='추가' and _v56_same_hunk(old,events[i+1])):
            new=events[i+1]
            old_text=str(old.get('text') or '').strip(); new_text=str(new.get('text') or '').strip()
            # Never turn canonically/visually identical text into a replacement event.
            # This specifically blocks false summaries such as “석” → “석”.
            if _v5189_semantic_equal(old_text, new_text):
                i += 2
                continue
            # Keep truly detached punctuation as separate events; lexical substitutions are
            # much easier to review as one replacement.
            lexical=any(ch.isalnum() or ('가'<=ch<='힣') for ch in old_text+new_text)
            if old_text and new_text and lexical:
                # Preserve an explicit item label such as [(2)] when the older message has it.
                prefix=''
                mm=re.match(r'^(\[[^\]]+\])\s*',str(old.get('message') or ''))
                if mm and not re.match(r'^\[\d+\]$',mm.group(1)):
                    prefix=mm.group(1)+' '
                merged=dict(new)
                merged.update({
                    'relative_doc': other_index,
                    'target_doc': other_index,
                    'action': '변경',
                    'text': new_text,
                    'endpoints': [
                        {'target_doc': base_index,
                         'char_start': int(old.get('char_start',0)),
                         'char_end': int(old.get('char_end',old.get('char_start',0)))},
                        {'target_doc': other_index,
                         'char_start': int(new.get('char_start',0)),
                         'char_end': int(new.get('char_end',new.get('char_start',0)))},
                    ],
                    'sort_key': tuple(old.get('sort_key') or ()),
                    'old_text': old_text,
                    'new_text': new_text,
                    'message': f'{prefix}변경: “{re.sub(r"\s+", " ", old_text).strip()}” → “{re.sub(r"\s+", " ", new_text).strip()}”',
                })
                out.append(merged); i += 2; continue
        out.append(old); i += 1
    return out


def _v56_pair_marker_events(base: Unit, other: Unit, base_index: int, other_index: int):
    events=_v56_previous_pair_marker_events(base,other,base_index,other_index)
    return _v56_coalesce_replacements(events,base_index,other_index)


_v31_pair_marker_events = _v56_pair_marker_events
APP_VERSION = "5.7-engine"

# ---------------- V5.7: structural item lineage guard ----------------
_V57_QUOTED_HEAD_RE = re.compile(r'^\s*["“‘]([^"”’]{1,96})["”’]')

def _v57_definition_head(part) -> str:
    core=str(part.get('core') or '').strip()
    m=_V57_QUOTED_HEAD_RE.match(core)
    return normalize_text(m.group(1)) if m else ''

def _v57_match_parts(ap, bp):
    used_a=set(); used_b=set(); pairs=[]
    def take(i,j,score):
        if i in used_a or j in used_b: return False
        pairs.append((i,j,score)); used_a.add(i); used_b.add(j); return True

    # 1) Exact content is definitive, even if item number/location changed.
    by_text=defaultdict(list)
    for j,b in enumerate(bp):
        n=normalize_text(b.get('core',''))
        if n: by_text[n].append(j)
    for i,a in enumerate(ap):
        n=normalize_text(a.get('core',''))
        if not n: continue
        cand=[j for j in by_text.get(n,[]) if j not in used_b]
        if cand: take(i,min(cand,key=lambda j:abs(i-j)),1.0)

    # 2) Definition head is an identity key: "Company" stays Company, "Member" stays Member.
    by_head=defaultdict(list)
    for j,b in enumerate(bp):
        if j in used_b: continue
        h=_v57_definition_head(b)
        if h: by_head[h].append(j)
    for i,a in enumerate(ap):
        if i in used_a: continue
        h=_v57_definition_head(a)
        if not h: continue
        cand=[j for j in by_head.get(h,[]) if j not in used_b]
        if cand:
            j=max(cand,key=lambda j:(_v35_part_similarity(a,bp[j]),-abs(i-j)))
            take(i,j,max(.90,_v35_part_similarity(a,bp[j])))

    # 3) Same enumerator is the normal legal anchor and outranks generic fuzzy cross-number text.
    for i,a in enumerate(ap):
        if i in used_a or a.get('label')=='본문': continue
        cand=[]
        for j,b in enumerate(bp):
            if j in used_b or b.get('label')!=a.get('label'): continue
            cand.append((_v35_part_similarity(a,b),j))
        if cand:
            ss,j=max(cand)
            if ss>=0.24:
                take(i,j,max(ss,.55))

    # 4) Residual moves need strong evidence; conflicting definition heads may never match.
    cands=[]
    for i,a in enumerate(ap):
        if i in used_a: continue
        ha=_v57_definition_head(a)
        wa=set(token_words(a.get('core','')))
        for j,b in enumerate(bp):
            if j in used_b: continue
            hb=_v57_definition_head(b)
            if ha and hb and ha!=hb: continue
            ss=_v35_part_similarity(a,b)
            wb=set(token_words(b.get('core','')))
            shared=len(wa & wb)
            contain=shared/max(1,min(len(wa),len(wb)))
            if ss>=0.72 and (shared>=2 or contain>=0.58):
                cands.append((ss,contain,-abs(i-j),i,j))
    for ss,contain,_dist,i,j in sorted(cands,reverse=True):
        take(i,j,ss)
    return pairs,used_a,used_b

_v35_match_parts = _v57_match_parts

def _v57_subpart_changes(old: str, new: str, limit: int=14):
    ap0=_v30_split_explicit_parts(old); bp0=_v30_split_explicit_parts(new)
    if len(ap0)==1 and len(bp0)==1 and ap0[0][0]=='본문' and bp0[0][0]=='본문':
        lines=_v30_compact_ops(ap0[0][1],bp0[0][1],limit=limit)
        return {'lines':lines,'count':len(lines),'moves':0,'changes':len(lines),'truncated':0}
    a_struct=any(_v30_structural_label(l) for l,_ in ap0)
    b_struct=any(_v30_structural_label(l) for l,_ in bp0)
    if a_struct != b_struct:
        lines=['항/호 구조 변경']
        lines.extend(_v30_compact_ops(old,new,limit=max(0,limit-1))[:max(0,limit-1)])
        return {'lines':lines,'count':len(lines),'moves':1,'changes':max(0,len(lines)-1),'truncated':0}
    ap=[{'label':l,'core':t} for l,t in ap0]
    bp=[{'label':l,'core':t} for l,t in bp0]
    pairs,used_a,used_b=_v57_match_parts(ap,bp)
    lines=[]; move_count=0; change_count=0
    for i,j,ss in sorted(pairs,key=lambda x:x[0]):
        la,ta=ap0[i]; lb,tb=bp0[j]
        moved=(i!=j) or (la!=lb and (_v30_structural_label(la) or _v30_structural_label(lb)))
        if moved:
            move_count+=1
            if len(lines)<limit:
                if _v30_structural_label(la) or _v30_structural_label(lb):
                    lines.append(f'항/호 번호·위치 변경: {la} → {lb}')
                else:
                    lines.append('문장 위치 이동')
        if normalize_text(ta)!=normalize_text(tb):
            ops=_v30_compact_ops(ta,tb,limit=4)
            if ops: change_count+=1
            for d in ops:
                if len(lines)>=limit: break
                prefix=f'[{la}→{lb}] ' if la!=lb else ('' if la=='본문' else f'[{la}] ')
                lines.append(prefix+d)
    for i,(la,ta) in enumerate(ap0):
        if i not in used_a:
            change_count+=1
            if len(lines)<limit: lines.append(f'[{la}] 삭제 “{_short_v15(ta,150)}”')
    for j,(lb,tb) in enumerate(bp0):
        if j not in used_b:
            change_count+=1
            if len(lines)<limit: lines.append(f'[{lb}] 추가 “{_short_v15(tb,150)}”')
    total=move_count+change_count
    return {'lines':lines,'count':total,'moves':move_count,'changes':change_count,'truncated':max(0,total-len(lines))}

_v30_subpart_changes = _v57_subpart_changes

# Keep the item-number movement visible in the same marker message.
_v57_previous_pair_marker_events = _v31_pair_marker_events

def _v57_pair_marker_events(base: Unit, other: Unit, base_index: int, other_index: int):
    events=_v57_previous_pair_marker_events(base,other,base_index,other_index)
    old_parts=_v35_parts_with_spans(base.body or '')
    new_parts=_v35_parts_with_spans(other.body or '')
    pairs,_,_=_v57_match_parts(old_parts,new_parts)
    moved_labels={i:(old_parts[i].get('label',''),new_parts[j].get('label','')) for i,j,_ in pairs
                  if old_parts[i].get('label') not in ('','본문') and new_parts[j].get('label') not in ('','본문')
                  and old_parts[i].get('label')!=new_parts[j].get('label')}
    for e in events:
        sk=tuple(e.get('sort_key') or ())
        # Matched body hunks have (doc, 1, source_item, hunk, role). Unmatched add/delete
        # events use hunk=999 and must never inherit a moved item's label by index coincidence.
        source_item=(sk[2] if e.get('part')=='body' and len(sk)>=4 and sk[1]==1
                     and isinstance(sk[2],int) and sk[3]!=999 else None)
        if source_item in moved_labels:
            la,lb=moved_labels[source_item]
            msg=re.sub(r'^\[[^\]]+\]\s*','',str(e.get('message') or ''))
            e['message']=f'항/호 {la}→{lb} · {msg}'
    return events

_v31_pair_marker_events = _v57_pair_marker_events
APP_VERSION = "5.7-engine"



# ---------------- V5.8: review layout support + punctuation option ----------------
# The Avalonia UI can suppress punctuation-only edits without changing the source text.
# Lexical edits keep punctuation attached to the phrase; only standalone punctuation changes
# disappear when the option is off.
ACTIVE_INCLUDE_PUNCTUATION = True


def _v58_is_punctuation_only(value: str) -> bool:
    chars=[ch for ch in str(value or '') if not ch.isspace()]
    return bool(chars) and all(not ch.isalnum() and ch!='_' and not ('가'<=ch<='힣') for ch in chars)


_v58_previous_directional_pair_marks = _directional_pair_marks_v19

def _v58_directional_pair_marks(base_text: str, other_text: str):
    at,am,bt,bm=_v58_previous_directional_pair_marks(base_text,other_text)
    if globals().get('ACTIVE_INCLUDE_PUNCTUATION',True):
        return at,am,bt,bm
    for toks,marks in ((at,am),(bt,bm)):
        for i,tok in enumerate(toks):
            if marks[i] != 'normal' and _v58_is_punctuation_only(tok):
                marks[i]='normal'
    return at,am,bt,bm

_directional_pair_marks_v19 = _v58_directional_pair_marks


_v58_previous_pair_events = _v55_pair_events

def _v55_pair_events(members: List[Optional[Unit]], base_index: int, include_ac: bool=False):
    events=_v58_previous_pair_events(members,base_index,include_ac)
    if globals().get('ACTIVE_INCLUDE_PUNCTUATION',True):
        return events
    filtered=[]
    for e in events:
        txt=str(e.get('text') or '')
        msg=str(e.get('message') or '')
        # Pure punctuation insert/delete/change events are omitted.  A lexical replacement
        # containing punctuation is retained because the punctuation belongs to that phrase.
        if _v58_is_punctuation_only(txt):
            continue
        if ('기호 추가' in msg or '기호 삭제' in msg or '기호 변경' in msg or '기호 위치 변경' in msg) and not any(ch.isalnum() for ch in txt):
            continue
        filtered.append(e)
    for n,e in enumerate(filtered,1):
        e['num']=n
    return filtered


_v58_previous_pairwise_segments = _v55_pairwise_segments

def _v55_pairwise_segments(members: List[Optional[Unit]], base_index: int, attr: str='body', include_ac: bool=False):
    result=_v58_previous_pairwise_segments(members,base_index,attr,include_ac)
    if globals().get('ACTIVE_INCLUDE_PUNCTUATION',True):
        return result
    out=[]
    for segs in result:
        cleaned=[]
        for text,style in segs:
            if style!='normal' and _v58_is_punctuation_only(text):
                style='normal'
            if cleaned and cleaned[-1][1]==style:
                cleaned[-1]=(cleaned[-1][0]+text,style)
            else:
                cleaned.append((text,style))
        out.append(cleaned)
    return out


# Insertion of a new item before an existing item changes its list index but does not mean
# that an unchanged (2) became another (2).  Keep only real enumerator changes such as
# (2) -> (10); suppress the misleading '(2) -> (2)' structural note.
_v58_previous_subpart_changes = _v30_subpart_changes

def _v58_subpart_changes(old: str, new: str, limit: int=14):
    result=_v58_previous_subpart_changes(old,new,limit=limit)
    lines=[]; removed=0
    same_re=re.compile(r'^항/호 번호·위치 변경:\s*(\S+)\s*→\s*\1\s*$')
    for line in result.get('lines',[]):
        if same_re.match(str(line).strip()):
            removed+=1
            continue
        lines.append(line)
    result=dict(result)
    result['lines']=lines
    result['moves']=max(0,int(result.get('moves',0))-removed)
    result['count']=max(0,int(result.get('count',len(lines)))-removed)
    return result

_v30_subpart_changes = _v58_subpart_changes
APP_VERSION = "5.8-engine"



# ---------------- V5.9: paired zero-width anchor markers ----------------
# Insert/delete changes are one-sided text edits, but visually they still need the same
# change number on BOTH compared documents.  Add a zero-width counterpart endpoint at the
# corresponding unchanged boundary in the other document.  Avalonia renders that endpoint
# as a marker badge only (no phantom text highlight), so e.g. an inserted "(signature space)"
# in B gets the same [n] badge in A at the insertion anchor.
_v59_previous_pair_marker_events = _v31_pair_marker_events


def _v59_map_boundary(src: str, dst: str, pos: int) -> int:
    """Map a zero-width counterpart marker to a stable textual boundary.

    Character-level SequenceMatcher can align a deleted phrase against the middle of a
    neighbouring word.  For example, deleting ``but not specified`` after ``used`` once
    mapped to the middle of ``used`` in the revised document (``us[74]ed``).  When the
    source position is already between words, map by matched lexical neighbours first so
    the counterpart badge lands at the same *word boundary*.  Edits genuinely inside one
    word still use the older character-level mapping.
    """
    src = src or ''
    dst = dst or ''
    pos = max(0, min(int(pos), len(src)))

    def word_char(ch: str) -> bool:
        return bool(ch) and (ch.isalnum() or ch == '_' or ('가' <= ch <= '힣'))

    inside_word = 0 < pos < len(src) and word_char(src[pos - 1]) and word_char(src[pos])
    if not inside_word:
        src_spans = _v30_lex_spans(src)
        dst_spans = _v30_lex_spans(dst)
        if src_spans and dst_spans:
            sk = [normalize_text(t[0]) for t in src_spans]
            dk = [normalize_text(t[0]) for t in dst_spans]
            token_matcher = SequenceMatcher(None, sk, dk, autojunk=False)
            mapped_tokens = {}
            for block in token_matcher.get_matching_blocks():
                for k in range(block.size):
                    mapped_tokens[block.a + k] = block.b + k

            left = [i for i, (_, _a, b) in enumerate(src_spans) if b <= pos and i in mapped_tokens]
            right = [i for i, (_, a, _b) in enumerate(src_spans) if a >= pos and i in mapped_tokens]
            li = max(left, key=lambda i: src_spans[i][2], default=None)
            ri = min(right, key=lambda i: src_spans[i][1], default=None)
            if li is not None and ri is not None:
                dl = mapped_tokens[li]; dr = mapped_tokens[ri]
                left_end = dst_spans[dl][2]
                right_start = dst_spans[dr][1]
                if dl <= dr and left_end <= right_start:
                    # A deletion/insertion between two surviving words belongs immediately
                    # after the surviving left context, before any separating whitespace.
                    return max(0, min(len(dst), left_end))
            if li is not None:
                return max(0, min(len(dst), dst_spans[mapped_tokens[li]][2]))
            if ri is not None:
                return max(0, min(len(dst), dst_spans[mapped_tokens[ri]][1]))

    # Fallback for spelling/suffix edits that really occur inside a word.
    sm = SequenceMatcher(None, src, dst, autojunk=False)
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == 'equal' and i1 <= pos <= i2:
            return max(0, min(len(dst), j1 + min(pos - i1, j2 - j1)))
        if i1 <= pos <= i2:
            if i2 == i1:
                return max(0, min(len(dst), j1))
            frac = (pos - i1) / max(1, i2 - i1)
            mapped = round(j1 + frac * (j2 - j1))
            return max(0, min(len(dst), mapped))
        if pos < i1:
            return max(0, min(len(dst), j1))
    return len(dst)


def _v59_pair_marker_events(base: Unit, other: Unit, base_index: int, other_index: int):
    events = _v59_previous_pair_marker_events(base, other, base_index, other_index)
    old_parts = {
        'header': base.header or '',
        'body': base.body or '',
    }
    new_parts = {
        'header': other.header or '',
        'body': other.body or '',
    }
    old_offsets = {'header': 0, 'body': len(base.header or '') + 1}
    new_offsets = {'header': 0, 'body': len(other.header or '') + 1}

    for e in events:
        if e.get('endpoints'):
            continue
        action = str(e.get('action') or '')
        if action not in ('추가', '삭제'):
            continue
        part = str(e.get('part') or 'body')
        if part not in ('header', 'body'):
            continue

        start = int(e.get('char_start', 0))
        end = int(e.get('char_end', start))
        if action == '추가':
            # Actual text exists in the revised/other document.  Map its start boundary back
            # to the old/base document and add a zero-width visual anchor there.
            local_new = max(0, start - new_offsets[part])
            old_anchor = _v59_map_boundary(new_parts[part], old_parts[part], local_new)
            e['endpoints'] = [
                {'target_doc': base_index,
                 'char_start': old_offsets[part] + old_anchor,
                 'char_end': old_offsets[part] + old_anchor},
                {'target_doc': other_index,
                 'char_start': start,
                 'char_end': end},
            ]
        else:
            # Actual text exists in the old/base document.  Map its start boundary forward
            # to the revised document and add a zero-width visual anchor there.
            local_old = max(0, start - old_offsets[part])
            new_anchor = _v59_map_boundary(old_parts[part], new_parts[part], local_old)
            e['endpoints'] = [
                {'target_doc': base_index,
                 'char_start': start,
                 'char_end': end},
                {'target_doc': other_index,
                 'char_start': new_offsets[part] + new_anchor,
                 'char_end': new_offsets[part] + new_anchor},
            ]
    return events


_v31_pair_marker_events = _v59_pair_marker_events

# V5.19.0 final guard: the marker source is the single source of truth for GUI summaries.
# Filter semantic no-ops HERE so A/B/C badges and the right-hand change list stay consistent.
_v5189_previous_pair_marker_events = _v31_pair_marker_events

def _v5189_pair_marker_events(base: Unit, other: Unit, base_index: int, other_index: int):
    # If a row differs only by Unicode decomposition / ignorable formatting, it is unchanged.
    if (_v5189_semantic_equal(base.header or '', other.header or '') and
            _v5189_semantic_equal(base.body or '', other.body or '')):
        return []
    events=_v5189_previous_pair_marker_events(base,other,base_index,other_index)
    out=[]
    for e in events:
        if str(e.get('action') or '') == '변경':
            old_text=e.get('old_text')
            new_text=e.get('new_text')
            if old_text is not None and new_text is not None and _v5189_semantic_equal(str(old_text),str(new_text)):
                continue
            # Compatibility with older event producers that only encoded old/new in message.
            if old_text is None or new_text is None:
                m=re.search(r'변경:\s*[“"](.*?)[”"]\s*→\s*[“"](.*?)[”"]\s*$', str(e.get('message') or ''))
                if m and _v5189_semantic_equal(m.group(1),m.group(2)):
                    continue
        out.append(e)
    return out

_v31_pair_marker_events = _v5189_pair_marker_events
# ---------------- V5.19.4: continuous edit decoration across internal spaces ----------------
# SequenceMatcher can align repeated whitespace tokens independently from the surrounding
# inserted/deleted words.  That is useful for diff matching, but visually it used to leave
# tiny gaps in an otherwise continuous underline/strikethrough phrase.  Preserve the exact
# source text and offsets, but inherit the edit style for whitespace that sits BETWEEN two
# changed segments belonging to a compatible edit direction.
_v5194_previous_pairwise_segments = _v55_pairwise_segments


def _v5194_style_roles(style: str):
    if style == 'insert': return {'insert'}
    if style == 'delete': return {'delete'}
    if style == 'both': return {'insert','delete'}
    return set()


def _v5194_join_internal_edit_whitespace(segs):
    items=[list(x) for x in (segs or [])]
    for i,(txt,style) in enumerate(items):
        if style != 'normal' or not txt or not txt.isspace():
            continue
        li=i-1
        while li>=0 and not items[li][0]: li-=1
        ri=i+1
        while ri<len(items) and not items[ri][0]: ri+=1
        if li<0 or ri>=len(items):
            continue
        shared=_v5194_style_roles(items[li][1]) & _v5194_style_roles(items[ri][1])
        if shared == {'insert'}:
            items[i][1]='insert'
        elif shared == {'delete'}:
            items[i][1]='delete'
        elif shared == {'insert','delete'}:
            items[i][1]='both'

    merged=[]
    for txt,style in items:
        if not txt: continue
        if merged and merged[-1][1] == style:
            merged[-1]=(merged[-1][0]+txt,style)
        else:
            merged.append((txt,style))
    return merged


def _v55_pairwise_segments(members: List[Optional[Unit]], base_index: int, attr: str='body', include_ac: bool=False):
    result=_v5194_previous_pairwise_segments(members,base_index,attr,include_ac)
    return [_v5194_join_internal_edit_whitespace(segs) for segs in result]


APP_VERSION = "5.19.4-engine"

# The entry point MUST be last so every version patch above is installed before the GUI starts.
if __name__=='__main__':
    try: main()
    except Exception as exc:
        logging.exception('fatal GUI error')
        try:
            tk,ttk,filedialog,messagebox=_safe_import_tk(); r=tk.Tk(); r.withdraw(); messagebox.showerror('문서 비교기 실행 오류',f'{exc}\n\n로그: {LOG_PATH}'); r.destroy()
        except Exception: pass
        raise
